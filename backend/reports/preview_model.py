# -*- coding: utf-8 -*-
"""把导出的 xlsx 字节流转成界面能画的网格（导出 Excel 可视化）。

为什么要它：用户 2026-10-16 要求「把导出 Excel 可视化做进系统中」。
成果文件是发给甲方的，界面上却只能"下载—打开 Excel—对着看"，出一次错要来回好几趟。

【为什么在后端出数据、而不是前端塞一个 SheetJS/Univer】
· 前端目前只有 antd/axios/dayjs/echarts，为预览多背 1MB 依赖不划算（还要打进免安装包）；
· 更关键的是**校验**：预览数据由后端生成时，可以顺手拿预警规则去比对「控制值」那些格子，
  界面才能说"这一列 26 行全部等于规则里的 35/5"，而不只是"画个表给人看"。
  前端自己解析文件就做不到这件事（它不知道系统规则、也不知道哪些格子是控制值）。

【只画"看得见的东西"】
· 合并单元格给出跨行跨列，前端按 colSpan/rowSpan 还原；
· 数字按单元格的**显示格式**输出文本（0.0 就是 2.0 而不是 2），所见即打印所见；
· 公式格导出件里没有算好的缓存值（文件是程序生成的，Excel 从没打开过），
  所以原样带上公式文本并标成公式，由界面注明"这一格由 Excel 打开后计算"。
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

#: 预览的规模上限：母版整卷也就几十行 × 十几列，超了就是异常文件，截断别把界面拖死
MAX_ROWS = 220
MAX_COLS = 30
#: 单元格文本截断长度（说明段落很长，全塞进 JSON 没意义）
_MAX_TEXT = 180
_DECIMALS_RE = re.compile(r"^0\.(0+)")


def _decimals(fmt: str | None) -> int | None:
    """从 Excel 数字格式里读小数位数；读不出来（General/日期）返回 None = 原样显示。

    只处理本仓真正用到的这几种：`General`、`0.0`、`0.00`、`0.000`、`0.00000_ `。
    """
    if not fmt or fmt == "General":
        return None
    m = _DECIMALS_RE.match(fmt)
    return len(m.group(1)) if m else None


def _display(value, fmt: str | None) -> tuple[str, bool]:
    """(显示文本, 是否数值)。数值按显示格式补小数位 —— 报表上 2.0 就该看着是 2.0。"""
    if value is None:
        return "", False
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), False
    if isinstance(value, (int, float)):
        d = _decimals(fmt)
        if d is None:
            # General 格式：Excel 最多显 11 位有效数字，超了才转科学计数。
            # 用 `:g`（只有 6 位）会把坐标 3413090.672 印成 `3.41309e+06` —— 看着像数据坏了。
            return (f"{value:.11g}" if isinstance(value, float) else str(value)), True
        return f"{value:.{d}f}", True
    if isinstance(value, str) and value.startswith("="):
        return value, False
    # 日期：openpyxl 读文件时就把日期格换算成了 datetime（导出写的是序列 46255），
    # 这里只把它印回人看得懂的样子 —— 预览里出现一片五位数，用户会以为日期写坏了。
    if isinstance(value, datetime):
        return (value.strftime("%Y-%m-%d") if not (value.hour or value.minute)
                else value.strftime("%Y-%m-%d %H:%M")), False
    if isinstance(value, date):
        return str(value), False
    return str(value), False


def _cell_payload(cell) -> dict | None:
    v, is_num = _display(cell.value, cell.number_format)
    if v == "":
        return None
    formula = isinstance(cell.value, str) and cell.value.startswith("=")
    out: dict = {"r": cell.row, "c": cell.column, "v": v[:_MAX_TEXT]}
    if is_num:
        out["n"] = True
    if formula:
        out["f"] = True
    font = cell.font
    if font is not None and font.b:
        out["b"] = True
    fill = cell.fill
    if fill is not None and fill.patternType == "solid":
        rgb = getattr(getattr(fill, "fgColor", None), "rgb", None)
        if isinstance(rgb, str) and rgb not in ("00000000", "FFFFFFFF"):
            out["bg"] = rgb[-6:]
    al = cell.alignment
    if al is not None:
        if al.horizontal:
            out["ha"] = al.horizontal
        if al.wrap_text:
            out["wrap"] = True
    if len(v) > _MAX_TEXT:
        out["cut"] = True
    return out


def workbook_to_preview(wb) -> dict:
    """工作簿 → 预览网格（稀疏单元格 + 合并区 + 行列规模）。"""
    sheets: list[dict] = []
    for ws in wb.worksheets:
        max_r = min(ws.max_row or 1, MAX_ROWS)
        max_c = min(ws.max_column or 1, MAX_COLS)
        cells: list[dict] = []
        for row in ws.iter_rows(min_row=1, max_row=max_r, min_col=1, max_col=max_c):
            for cell in row:
                payload = _cell_payload(cell)
                if payload:
                    cells.append(payload)
        merges = [
            {"ref": str(rng), "r": rng.min_row, "c": rng.min_col,
             "rs": rng.max_row - rng.min_row + 1, "cs": rng.max_col - rng.min_col + 1}
            for rng in ws.merged_cells.ranges
            if rng.min_row <= max_r and rng.min_col <= max_c
        ]
        widths = {k: round(v.width, 2) for k, v in (ws.column_dimensions or {}).items()
                  if v.width and _col_to_int(k) and _col_to_int(k) <= max_c}
        sheets.append({
            "name": ws.title,
            "rows": max_r,
            "cols": max_c,
            "cells": cells,
            "merges": merges,
            "widths": widths,
            "state": ws.sheet_state,  # 隐藏页也带着，界面上注明"这张表在 Excel 里是隐藏的"
            # 截断是**每张表各自**的事（哪张超了一目了然），别用一句全局提示糊过去
            "truncated": (ws.max_row or 0) > MAX_ROWS or (ws.max_column or 0) > MAX_COLS,
        })
    return {"sheets": sheets, "truncated": any(s["truncated"] for s in sheets)}


def _col_to_int(letters: str) -> int | None:
    from openpyxl.utils import column_index_from_string

    try:
        return column_index_from_string(letters)
    except Exception:  # noqa: BLE001
        return None


def bytes_to_preview(raw: bytes) -> dict:
    """xlsx 字节流 → 预览网格。"""
    wb = load_workbook(io.BytesIO(raw), data_only=False)
    try:
        return workbook_to_preview(wb)
    finally:
        wb.close()


# ─────────────────────── 预览附带的「控制值」核对 ───────────────────────


def check_control_values(raw: bytes, layout: dict | None, rules: dict[str, tuple]) -> list[dict]:
    """逐格核对报表上的「控制值」是不是**系统预警规则**里的那对数。

    这是预览真正值钱的地方：光画表格，用户还是得自己对着规则核；
    这里直接把每个 Sheet 的数据区扫一遍，不一致的格子点名出来。

    判据与导出端同一个（`template_check_service.data_area_rows`），免得出现「预览说没问题、
    文件里却留着旧值」这种两头各说一套的检查。
    数据区界线之外还印着数字的格子单独列成 residual —— 那是母版遗留、导出够不着的部分。
    """
    from backend.services.template_check_service import data_area_rows

    wb = load_workbook(io.BytesIO(raw), data_only=False)
    try:
        return _check_on(wb, layout, rules, data_area_rows)
    finally:
        wb.close()


def _check_on(wb, layout: dict | None, rules: dict[str, tuple], data_area_rows) -> list[dict]:
    out: list[dict] = []
    for scfg in (layout or {}).get("sheets") or []:
        sn = scfg.get("sheet_name")
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        f = scfg.get("layout") or {}
        cols = [(k, f.get(k)) for k in ("control_cum", "control_rate")]
        cols = [(k, c) for k, c in cols if c]
        if not cols:
            continue
        kind = scfg.get("kind") or ("displacement" if "水平" in sn else "settlement")
        cum, rate = rules.get(kind, (None, None))
        expected = {"control_cum": cum, "control_rate": rate}
        lo, hi = data_area_rows(ws, scfg)
        # 数据区界线是**按行**算的，碰到横向合并的大块（说明/结论/签字）会整块穿过去。
        # 十号基坑实测：导出把母版那句「变化量正负号说明」清掉了，于是 A31:L33 这块的
        # 左上格里再也找不到"点号列第一次出现文字"的行，界线从 30 漂到 33 ——
        # 那三行的 I/J 本就不是独立的格子（Excel 显示不出、导出也写不进），
        # 报成「控制值空白」是假警报，得跳过。
        def covered(r: int, col: int) -> bool:
            return isinstance(ws.cell(row=r, column=col), MergedCell)

        mismatch: list[str] = []
        blank: list[str] = []
        residual: list[str] = []
        rows_ok: set[int] = set()
        for key, col in cols:
            want = expected[key]
            for r in range(lo, hi + 1):
                if covered(r, col):
                    continue
                rows_ok.add(r)
                cell = ws.cell(row=r, column=col)
                v = cell.value
                if v is None or (isinstance(v, str) and not v.strip()):
                    # 只配了一项时那一格**就该留空**（缺的不兜底），反过来规则有值却空着才是问题
                    if want is not None:
                        blank.append(f"{sn}!{cell.coordinate}")
                    continue
                if want is None:
                    mismatch.append(f"{sn}!{cell.coordinate}={v}（规则里没配这一项）")
                    continue
                try:
                    if abs(float(v) - float(want)) > 1e-6:
                        mismatch.append(f"{sn}!{cell.coordinate}={v}（规则是 {want:g}）")
                except (TypeError, ValueError):
                    mismatch.append(f"{sn}!{cell.coordinate}={v}（不是数）")
            for r in range(hi + 1, (ws.max_row or hi) + 1):
                cell = ws.cell(row=r, column=col)
                if isinstance(cell.value, (int, float)):
                    residual.append(f"{sn}!{cell.coordinate}={cell.value}")
        out.append(
            {
                "sheet": sn,
                "kind": kind,
                "expected": expected,
                "area": [lo, hi],
                "control_cols": dict(cols),
                "checked_rows": len(rows_ok),
                "ok": not mismatch and not blank and not residual,
                "mismatch": mismatch[:20],
                "blank": blank[:20],
                "residual": residual[:20],
            }
        )
    return out
