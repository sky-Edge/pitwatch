# -*- coding: utf-8 -*-
"""母版导出：把"母版里没有的测点"自动补进来（2026-09-30 用户需求）。

【问题】母版是拿别的基坑的报表当底模的（只有 8 个点位行），而项目有 12 个测点。
`fill_data_sheets` 是**按母版固有行位**逐点驱动的，多出来的 4 个点从来没被写进去过 ——
导出文件里整条消失，而且**没有任何提示**。

【用户确认的做法】自动补行 + 下移说明行：
- 接着母版最后一个点位行往下补，行格式沿用母版的点位行（边框/行高/控制值）；
- 位置不够时，把下面的「变化量说明 / 监测结论 / 签字栏」整块往下移；
- 反过来（母版点位比项目多）**保留空行**：点号照写、数据留空。

【为什么手写搬行，不用 `ws.insert_rows()`】
openpyxl 的 `insert_rows` 只搬单元格的值/样式，**不搬合并单元格、不搬行高、不挪数据验证**。
实测母版第 34~37 行是签字栏（含 `A35:B35`、`A37:F37`、`G37:L37` 三处合并），
直接用它会把签字栏的合并拆散 —— 所以这里自己写：从下往上搬，值+样式+行高+合并区一起走。
"""
from __future__ import annotations

from copy import copy

from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.cell_range import CellRange


def tail_start_row(ws, point_col: int, after_row: int, limit: int = 200) -> int:
    """数据区之后「说明 / 结论 / 签字」块的起始行。

    判据：点位行之后**点号列第一次出现文字**的那一行（母版里是「变化量正负号说明：」）。
    中间那些只有控制值、或完全空白的行，都是母版预留的空白点位行 —— 可以拿来放新点。
    找不到就返回 `max_row + 1`（整张表下面都是空的）。
    """
    end = min(ws.max_row or after_row, after_row + limit)
    for r in range(after_row + 1, end + 1):
        v = ws.cell(row=r, column=point_col).value
        if v is not None and str(v).strip():
            return r
    return end + 1


def _style_source(ws, row: int, col: int):
    """取样式来源格。

    ⚠ 踩过的坑：**合并区的从属格（MergedCell）不持有样式** —— 样式全在左上角那一格。
    母版里位移表每个点的两行是把点号列纵向合并的（`A21:A22`），
    照搬"原型行"的样式时如果拿到 `A22` 这种从属格，就会把边框复制成 None，
    新补的行看上去缺边框。所以这里遇到合并格就回到左上角去取。
    """
    cell = ws.cell(row=row, column=col)
    if not isinstance(cell, MergedCell):
        return cell
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(row=rng.min_row, column=rng.min_col)
    return cell


def vertical_merge_of(ws, row: int, col: int) -> tuple[int, int] | None:
    """(row, col) 所在的那个**同列纵向合并区**，没有就返回 None。

    用来判断母版有没有"点号跨行合并"的惯例（位移表 X/Y 两行共用点号）。
    """
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            if rng.min_col == rng.max_col == col and rng.max_row > rng.min_row:
                return (rng.min_row, rng.max_row)
    return None


def _copy_style(src, dst) -> None:
    if isinstance(dst, MergedCell) or isinstance(src, MergedCell):
        return
    if src.has_style:
        dst._style = copy(src._style)  # noqa: SLF001 —— openpyxl 没有公开的"整份样式"接口


def copy_row_style(ws, src_row: int, dst_row: int, max_col: int) -> None:
    """把一整行的**样式**（边框/字体/对齐/数字格式）和行高复制到另一行（不复制值）。"""
    for c in range(1, max_col + 1):
        _copy_style(_style_source(ws, src_row, c), ws.cell(row=dst_row, column=c))
    h = ws.row_dimensions[src_row].height if src_row in ws.row_dimensions else None
    if h is not None:
        ws.row_dimensions[dst_row].height = h


def plan_raw_slots(raw: dict, bottom_row: int, count: int, rows_per_block: int = 3) -> list[tuple[str, int]]:
    """给"原始数据区"里多出来的测点找空位 → [(left/right, start_row), ...]。

    这一区和上面的数据区不一样：它是母版里**两列 × 每点 3 行（X/Y/Z）**的固定网格。
    母版来自别的基坑时只填了 8 个块（左列 6、右列 2），但**右列下面往往已经预留好了
    带边框的空格**（实测母版右列 row 10~27 全有细边框）—— 12 个点正好塞得下，不用挪动任何东西。

    排布沿用母版的习惯：**先把左列排满，再排右列**（母版本身就是左 6 右 2）。
    `bottom_row` 是网格底部行（两列里最低的那个块的最后一行），保证新块不越出网格、
    不压到下面的「变化量说明 / 监测结论 / 签字栏」。

    放不下就少返回几个 —— 调用方**必须**把它们报出去（绝不静默丢点）。
    """
    slots: list[tuple[str, int]] = []
    if count <= 0 or bottom_row <= 0:
        return slots
    left = list(raw.get("left") or [])
    right = list(raw.get("right") or [])
    if not left and not right:
        return slots
    # 列对齐基准：母版两列的第一块是齐平的
    head = min(e["start_row"] for e in (left or right))

    for side, entries in (("left", left), ("right", right)):
        start = (max(e["start_row"] for e in entries) + rows_per_block) if entries else head
        while len(slots) < count and start + rows_per_block - 1 <= bottom_row:
            slots.append((side, start))
            start += rows_per_block
    return slots[:count]


def shift_rows_down(ws, first_row: int, count: int) -> None:
    """把 [first_row, max_row] 整块下移 count 行：值、样式、行高、合并单元格一起搬。

    顺序要紧：**先解除合并再搬值**（合并区里的从属格是只读的），搬完在目标位置重建合并；
    搬的时候**从下往上**，否则会自己覆盖自己。
    """
    if count <= 0:
        return
    last = ws.max_row or first_row
    max_col = ws.max_column or 1

    # ① 落在这个范围里的合并区，先解除（记下原位置）
    moved: list[str] = []
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= first_row and rng.max_row <= last:
            moved.append(str(rng))
    for ref in moved:
        ws.unmerge_cells(ref)

    # ② 从下往上搬单元格（值 + 样式）与行高
    for r in range(last, first_row - 1, -1):
        height = ws.row_dimensions[r].height if r in ws.row_dimensions else None
        for c in range(1, max_col + 1):
            src = ws.cell(row=r, column=c)
            dst = ws.cell(row=r + count, column=c)
            if isinstance(dst, MergedCell):
                continue
            if isinstance(src, MergedCell):
                dst.value = None
                continue
            dst.value = src.value
            if src.has_style:
                dst._style = copy(src._style)  # noqa: SLF001
        if height is not None:
            ws.row_dimensions[r + count].height = height
        # 源行清空（值搬走了，样式留着给新点位行用）
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            if not isinstance(cell, MergedCell):
                cell.value = None

    # ③ 在原位置 + count 处重建合并区
    for ref in moved:
        rng = CellRange(ref)
        rng.shift(row_shift=count)
        ws.merge_cells(str(rng))
