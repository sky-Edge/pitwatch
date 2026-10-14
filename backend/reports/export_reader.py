# -*- coding: utf-8 -*-
"""把**导出的日报表**读回来 —— 与导出（`reports/master_report.py`）严格镜像的读法。

用途：「文件检查」要把用户导出去的那份 Excel 和系统里的数比一比，
所以必须能**准确还原**导出时写进去的东西。这里每一条读法都对应导出端的一条写法：

| 读 | 对应导出端 |
|---|---|
| 数据区从 `data_start_row` 起、**行首点号列为空即结束** | `master_report._master_point_blocks`（同一判据）|
| 竖向位移每点 1 行；水平位移每点 2 行（次行点号列空）| `master_report.fill_data_sheets` 的 `is_main` 分支 |
| 列号走 `export_template_service.map_columns` | 同一函数（列判据唯一出处）|
| 工程名称取「标签后的最后一段」 | `master_report._extract_master_project_name`（同一套标签）|
| 日期取「本次监测时间：YYYY-MM-DD」 | `master_report._replace_date_label`（ISO 格式）|

⚠ 这里**不做任何业务判断**（不知道"对不对"），只把文件里的数摊平。
比对逻辑在 `services/export_audit_service.py`。
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from backend.services.export_template_service import is_data_sheet_name, map_columns

#: 导出文件名形如「十号基坑第24期2026.10.09.xlsx」—— 内容认不出来时靠它兜底
_FILENAME_DATE = re.compile(r"(20\d{2})[.\-/](\d{1,2})[.\-/](\d{1,2})")
#: 表里「本次监测时间：2026-10-09」/「2026.10.9」都认
_TEXT_DATE = re.compile(r"(20\d{2})\s*[.\-/年]\s*(\d{1,2})\s*[.\-/月]\s*(\d{1,2})")
#: 表头/名称行只在这些行里找就够了（导出端也只看前几行）
_HEAD_SCAN_ROWS = 25
#: 数据区最多往下看多少行（导出端是 start+60，这里给宽一点）
_MAX_DATA_ROWS = 120


@dataclass
class ReportRow:
    """报表上一行（= 一个测点的一个分量）里我们关心的那几个数。"""

    point_no: str
    component: str  # value / X / Y
    excel_row: int
    prev_cum_mm: float | None = None
    cum_mm: float | None = None
    delta_mm: float | None = None
    rate_mm_per_day: float | None = None
    control_cum_mm: float | None = None
    control_rate_mm_per_day: float | None = None


@dataclass
class ReportSheet:
    sheet_name: str
    kind: str  # settlement / displacement
    mode: str  # per_row / x_y_two_rows
    header_row: int
    columns: dict
    rows: list[ReportRow] = field(default_factory=list)


@dataclass
class ReadReport:
    """一份解析出来的导出日报。"""

    path: Path
    filename: str
    #: 表里出现的工程名称候选（"工程名称：兴善公路 十号基坑" → ["兴善公路 十号基坑", "十号基坑"]）
    name_candidates: list[str] = field(default_factory=list)
    obs_date: str | None = None
    date_source: str = ""  # text / filename
    sheets: list[ReportSheet] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)


def _text(v) -> str:
    return "" if v is None else str(v)


def _num(v) -> float | None:
    """单元格 → 数值。文本型数字（"1.2"）也认；空/非数字返回 None。"""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if not s or s in {"-", "—", "/", "无"}:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _row_texts(ws, row: int, limit_cols: int = 40) -> list[str]:
    return [_text(ws.cell(row=row, column=c).value) for c in range(1, limit_cols + 1)]


def _find_header(ws) -> tuple[int, list[str], list[str]] | None:
    """找出含「点号」的表头行，返回 (行号, 该行文本, 下一行文本)。"""
    max_row = ws.max_row or 1
    for r in range(1, min(10, max_row + 1)):
        cells = _row_texts(ws, r)
        if any("点号" in t for t in cells):
            nxt = _row_texts(ws, r + 1) if r + 1 <= max_row else []
            return r, cells, nxt
    return None


def _name_candidates(wb) -> list[str]:
    """表里的工程名称候选。

    与 `master_report._extract_master_project_name` 用**同一套标签**取法：
    标签后按空白切分、最后一段是名字（前面是"兴善公路"这类地名前缀）。
    这里额外把"整段"也收进来 —— 检查时是拿去跟**已知项目名**对，
    多给几个候选没有副作用，漏给才会认不出。
    """
    from backend.reports.master_report import NAME_LABELS

    out: list[str] = []
    for ws in wb.worksheets:
        max_row = ws.max_row or 1
        for row in ws.iter_rows(min_row=1, max_row=min(_HEAD_SCAN_ROWS, max_row)):
            for cell in row:
                t = _text(cell.value)
                if not t or not t.strip():
                    continue
                for lb in NAME_LABELS:
                    i = t.find(lb)
                    if i < 0:
                        continue
                    rest = t[i + len(lb):].strip()
                    if rest:
                        out.append(rest)
                        parts = [p for p in re.split(r"[\s\u3000]+", rest) if p]
                        if parts and parts[-1] != rest:
                            out.append(parts[-1].strip("（）()【】[]"))
                    break
    # 去重保序
    seen: set[str] = set()
    uniq: list[str] = []
    for c in out:
        c = c.strip()
        if c and c not in seen:
            seen.add(c)
            uniq.append(c)
    return uniq


def _obs_date(wb) -> tuple[str | None, str]:
    """从表里读「本次监测时间」；读不到返回 (None, "")。"""
    for ws in wb.worksheets:
        max_row = ws.max_row or 1
        for row in ws.iter_rows(min_row=1, max_row=min(6, max_row)):
            for cell in row:
                t = _text(cell.value)
                if "本次监测时间" not in t:
                    continue
                m = _TEXT_DATE.search(t)
                if m:
                    y, mo, d = (int(x) for x in m.groups())
                    try:
                        return f"{y:04d}-{mo:02d}-{d:02d}", "text"
                    except ValueError:
                        pass
    return None, ""


def _date_from_filename(name: str) -> str | None:
    m = _FILENAME_DATE.search(name)
    if not m:
        return None
    y, mo, d = (int(x) for x in m.groups())
    return f"{y:04d}-{mo:02d}-{d:02d}"


def _is_note_label(s: str) -> bool:
    """像不像"数据区下面的说明/结论行"。

    实测（2026-10-12）：两张数据表的数据区末尾都有
    「变化量正负号说明：…」「监测结论及建议：」「本次监测时间：…」三行，
    它们的文字落在 A 列（点号列）。不认出来就会当成测点去查，报一堆假的"点号不存在"。
    """
    return ("：" in s) or (":" in s) or len(s) > 12


#: 用来判断"这一行到底有没有值"的列（控制值也算 —— 母版套打时每个点位行都会写控制值）
_VALUE_KEYS = (
    "v0_m",
    "prev_m",
    "current_m",
    "prev_cum_mm",
    "cum_mm",
    "delta_mm",
    "rate",
    "control_cum",
    "control_rate",
)


def _has_any_value(ws, row: int, columns: dict) -> bool:
    for key in _VALUE_KEYS:
        col = columns.get(key)
        if col and ws.cell(row=row, column=col).value is not None:
            return True
    return False


def _read_sheet_rows(ws, columns: dict, is_displacement: bool) -> list[ReportRow]:
    """按导出端的**同一判据**走数据区。

    位移表一个测点可能占两行（X 行 + Y 行），而**两种导出实现的写法不一样**：
      · 母版套打（`master_report`）：只在 X 行写点号，Y 行点号列留空；
      · 系统标准版式（`excel_report`）：**每行都写点号**；同一点 X/Y 都存在时，
        是两行**同名**点号。
    所以配对判据是「下一行点号列为空 **或** 与当前行相同」。

    ⚠ 但"空点号"有两种含义，必须分清（第一版就栽在这儿）：
      · 本块的 Y 行（母版套打那种）；
      · **数据区与说明行之间的空行** —— 标准版式在数据区末尾留了一整行空白，
        再往下是「变化量正负号说明：…」等三行。把空行当成 Y 行吞掉，
        下一轮就会把说明行读成测点，报出一堆假的"点号不存在"。
      区分办法：这一行**整行都没值**、而且**再下一行也是空的或像说明** → 那是分隔行，数据区到此为止。

    结束条件还有一条：点号列的文字像说明（含冒号/过长）→ 直接停。
    """
    pc = columns.get("point_col") or 1
    start = columns.get("data_start_row") or 7
    max_row = ws.max_row or start
    end = min(max_row, start + _MAX_DATA_ROWS)

    def _label_at(r: int) -> str:
        v = ws.cell(row=r, column=pc).value
        return "" if v is None else str(v).strip()

    def _row_is_blank(r: int) -> bool:
        return _label_at(r) == "" and not _has_any_value(ws, r, columns)

    out: list[ReportRow] = []
    r = start
    while r <= end:
        point_no = _label_at(r)
        if not point_no or _is_note_label(point_no):
            break  # 空行 / 说明行 → 数据区结束
        rows = [r]
        if is_displacement and r + 1 <= end:
            nxt = _label_at(r + 1)
            if nxt == point_no and nxt:
                rows.append(r + 1)  # 标准版式：同一点写在相邻两行（X/Y）
            elif nxt == "":
                after = _label_at(r + 2) if r + 2 <= end else ""
                after_blank = r + 2 > end or _row_is_blank(r + 2)
                is_separator = _row_is_blank(r + 1) and (after_blank or _is_note_label(after))
                if not is_separator:
                    rows.append(r + 1)  # 母版套打：Y 行点号留空
        for i, rr in enumerate(rows):
            comp = "value" if not is_displacement else ("X" if i == 0 else "Y")
            out.append(
                ReportRow(
                    point_no=point_no,
                    component=comp,
                    excel_row=rr,
                    prev_cum_mm=_num(ws.cell(row=rr, column=columns["prev_cum_mm"]).value) if columns.get("prev_cum_mm") else None,
                    cum_mm=_num(ws.cell(row=rr, column=columns["cum_mm"]).value) if columns.get("cum_mm") else None,
                    delta_mm=_num(ws.cell(row=rr, column=columns["delta_mm"]).value) if columns.get("delta_mm") else None,
                    rate_mm_per_day=_num(ws.cell(row=rr, column=columns["rate"]).value) if columns.get("rate") else None,
                    control_cum_mm=_num(ws.cell(row=rr, column=columns["control_cum"]).value) if columns.get("control_cum") else None,
                    control_rate_mm_per_day=_num(ws.cell(row=rr, column=columns["control_rate"]).value) if columns.get("control_rate") else None,
                )
            )
        r += len(rows)
    return out


def read_report(path: Path) -> ReadReport:
    """读一份导出的日报。**不抛异常**：读不了就写进 problems，让检查页如实说"这个文件读不了"。"""
    rep = ReadReport(path=path, filename=path.name)
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:  # noqa: BLE001
        rep.problems.append(f"打不开这个文件：{e}")
        return rep

    try:
        rep.name_candidates = _name_candidates(wb)
        rep.obs_date, rep.date_source = _obs_date(wb)
        if rep.obs_date is None:
            # 兜底：导出文件名里就带着日期（「十号基坑第24期2026.10.09.xlsx」）
            fn_date = _date_from_filename(path.name)
            if fn_date:
                rep.obs_date, rep.date_source = fn_date, "filename"

        for name in wb.sheetnames:
            if not is_data_sheet_name(name):
                continue
            ws = wb[name]
            found = _find_header(ws)
            if not found:
                continue
            header_row, top, sub = found
            columns = map_columns(top, sub, header_row)
            is_disp = ("水平" in name or "r." in name.lower()) and "竖向" not in name
            sheet = ReportSheet(
                sheet_name=name,
                kind="displacement" if is_disp else "settlement",
                mode="x_y_two_rows" if is_disp else "per_row",
                header_row=header_row,
                columns=columns,
                rows=_read_sheet_rows(ws, columns, is_disp),
            )
            if sheet.rows:
                rep.sheets.append(sheet)
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass

    if not rep.sheets:
        rep.problems.append("没找到监测数据表（不知道哪几张表是数据表）")
    if rep.obs_date is None:
        rep.problems.append("认不出观测日期（表里没写「本次监测时间」，文件名里也没有日期）")
    return rep
