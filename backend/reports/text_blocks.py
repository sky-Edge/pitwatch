# -*- coding: utf-8 -*-
"""按**本期数据**自动生成整卷日报里的文字页（概况 / 分析 / 成果提交）。

【为什么要做】母版是拿**别的工程**那期报表另存的 —— 除了版式，还带着别人的整段文字。
用户 2026-10-08 的实测：五份母版（十五号/十四号/十三号/十一号/十号）的
`概况`/`分析`/`成果提交` 三页**逐字完全相同**（文字指纹一致），内容却是另一个工程的：

    分析!C3  = "(1)地表沉降监测：…DBC05-01，DBC06-05…"      ← 用户根本没这类测点
    分析!C14 = "(12)测斜QCX1监测：…"
    成果提交!B3  = "本周通过对**屿前站**的监测数据和巡视信息综合分析"
    成果提交!B9~B20 = 12 类监测项的成果清单（用户只有竖向位移 + 水平位移两类）

也就是说：**这三页每期都在往外发别人的东西**。让用户每期手改 12 段不现实（他自己也没改过），
所以改为**按本期真实数据生成**。

【生成边界 —— 数据能说的就说，不能说的绝不编】
- ✅ 能生成：累计变化量极值、极值对应的测点、达到预警/报警指标的测点数与点号、
  本期监测的类别与点次、期次与日期。
- ❌ 不生成：**现场巡视是否发现异常**（那是人的判断，系统没有巡视数据）。
  巡视那句只写"本周对××现场巡视 1 次，巡视情况详见现场巡视记录"，不写"未发现异常"。
  同理"累积完成工作量"母版本来就是空的，也不编。

【写入口径】**先清空、再写入、多出来的格子连行一起隐藏**
- 母版有 12 段、本项目只有 2 类 → 生成的 2 段写前两格，**其余 10 格清空并隐藏行**。
  清空是最要紧的一步：不清就是在发别人的内容；隐藏是为了不留大片空白。
- 生成的段数**多于**母版格子数（以后监测类别变多）→ 塞不下的**报给用户**，绝不静默丢。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date

from openpyxl.cell.cell import MergedCell

from backend.reports.export_format import MM_DECIMALS
from backend.warnings.engine import LEVEL_ALARM, LEVEL_CONTROL

# ── 用途（存在 layout_json["text_blocks"][*]["purpose"]）──
PURPOSE_OVERVIEW = "overview"            # 分析页 4.1 现场巡视（整格：保留标题行、替换正文）
PURPOSE_ANALYSIS = "analysis_items"      # 分析页 4.2 各监测项段落（逐格一段）
PURPOSE_CONCLUSION = "conclusion"        # 成果提交页 5.1 监测结论（逐格一段）
PURPOSE_DELIVERABLES = "deliverables"    # 成果提交页 6 监测成果清单（逐条一格）
PURPOSE_PROGRESS = "progress"            # 概况页 1.3 监测完成情况（逐格一段）
PURPOSE_WORKTABLE = "worktable"          # 概况页 1.3 下方「工作量」表（行块）

#: 系统里的监测类别 → 报表里的中文名（与数据表 Sheet 名对齐）
KIND_LABEL = {"settlement": "竖向位移", "displacement": "立柱结构水平位移"}

#: 分析段落里极值保留的小数位 —— 与数据表**同一个口径**（`export_format.MM_DECIMALS`）。
#: 数据行已在 `paper_values` 收过口径，这里只负责把它排成固定 1 位小数的字符串。
_MM_DECIMALS = MM_DECIMALS


@dataclass
class GenContext:
    """生成所需的一切输入 —— 全是"本期已知事实"，此处不做任何查询。"""

    project_name: str
    obs_date: date
    period_text: str                 # "第X期" / "初始值"（调用方已算好，避免重复实现中文期次）
    is_initial: bool
    rows: list[dict] = field(default_factory=list)   # _series_rows 的结果
    kinds: list[str] = field(default_factory=list)   # 本期实际有数据的类别（有序，去重）

    @property
    def kind_labels(self) -> list[str]:
        return [KIND_LABEL.get(k, k) for k in self.kinds]


# ─────────────────────── 文本小工具 ───────────────────────


def _mm(v: float | None) -> str:
    """"-4.0" / "8.3"：极值统一 1 位小数（-0.0 要归一成 0.0）。"""
    if v is None:
        return "—"
    x = round(float(v), _MM_DECIMALS) + 0.0
    return f"{x:.{_MM_DECIMALS}f}"


def _point_label(row: dict) -> str:
    """测点标签：水平位移带 X/Y 后缀（母版里也是"点号-序号"这种写法，不带后缀分不清）。"""
    if row.get("series_type") == "displacement":
        return f"{row['point_no']}-{row.get('component', '')}"
    return str(row["point_no"])


def _with_warning(rows: list[dict]) -> list[dict]:
    return [r for r in rows if r.get("warning_level") in (LEVEL_CONTROL, LEVEL_ALARM)]


def _warning_sentence(rows: list[dict]) -> str:
    """预警描述。没有超限点 → "数据未发现异常情况。"（这是**数据结论**，可以下）。"""
    hit = _with_warning(rows)
    if not hit:
        return "数据未发现异常情况。"
    groups: list[tuple[str, list[dict]]] = [
        (word, [r for r in hit if r.get("warning_level") == level])
        for level, word in ((LEVEL_ALARM, "报警"), (LEVEL_CONTROL, "预警"))
    ]
    groups = [(w, g) for w, g in groups if g]
    if len(groups) == 1:
        # 只有一种等级 → 别写成"共 N 个:达到预警指标的有 N 个"（同义反复，报表上很啰嗦）
        word, group = groups[0]
        names = "、".join(_point_label(r) for r in group)
        return f"达到{word}指标的测点共 {len(group)} 个，点号为：{names}。"
    parts = [
        f"达到{word}指标的有 {len(group)} 个，点号为：{'、'.join(_point_label(r) for r in group)}"
        for word, group in groups
    ]
    return f"达到预警指标的测点共 {len(hit)} 个:" + "；".join(parts) + "。"


# ─────────────────────── 生成 ───────────────────────


def _gen_overview(ctx: GenContext) -> list[str]:
    """分析页 4.1 现场巡视正文。

    ⚠ 刻意**不写**"未发现异常"：系统没有巡视数据，写成"未发现异常"就是替人下结论。
    """
    return [f"        本周对{ctx.project_name}现场巡视 1 次，巡视情况详见现场巡视记录。"]


def _gen_analysis(ctx: GenContext) -> list[str]:
    """分析页 4.2：每个监测类别一段（外来的 12 段被替换成本项目实际的类别）。"""
    out: list[str] = []
    for i, kind in enumerate(ctx.kinds, start=1):
        rows = [r for r in ctx.rows if r.get("series_type") == kind]
        label = KIND_LABEL.get(kind, kind)
        usable = [r for r in rows if r.get("cum_mm") is not None]
        if not usable:
            out.append(
                f"({i}){label}监测：\n"
                f"   本期无{label}有效读数，未能给出分析结论，请核对原始记录。"
            )
            continue
        lo = min(usable, key=lambda r: r["cum_mm"])
        hi = max(usable, key=lambda r: r["cum_mm"])
        out.append(
            f"({i}){label}监测：\n"
            f"   经分析本期累积变化量介于：[{_mm(lo['cum_mm'])},{_mm(hi['cum_mm'])}]之间"
            f"(单位为：mm，点号分别为：{_point_label(lo)}，{_point_label(hi)}),"
            f"{_warning_sentence(rows)}"
        )
    return out


def _gen_conclusion(ctx: GenContext) -> list[str]:
    """成果提交页 5.1：整期一句话结论（有超限点就说有，没有就说稳定）。"""
    hit = _with_warning(ctx.rows)
    if not hit:
        body = (
            f"本周通过对{ctx.project_name}的监测数据和巡视信息综合分析："
            "本期监测数据处于稳定状态，未发现新增的异常情况。"
        )
    else:
        names = "、".join(_point_label(r) for r in hit)
        body = (
            f"本周通过对{ctx.project_name}的监测数据综合分析："
            f"本期共有 {len(hit)} 个测点达到预警指标（{names}），"
            "建议及时复核原始记录、加密监测频次并关注其变化趋势。"
        )
    return ["     " + body]


def _gen_deliverables(ctx: GenContext) -> list[str]:
    """成果提交页 6：按**本项目实际有的类别**列成果清单（母版的 12 类里没产出的一律不写）。"""
    return [f"（{i}）{label}监测日报表（共 1 次）；" for i, label in enumerate(ctx.kind_labels, start=1)]


def _gen_progress(ctx: GenContext) -> list[str]:
    """概况页 1.3 监测完成情况。第(2)句是通用流程句，沿用母版说法。"""
    labels = "、".join(ctx.kind_labels) or "本期项目"
    tail = "测点初始状态" if ctx.is_initial else ""
    day = f"{ctx.obs_date.year}年{ctx.obs_date.month}月{ctx.obs_date.day}日"
    return [
        f"（1）本周（{day}）{ctx.project_name}{labels}共监测 1 次{tail}。\n"
        f"（2）本周共巡查 1 次，监测项目工作量如下表："
    ]


def _gen_worktable(ctx: GenContext) -> list[str]:
    """概况页工作量表：**一行汇总**（母版本来就只有一行数据行，不去补行改版式）。

    "点次"按**本期实际读数条数**算（沉降每点 1 条、水平位移每点 X/Y 各 1 条）——
    这是有依据的数；"累积完成工作量"母版本来就是空的，不编。
    """
    return [
        "1",
        "、".join(ctx.kind_labels) or "—",
        f"{len(ctx.rows)}点次",
        "",                     # 累积完成工作量：无依据，留空
        "初始值" if ctx.is_initial else ctx.period_text,
    ]


_GENERATORS = {
    PURPOSE_OVERVIEW: _gen_overview,
    PURPOSE_ANALYSIS: _gen_analysis,
    PURPOSE_CONCLUSION: _gen_conclusion,
    PURPOSE_DELIVERABLES: _gen_deliverables,
    PURPOSE_PROGRESS: _gen_progress,
    PURPOSE_WORKTABLE: _gen_worktable,
}


def generate(purpose: str, ctx: GenContext) -> list[str]:
    fn = _GENERATORS.get(purpose)
    return list(fn(ctx)) if fn else []


# ─────────────────────── 写入 ───────────────────────


def _split_ref(ref: str) -> tuple[str, int, int]:
    """'分析!C3' / 'C3' → (sheet, row, col)。"""
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

    sheet, _, coord = ref.rpartition("!")
    letters, row = coordinate_from_string(coord)
    return sheet, int(row), column_index_from_string(letters)


def write_block(ws, refs: list[str], texts: list[str]) -> tuple[int, int, list[str]]:
    """把一个文字块写进 worksheet。返回 (写入格数, 清空格数, 塞不下的文本)。

    - 合并单元格的值在**左上角**，`refs` 存的就是锚点格，所以直接赋值即可；
    - 多出来的格子（母版段数 > 本项目类别数）**清空 + 隐藏所在行** ——
      不清就是在发别人的内容，不隐藏就会留一大片空白。
    """
    written = cleared = 0
    overflow: list[str] = []
    for i, ref in enumerate(refs):
        _sheet, row, col = _split_ref(ref)
        cell = ws.cell(row=row, column=col)
        if isinstance(cell, MergedCell):     # 防御：refs 只存锚点格，正常不会命中
            continue
        if i < len(texts):
            cell.value = texts[i]
            written += 1
        else:
            if cell.value is not None:
                cell.value = None
            cleared += 1
            try:
                ws.row_dimensions[row].hidden = True
            except Exception:  # noqa: BLE001
                pass
    overflow = [t for t in texts[len(refs):]]
    return written, cleared, overflow


def apply_text_blocks(wb, blocks: list[dict], ctx: GenContext) -> dict:
    """把 `layout["text_blocks"]` 里识别到的区块按本期数据改写。

    返回 {"generated": {sheet: 段数}, "cleared": {sheet: 格数}, "overflow": [sheet: 文本]}；
    没有识别到任何区块（母版里没这几种标题）时返回空 dict —— **保持原样**，
    宁可不改也不要瞎改。
    """
    summary: dict = {"generated": {}, "cleared": {}, "overflow": {}}
    for blk in blocks or []:
        purpose = blk.get("purpose")
        sheet = blk.get("sheet")
        if not purpose or not sheet or sheet not in wb.sheetnames:
            continue
        texts = generate(purpose, ctx)
        if not texts:
            continue
        ws = wb[sheet]

        if purpose == PURPOSE_WORKTABLE:
            written, cleared, overflow = _write_worktable(ws, blk, texts)
        elif purpose == PURPOSE_OVERVIEW:
            written, cleared, overflow = _write_overview(ws, blk, texts)
        else:
            written, cleared, overflow = write_block(ws, blk.get("cells") or [], texts)

        if written:
            summary["generated"][sheet] = summary["generated"].get(sheet, 0) + written
        if cleared:
            summary["cleared"][sheet] = summary["cleared"].get(sheet, 0) + cleared
        if overflow:
            summary["overflow"][sheet] = overflow
    return {k: v for k, v in summary.items() if v}


def _write_overview(ws, blk: dict, texts: list[str]) -> tuple[int, int, list[str]]:
    """巡视那块：格子里"标题 + 正文"是同一个单元格，只替换正文、标题保留。"""
    refs = blk.get("cells") or []
    if not refs:
        return 0, 0, texts
    _sheet, row, col = _split_ref(refs[0])
    cell = ws.cell(row=row, column=col)
    old = str(cell.value or "")
    keep_until = blk.get("keep_until_line_containing") or ""
    lines = old.splitlines()
    keep = 0
    if keep_until:
        keep = next((i + 1 for i, ln in enumerate(lines) if keep_until in ln), 0)
    prefix = lines[:keep]
    # 只保留标题那几行，正文整段换成生成的（母版的旧正文一律丢弃 —— 那是别的工程的）
    rest = [*prefix, *texts]
    # 去掉尾部空行，避免单元格里堆一串 \n
    while rest and not rest[-1].strip():
        rest.pop()
    cell.value = "\n".join(rest)
    return 1, 0, []


def _write_worktable(ws, blk: dict, texts: list[str]) -> tuple[int, int, list[str]]:
    """工作量表：表头不动，只写数据行（母版几行就写几行，多出来的清空）。"""
    cols = blk.get("cols") or {}
    first = blk.get("first_row")
    last = blk.get("last_row") or first
    if not (cols and first):
        return 0, 0, texts
    order = [cols.get(k) for k in ("index", "item", "week", "total", "note")]
    written = cleared = 0
    for offset, row in enumerate(range(int(first), int(last) + 1)):
        if offset == 0:
            for col, val in zip(order, texts):
                if col:
                    ws.cell(row=row, column=int(col)).value = val or None
            written += 1
        else:
            touched = False
            for col in order:
                if col and ws.cell(row=row, column=int(col)).value is not None:
                    ws.cell(row=row, column=int(col)).value = None
                    touched = True
            if touched:
                cleared += 1
    return written, cleared, []
