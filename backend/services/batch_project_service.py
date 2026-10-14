# -*- coding: utf-8 -*-
"""批量建项目：一份「点号 + X/Y/Z」清单 → 拆成多个基坑项目。

现场一次采集常常一口气测好几个基坑，仪器导出来就是一张没有表头的
「点号 + 坐标 + 高程」清单（形态见 `parser._try_parse_headerless_points` 的注释）。
这些点要分别落到各自基坑的项目里，而系统原有的 Excel 导入
**必须先有项目、且一次只导一个项目** —— 本服务补的就是这一段。

两条硬约束（都踩过，改之前先看完）：

1. **不要写 `MeasurementPoint.initial_value_mm`。**
   它是**测点级单值**，而 `calc_service.compute_point_series` 会把这**同一个值**
   传给沉降、位移 X、位移 Y 三条序列当 V0。把高程 3.119 m 换算成 3119 写进去，
   位移 X 的 V0 就变成 3119，而 X 的实际读数是 3413192929 ——
   累计变化量直接变成天文数字，而且看不出是哪一步错的。
   正确做法：**只写「第 0 期」观测记录**，让 engine 的
   「首期有效读数自动成为 V0」生效（首期累计 = 0，后续滚动正确）。

2. **点号必须允许中文。** 现场点名就是「15号1」这种。`point_service` 里那套
   批量建点的编号校验（`_POINT_NO_PATTERN`）只放行字母数字横线，这里不能复用，
   改用与导入解析同一套的 `_accept_point_no`（中文点号额外要求该行有读数）。

分组**不在本服务里做**：用户在界面上手动区分哪些点属于哪个基坑，
本服务只负责按分好的组落库。
"""
from __future__ import annotations

import re
from collections import OrderedDict
from pathlib import Path
from typing import Any

from sqlalchemy.orm import Session

from backend.importers.parser import (
    KIND_LEDGER,
    KIND_POINT_LIST,
    _accept_point_no,
    is_period_ledger,
    parse_workbook,
)
from backend.repositories import observation_repo, point_repo, project_repo
from backend.schemas.observation_record import RecordCreate
from backend.schemas.point import PointCreate
from backend.schemas.project import ProjectCreate
from backend.services import project_service

# 新建项目默认套正负号规则（与示例项目同口径，基坑习惯）：
# 沉降正值=下沉、负值=隆起；位移正值=向基坑内。
DEFAULT_SIGN_CONFIG = {
    "settlement": {"positive_means": "down"},
    "displacement": {"positive_means": "toward_pit"},
}

# 点号尾部的序号：15号1 → 15号，JJC12 → JJC
_TAIL_INDEX = re.compile(r"[\s\-_]*\d+$")

# 传错入口（把「某一天的观测台账」传到这里建点）时的提示。
# 必须同时说清"为什么不行"和"下一步点哪里" —— 只说"格式不对"用户就只会反复换文件。
WRONG_ENTRY_MSG = (
    "这份文件是「逐期观测台账」（表里写着观测日期、填的是「本次测值」），"
    "不是「点号 + 坐标/高程」点位清单 —— 在这里建点会把几毫米的读数当成高程写成初始值。"
    "要录入某一天的观测数据：回项目列表 → 进对应项目 → 用「Excel 导入」；"
    "要在这个页面建项目/测点：改传「点号 + X/Y/Z」清单（仪器导出的点位表，或初始值清单）。"
    "手上只有这张台账又想用它建点的话，把表里的「初始值」列单独整理成清单再传。"
)


def _suggest_groups(points: list[dict]) -> list[dict]:
    """按点号尾部序号猜分组（15号1 → 15号）。

    只在**猜得有意义**时才返回：组数 ≥2，且确实把点合到了一起（组数 < 点数）。
    否则返回空 —— 宁可让用户自己分，也不要甩一堆"一组一个点"的假分组添乱。
    """
    buckets: "OrderedDict[str, list[str]]" = OrderedDict()
    for p in points:
        key = _TAIL_INDEX.sub("", p["point_no"]) or p["point_no"]
        buckets.setdefault(key, []).append(p["point_no"])
    if len(buckets) < 2 or len(buckets) >= len(points):
        return []
    return [
        {"name": name, "point_count": len(nos), "point_nos": nos}
        for name, nos in buckets.items()
    ]


def _diagnose_no_rows(path: str | Path) -> str:
    """读不到数据行时，尽量说清"为什么"，而不是甩一句通用错误。

    用户实测（2026-10-06）：拿一份**点号是纯数字**（1、2、3…）的清单来传，
    界面只说"没有读到数据行"，他完全不知道该改什么，连着试了 7 次（还换了 .xlsm）。
    失败提示必须能直接指导下一步动作，这跟"故意拒绝纯数字点号"是两件事。
    """
    import re as _re

    from openpyxl import load_workbook

    from backend.importers.parser import _accept_point_no

    default = (
        "没有从这份文件里读到数据行。请确认它是「点号 + X/Y/Z」清单"
        "（每行：点名、坐标、坐标、高程）。"
    )
    try:
        wb = load_workbook(path, data_only=True)
    except Exception:  # noqa: BLE001 —— 诊断失败就退回通用提示
        return default
    try:
        for ws in wb.worksheets:
            filled: list[str] = []
            for r in range(1, (ws.max_row or 0) + 1):
                v = ws.cell(row=r, column=1).value
                if v is None or not str(v).strip():
                    continue
                filled.append(str(v).strip())
            if len(filled) < 2:
                continue  # 这个 sheet 没内容，看下一个
            # 首行若是表头（"点号"、"X"…），先摘掉再判 —— 否则会把表头算进去，
            # 提示就从"点号是纯数字"退化成含糊的"都不能当点号"。
            if filled and (
                any(k in filled[0] for k in ("点号", "名字", "编号", "测点", "点名"))
                or filled[0].upper() in ("X", "Y", "Z")
            ):
                filled = filled[1:]
            if len(filled) < 2:
                continue
            accepted = sum(1 for t in filled if _accept_point_no(t, has_reading=True))
            if accepted > 0:
                continue  # 有可用点号 → 问题不在这，交给默认提示
            sample = "、".join(filled[:5])
            if all(_re.fullmatch(r"[+\-]?\d+(\.\d+)?", t) for t in filled):
                return (
                    f"这份文件第 1 列是**纯数字**（如 {sample}…），系统不敢把它当测点号 —— "
                    "真实台账下方常附一块「差值／坐标小表」，那种表里的数字会被误当成新测点。"
                    "请把点号改成带字母或中文的写法（例如 JJC1、15号1、钢便桥1），再传一次。"
                )
            return (
                f"第 1 列读到了 {len(filled)} 个值，但没有一个能当测点号（如 {sample}…）。"
                "测点号要求：含数字、不超过 24 个字符、不含冒号。"
            )
        return default
    finally:
        wb.close()


def parse_point_sheet(path: str | Path) -> dict:
    """解析清单 → 点列表。**不落库、不分组**（分组由用户在界面上定）。

    复用导入解析器：X/Y 谁是 X 由它按数值量级判定（本项目 X≈3413xxx、Y≈510xxx）。
    列映射一并返回，界面上会把它写出来让用户核对 —— 真有判反，当场能看见。

    只接受「点位清单」与「只填初始值列的 V0 期台账」；逐期观测台账会被挡下
    （见 WRONG_ENTRY_MSG）—— 那是「Excel 导入」入口的东西。
    """
    sheets = parse_workbook(path)
    data_sheets = [s for s in sheets if s.rows]
    if not data_sheets:
        raise ValueError(_diagnose_no_rows(path))

    # 这个页面只认两种东西：点位清单（点号 + 坐标/高程）、
    # 以及「只填了初始值列」的 V0 期台账（初始值本身就是坐标/高程）。
    # 其余（某一天的观测台账）在这里没有正确用法 —— 挑出来直接说清楚该去哪，
    # 别让它悄悄把「本次测值几毫米」当高程建出一批错测点。
    point_lists = [s for s in data_sheets if s.kind == KIND_POINT_LIST]
    baselines = [s for s in data_sheets if s.kind == KIND_LEDGER and s.is_baseline]
    usable = point_lists or baselines
    if not usable:
        raise ValueError(WRONG_ENTRY_MSG)
    first = usable[0]

    # 逐张可用的表都读进来，按点号合并。台账形态的初始值文件本来就是两张表
    # （「竖向位移」给高程、「水平位移」给 X/Y）—— 只看第一张会白丢一半点位信息。
    points: "OrderedDict[str, dict]" = OrderedDict()
    for sheet in usable:
        for row in sheet.rows:
            item = points.setdefault(
                row.point_no, {"point_no": row.point_no, "x": None, "y": None, "z": None}
            )
            if row.series_type == "settlement":
                item["z"] = row.current_mm
            elif row.component == "X":
                item["x"] = row.current_mm
            elif row.component == "Y":
                item["y"] = row.current_mm

    ordered = list(points.values())
    return {
        # 可能来自多张表（见上），都写出来，免得用户以为只读了第一张
        "sheet_name": " + ".join(s.sheet_name for s in usable),
        "column_map": first.column_map,
        "has_date": first.obs_date is not None,
        "obs_date": first.obs_date.isoformat() if first.obs_date else None,
        "points": ordered,
        "suggested_groups": _suggest_groups(ordered),
    }


def _infer_point_type(pt: dict) -> str:
    """有高程也有坐标 → both；只有一个方向 → 对应类型。"""
    has_z = pt.get("z") is not None
    has_xy = pt.get("x") is not None or pt.get("y") is not None
    if has_z and has_xy:
        return "both"
    if has_z:
        return "settlement"
    return "displacement"


def _record_specs(pt: dict, obs_date) -> list[RecordCreate]:
    """该点的第 0 期记录：Z→沉降读数，X/Y→位移坐标。单位已是 mm。"""
    specs: list[RecordCreate] = []
    if pt.get("z") is not None:
        specs.append(RecordCreate(series_type="settlement", component="value",
                                  obs_date=obs_date, value_mm=pt["z"]))
    if pt.get("x") is not None:
        specs.append(RecordCreate(series_type="displacement", component="X",
                                  obs_date=obs_date, value_mm=pt["x"]))
    if pt.get("y") is not None:
        specs.append(RecordCreate(series_type="displacement", component="Y",
                                  obs_date=obs_date, value_mm=pt["y"]))
    return specs


def commit_groups(db: Session, groups: list[dict]) -> dict:
    """按用户分好的组落库：项目（复用/新建）→ 测点 → 第 0 期记录。

    幂等取向（重复提交不会造重复数据）：
    - 项目编号已存在 → **复用**该项目，不新建、不报错；
    - 项目内已有同名测点 → 跳过并报出（**绝不覆盖初始值**：它是该点所有累计量的基准，
      改了会把历史数据一起改歪）；
    - 同点同分量同日已有记录 → 跳过该分量。
    """
    results: list[dict[str, Any]] = []
    totals = {
        "projects_created": 0,
        "projects_reused": 0,
        "points_created": 0,
        "points_skipped": 0,
        "records_created": 0,
    }

    for g in groups:
        group_name = (g.get("group_name") or "").strip()
        code = (g.get("project_code") or group_name).strip()
        name = (g.get("project_name") or code).strip()
        obs_date = g.get("obs_date")
        pts = g.get("points") or []

        entry: dict[str, Any] = {
            "group_name": group_name or code,
            "project_code": code,
            "project_name": name,
            "obs_date": obs_date.isoformat() if hasattr(obs_date, "isoformat") else obs_date,
            "created_project": False,
            "points_created": 0,
            "records_created": 0,
            "skipped": [],
        }

        if not code:
            entry["error"] = "项目编号为空，整组跳过"
            results.append(entry)
            continue
        if not pts:
            entry["error"] = "这组一个测点都没有，整组跳过"
            results.append(entry)
            continue
        if obs_date is None:
            entry["error"] = "没有指定初始值日期，整组跳过"
            results.append(entry)
            continue

        project = project_repo.get_by_code(db, code)
        if project is None:
            project = project_service.create_project(
                db,
                ProjectCreate(code=code, name=name, location=None, sign_config=DEFAULT_SIGN_CONFIG),
            )
            entry["created_project"] = True
            totals["projects_created"] += 1
        else:
            totals["projects_reused"] += 1

        entry["project_id"] = project.id
        entry["project_name"] = project.name

        existing_nos = {
            p.point_no for p in point_repo.list_by_project(db, project.id, include_inactive=True)
        }
        for pt in pts:
            no = (pt.get("point_no") or "").strip()
            if not no:
                entry["skipped"].append({"point_no": "(空)", "reason": "点号为空"})
                continue
            if no in existing_nos:
                entry["skipped"].append({"point_no": no, "reason": "该项目里已有同名测点，跳过（不改动已有初始值）"})
                continue
            if not _accept_point_no(no, has_reading=True):
                entry["skipped"].append({"point_no": no, "reason": "点号格式不合法"})
                continue
            specs = _record_specs(pt, obs_date)
            if not specs:
                entry["skipped"].append({"point_no": no, "reason": "这行没有任何读数（X/Y/Z 都空）"})
                continue
            try:
                point = point_repo.create(
                    db, project.id, PointCreate(point_no=no, point_type=_infer_point_type(pt))
                )
            except Exception as e:  # noqa: BLE001 —— 单点失败不能拖垮整组
                db.rollback()
                entry["skipped"].append({"point_no": no, "reason": f"建点失败：{e}"})
                continue
            existing_nos.add(no)
            entry["points_created"] += 1
            totals["points_created"] += 1

            # 第 0 期记录。这里刻意**不写 point.initial_value_mm**，理由见模块开头。
            for spec in specs:
                if observation_repo.get_active_by_key(
                    db, point.id, spec.series_type, spec.component, obs_date
                ):
                    continue
                observation_repo.create(db, point.id, spec, source="excel")
                entry["records_created"] += 1
                totals["records_created"] += 1

        totals["points_skipped"] += len(entry["skipped"])
        results.append(entry)

    return {"groups": results, "totals": totals}
