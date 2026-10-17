# -*- coding: utf-8 -*-
"""母版体检：导出**之前**把这份母版可能出的问题查出来（用户 2026-10-16 要求「做个模版检查吧」）。

【为什么要它】母版是拿别的基坑那期报表另存来的，系统只重写它**认得的那几行**。
2026-10-16 那次就是这样：十号母版竖向位移表 7~18 行有点号、19~26 行是空着的预留行，
可预留行也印着旧控制值 30/10 —— 导出件同一列出现两个数，用户只能靠肉眼发现。
「导出文件检查」查的是**已经导出的文件**，这里查的是**还没导出的母版**，两头都要有。

【口径】只读母版原件与 layout，**一个字节都不写回去**；判据与导出端**同一个函数**
（`_master_point_blocks` / `tail_start_row` / `report_control_values`），
免得体检说没问题、导出却出问题。

【分级】
· error：这样导出会被拦住或导不出（没配控制值规则、原件缺失、没学到数据表）；
· warn：能导出，但成果文件上会留下系统里查不到的数/别的基坑名（必须人工改母版）；
· info：导出会怎么处理这块（补行、清标签、替换工程名），说清楚让用户放心。
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from backend.models.export_template import ExportTemplate
from backend.services import point_service
from backend.services.warning_service import report_control_values

#: 一条问题的形状：code 给机器看，title/detail 直接给用户读
LEVELS = ("error", "warn", "info")


def _issue(code: str, level: str, title: str, detail: str = "", cells: list[str] | None = None) -> dict:
    return {"code": code, "level": level, "title": title, "detail": detail, "cells": cells or []}


def check_master_template(db: Session, project_id: int, tpl: ExportTemplate) -> list[dict]:
    """跑一遍体检，返回问题清单（空列表 = 没查出问题）。"""
    from backend.services.project_service import get_project

    project = get_project(db, project_id)
    issues: list[dict] = []

    if tpl.is_builtin or not tpl.file_path:
        return [_issue("builtin", "info", "内置标准版式不需要母版", "系统自己排版，没有母版残留值这回事。")]
    if not Path(tpl.file_path).exists():
        return [_issue("file_missing", "error", "母版原件不在盘上",
                       f"记录的地址找不到文件：{tpl.file_path}。请重新上传这份母版，否则导出会直接失败。")]

    layout = tpl.layout_json or {}
    sheets_cfg = layout.get("sheets") or []
    if not sheets_cfg:
        return [_issue("layout_not_learned", "error", "这份母版没学到数据表结构",
                       "上传时识别失败（或被改坏了），导出只会退回旧式模板填充，整卷版式不会更新。"
                       "建议重新上传一次母版。")]

    wb = load_workbook(tpl.file_path, data_only=False)  # data_only=False：公式也要看见
    try:
        issues += _check_control_rules(db, project_id, sheets_cfg)
        issues += _check_control_cells_reachable(wb, sheets_cfg)
        issues += _check_extra_tables(wb, sheets_cfg)
        issues += _check_point_rows(db, project_id, wb, sheets_cfg)
        # 放最后：这一步会在内存副本上真跑一遍名称替换（不落盘），前面的判读不该受它影响
        issues += _check_project_name(wb, project.name)
    finally:
        wb.close()
    return issues


# ────────────────────────────── ① 规则：导出会不会被拦 ──────────────────────────────


def _check_control_rules(db: Session, project_id: int, sheets_cfg: list[dict]) -> list[dict]:
    """项目没配控制值 → 导出在 `ensure_control_rules` 就被 422 拦下（用户 2026-10-14 定的口径）。

    体检把它提前说出来，而不是让用户点导出后只看到一个红叉。
    """
    out: list[dict] = []
    cn = {"settlement": "竖向位移（沉降）", "displacement": "水平位移"}
    for kind in {s.get("kind") for s in sheets_cfg if s.get("kind")}:
        cum, rate = report_control_values(db, project_id, kind)
        if cum is None and rate is None:
            out.append(
                _issue(
                    "no_control_rule",
                    "error",
                    f"{cn.get(kind, kind)}没有配「控制值」",
                    "报表上的「控制值」两列印的就是预警规则里的控制值，没配就没有数可印 —— "
                    "导出会被拦下。请到「预警规则」里给这类观测填上累计/速率控制值。",
                )
            )
        elif cum is None or rate is None:
            missing = "累计变化量" if cum is None else "变化速率"
            out.append(
                _issue(
                    "half_control_rule",
                    "info",
                    f"{cn.get(kind, kind)}只配了一项控制值，「{missing}」那一格会留空",
                    "缺的那项**不兜底、也不沿用母版上印的旧值**（宁可空着，也不在成果文件上留一个系统里查不到的数）。",
                )
            )
    return out


# ─────────────────────── ② 控制值格是否落在导出会重写的范围里 ───────────────────────


def data_area_rows(ws, scfg: dict) -> tuple[int, int]:
    """导出会重写「控制值」的行范围 [首, 尾]（含尾）—— 与 `fill_data_sheets` 同一个判据。

    母版的数据区由两段组成：`_master_point_blocks` 按点号认出的点位行（扫到第一个空点号就停），
    以及它停住之后、说明块之前的**预留空白行**（母版每行都印着控制值）。
    体检和预览都从这里取界线，谁也不许自己数一遍行。
    """
    from backend.reports.master_expand import tail_start_row
    from backend.reports.master_report import _master_point_blocks

    f = scfg.get("layout") or {}
    start = f.get("data_start_row") or 7
    pc = f.get("point_col") or 1
    kind = scfg.get("kind") or "settlement"
    blocks = _master_point_blocks(ws, start, f, kind == "displacement")
    last = max((r for _, rs in blocks for r in rs), default=start - 1)
    tail = tail_start_row(ws, pc, last) - 1  # 说明块之前 = 数据区
    return start, max(last, tail)


def _is_num(v) -> bool:
    return isinstance(v, (int, float)) or (isinstance(v, str) and v.strip().replace(".", "", 1).isdigit())


def _check_control_cells_reachable(wb, sheets_cfg: list[dict]) -> list[dict]:
    """母版上印着控制值的格子，有没有落在导出**够不着**的地方。

    够不着 = 那一格不会被规则刷新 = 成果文件上留着别的基坑那一期的限值。
    2026-10-16 的坑就是这个（预留空白行在识别范围之外），导出端已修；
    这条检查留着，是为了下次换个形状的母版（说明块下面还有一张表）能提前发现。
    """
    out: list[dict] = []
    for scfg in sheets_cfg:
        sn = scfg.get("sheet_name")
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        f = scfg.get("layout") or {}
        cols = [c for c in (f.get("control_cum"), f.get("control_rate")) if c]
        if not cols:
            continue
        lo, hi = data_area_rows(ws, scfg)
        stale: list[str] = []
        for col in cols:
            for r in range(7, (ws.max_row or 0) + 1):
                if lo <= r <= hi:
                    continue
                v = ws.cell(row=r, column=col).value
                if v is None or not _is_num(v):
                    continue
                # 说明块/表尾里出现的数字是别的内容（年份、桩号…），不当成控制值
                label = ws.cell(row=r, column=f.get("point_col") or 1).value
                if label and ("：" in str(label) or ":" in str(label)):
                    continue
                stale.append(f"{sn}!{ws.cell(row=r, column=col).coordinate}={v}")
        if stale:
            out.append(
                _issue(
                    "stale_control_out_of_range",
                    "warn",
                    f"「{sn}」有 {len(stale)} 个控制值格在数据区之外",
                    f"导出只重写第 {lo}~{hi} 行，这些格子留在母版上原样出去 —— "
                    "表里就会出现一个系统里查不到的限值。请把它们删掉，或把那几行并进数据区。",
                    stale[:20],
                )
            )
    return out


def _check_extra_tables(wb, sheets_cfg: list[dict]) -> list[dict]:
    """同一张 Sheet 里还有第二块「控制值」表头，或某个带控制值的 Sheet 根本没被认出。

    母版常是拿上一版报表改的，「竖向位移(2)」、表下面又接半张表都很常见。
    导出只按学到的那张表往下写，其余部分**整个不会被刷新**。
    """
    out: list[dict] = []
    known = {s.get("sheet_name"): s for s in sheets_cfg}
    for ws in wb.worksheets:
        hits = [
            (cell.row, f"{ws.title}!{cell.coordinate}")
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 250))
            for cell in row
            if cell.value is not None and "控制值" in str(cell.value)
        ]
        if not hits:
            continue
        scfg = known.get(ws.title)
        if scfg is None:
            out.append(
                _issue(
                    "sheet_not_learned",
                    "warn",
                    f"「{ws.title}」里有控制值表头，但导出认的是别的表",
                    "母版复制出了第二份数据表（表名不一样），导出不会往里写数，也不会刷新它的控制值 —— "
                    "这份文件发出去，那一页还是旧的。请把多余的表删掉，或把表名改成母版学习时用的那个。",
                    [c for _, c in hits][:20],
                )
            )
            continue
        # 已认出的表：数据起始行之后还出现「控制值」= 这张表下面又接了半张表
        start = (scfg.get("layout") or {}).get("data_start_row") or 7
        below = [c for r, c in hits if r >= start]
        if below:
            out.append(
                _issue(
                    "second_block_in_sheet",
                    "warn",
                    f"「{ws.title}」同一张表里还有第 {len(below)} 块「控制值」表头",
                    "导出按第一块的数据区往下写，下面那块不会被刷新 —— 建议一个 Sheet 只留一张数据表。",
                    below[:20],
                )
            )
    return out


# ─────────────────────────── ③ 工程名：会不会留着别的基坑名 ───────────────────────────


def _check_project_name(wb, project_name: str) -> list[dict]:
    """母版自称的基坑名（如「十一号」）会被导出端换成什么。

    用户 2026-10-08 报过：项目叫「十号基坑」，导出却是「十一号」。
    判据直接复用导出端那两个函数（取旧名 → 在**内存副本**上真跑一遍替换数处数），
    不另写一套识别，否则体检和导出又会各说各话。
    旧名认得出就一定换得掉（同一批格子、同一条正则），所以这里只有两种结果：
    会替换（报处数）/ 认不出旧名（提醒人工核对封面）。
    """
    from backend.reports.master_report import _extract_master_project_name, _sync_project_name

    old = _extract_master_project_name(wb)
    if not old:
        return [
            _issue(
                "name_not_detected",
                "info",
                "没认出母版自称的基坑名",
                "封面上的工程名称会保持母版原样。如果这份母版是从别的基坑另存的，导出后请核对封面名称；"
                "把名称写成「工程名称：某某基坑」这样标签和名字在同一格，系统才认得出这个槽。",
            )
        ]
    if old == project_name or old in project_name:
        return []
    changed = _sync_project_name(wb, project_name)  # 只改内存里这份 wb，不落盘
    return [
        _issue(
            "name_replaced",
            "info",
            f"导出会把母版上的「{old}」换成「{project_name}」（{changed} 处）",
            "封面、签字栏、结论叙述里出现的旧基坑名一起跟着换，这是正常行为。",
        )
    ]


# ─────────────────────────── ④ 点位行与项目测点是否够对得上 ───────────────────────────


def _check_point_rows(db: Session, project_id: int, wb, sheets_cfg: list[dict]) -> list[dict]:
    """母版每表有几个点位行、项目有几个测点 —— 说出来，别让"补行/清空"变成意外。"""
    from backend.reports.master_report import _master_point_blocks

    out: list[dict] = []
    pts = point_service.list_points(db, project_id, include_inactive=False)
    for scfg in sheets_cfg:
        sn = scfg.get("sheet_name")
        kind = scfg.get("kind") or "settlement"
        if sn not in wb.sheetnames:
            out.append(_issue("sheet_missing", "error", f"母版里没有「{sn}」这张表",
                              "layout 记着这个表名，但原件里找不到 —— 导出会跳过这一页。请重传母版。"))
            continue
        ws = wb[sn]
        f = scfg.get("layout") or {}
        blocks = _master_point_blocks(ws, f.get("data_start_row") or 7, f, kind == "displacement")
        tmpl_points = {p for p, _ in blocks}
        mine = {p.point_no for p in pts if p.point_type in ("both", kind)}
        missing = sorted(mine - tmpl_points)
        extra = sorted(tmpl_points - mine)
        if missing:
            out.append(
                _issue(
                    "rows_will_be_added",
                    "info",
                    f"「{sn}」母版只留了 {len(tmpl_points)} 个点位行，本项目有 {len(mine)} 个测点",
                    f"导出会自动往下补 {len(missing)} 行（{ '、'.join(missing[:6]) }"
                    f"{'…' if len(missing) > 6 else ''}），并把下面的说明/签字块整体下移 —— "
                    "版面会变长，发出去前看一眼有没有挤到签字栏。",
                )
            )
        if extra:
            out.append(
                _issue(
                    "rows_will_be_cleared",
                    "info",
                    f"「{sn}」母版上多出的 {len(extra)} 个点位行会被清空",
                    f"这些点号（{ '、'.join(extra[:6]) }{'…' if len(extra) > 6 else ''}）不是本项目的测点，"
                    "导出会把它们的点号和读数一起擦掉，不会留下别人的测点。",
                )
            )
    return out
