# -*- coding: utf-8 -*-
"""Excel 解析器：Sheet 识别 → 表头识别 → 列映射 → 数据行提取。

基于用户真实模板（八号基坑）实测的结构约定：
- 双层表头：R5 为主表头（监测点号/初始值/上次测值/本次测值/上次累计/本次累计/本次变化/速率/控制值/预警等级/备注），
  R6 为控制值子表头（累计变化量 mm / 变化速率 mm/d）。
- 观测日期在 R3：「本次监测时间：2026-07-31」。
- 竖向位移 Sheet：每测点 1 行。
- 水平位移 Sheet：每测点 2 行 —— 第一行点号=JJC1（X 分量），第二行点号空（Y 分量，继承上行点号）。
- 读值列（初始值/上次测值/本次测值）单位 m，入库存 mm（×1000）；累计/变化/速率列已是 mm，直接存。
"""
from __future__ import annotations

import datetime as _dt
import os
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import openpyxl

# 表头行搜索范围：真实台账标题/说明块有时占据前 20 多行，放宽到 30 行。
# 放宽后用"表头必须能解析出 点号 + 测值/初始值"来兜底，避免把说明页误判成数据表。
HEADER_SCAN_MAX = 30

# ---------------------------------------------------------------------------
# 常量与关键词表
# ---------------------------------------------------------------------------

# Sheet 名关键词 → 观测类型（白名单：只解析数据 Sheet；名称认不出时再看表头）
SHEET_KEYWORDS = {
    "settlement": ("竖向位移", "沉降", "垂直位移", "垂直沉降", "高程"),
    "displacement": ("水平位移", "立柱结构", "水平"),
}

# 列头关键词 → 字段名（去空白/换行后模糊匹配；同一字段取最左命中列）
# 兼容不同台账的说法差异（本次/本期、测值/读数/高程/坐标 等同义写法）
COLUMN_KEYWORDS = [
    ("point_no", ("点号", "测点", "监测点", "测点编号", "点编号")),
    ("initial_m", ("初始值", "初值", "基准值", "初始读数")),
    ("prev_m", ("上次测值", "上次高程", "上次坐标", "上次读数", "上期测值", "前次测值", "上次测量值")),
    (
        "current_m",
        ("本次测值", "本次高程", "本次坐标", "本次读数", "本期测值", "本期读数", "本次测量值", "本期高程", "本次观测值"),
    ),
    ("prev_cum_mm", ("上次累计", "上期累计", "前期累计")),
    ("cum_mm", ("本次累计", "累计变化量", "本期累计", "累计变化", "累计值")),
    ("delta_mm", ("本次变化", "本期变化", "变化量", "本次差值")),
    ("rate_mm_d", ("变化速率", "速率", "日变化", "变化速度")),
    ("control_cum_mm", ("控制值", "允许值", "限值", "报警值")),
    ("warning_level", ("预警等级", "预警", "报警等级")),
    ("remark", ("备注", "说明")),
]

# 观测日期：① 带标签（本次监测时间/观测日期/测量日期…）；② 退一步找含"日期/时间"字样的
DATE_PATTERN = re.compile(
    r"(?:本次|本期|此次)?(?:监测|观测|测量)(?:时间|日期)[:：]?\s*(\d{4})[-./年](\d{1,2})[-./月](\d{1,2})"
)
DATE_LOOSE_PATTERN = re.compile(r"(?:日期|时间)[:：]?\s*(\d{4})[-./年](\d{1,2})[-./月](\d{1,2})")

# 读值列单位：表头标注 (mm) → 已是毫米，不换算；否则（标注 m 或无标注）→ ×1000
_UNIT_MM_PATTERN = re.compile(r"[（(]\s*mm\s*[)）]", re.IGNORECASE)


def _norm(text: str | None) -> str:
    """去掉空白、换行、全角空格，便于关键词匹配。"""
    if text is None:
        return ""
    return str(text).replace("\n", "").replace(" ", "").replace("\u3000", "").strip()


# ---------------------------------------------------------------------------
# 老版 .xls（OLE2 复合文档）兼容：openpyxl 只认 xlsx，先转成 xlsx 再解析
# ---------------------------------------------------------------------------

_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def is_legacy_xls(path: str | Path) -> bool:
    """文件头是否 OLE2（老版 .xls）。这类文件即使扩展名写成 .xlsx 也不是 xlsx。"""
    try:
        with open(path, "rb") as fh:
            return fh.read(8) == _OLE2_MAGIC
    except OSError:
        return False


def xls_to_xlsx(path: str | Path) -> Path:
    """把老版 .xls 转成临时 .xlsx（仅取单元格值，供解析用）。

    临时文件写到系统临时目录，**不写进 uploads/**：否则进程中断会在上传目录里留下
    `*.converted.xlsx`，既不干净、也可能被后续扫描误当成用户上传的文件。
    """
    import tempfile

    import xlrd
    from openpyxl import Workbook

    book = xlrd.open_workbook(str(path))
    wb = Workbook()
    wb.remove(wb.active)
    for sh in book.sheets():
        ws = wb.create_sheet(str(sh.name)[:31])
        for r in range(sh.nrows):
            for c in range(sh.ncols):
                v = sh.cell_value(r, c)
                if v is None or (isinstance(v, str) and v.strip() == ""):
                    continue
                ws.cell(row=r + 1, column=c + 1, value=v)
    fd, tmp = tempfile.mkstemp(prefix="xls_convert_", suffix=".xlsx")
    os.close(fd)
    out = Path(tmp)
    wb.save(out)
    wb.close()
    return out


# 测点编号合法性：字母/数字/横线/下划线/斜杠组合，且至少含一个数字、长度 ≤ 24
_POINT_NO_PATTERN = re.compile(r"^[A-Za-z0-9_\-/]+$")
# 纯数值（含负数/小数，如 -1、0、2、3.5、2026）→ 是数据/序号，不是测点号。
# 真实台账常在数据表下方附一块差值/坐标小表，若不排除会把它们当成新测点。
_NUMERIC_LIKE = re.compile(r"^[+\-]?\d+(\.\d+)?$")


def is_valid_point_no(text: str | None) -> bool:
    """判断是否为合法测点编号（排除说明行/签名行/表格下方的数值区）。"""
    if not text:
        return False
    if len(text) > 24 or "：" in text or ":" in text:
        return False
    if _NUMERIC_LIKE.match(text):
        return False  # 纯数值（-1 / 0 / 2 / 2026 …）不是点号
    if not _POINT_NO_PATTERN.match(text):
        return False
    return any(ch.isdigit() for ch in text)


# 中文测点编号（如「左1号点」「2号右边蹲观测点1」）：现场台账常见。
# 判据：含数字、长度≤24、无冒号（冒号=说明/签字行）、不含页脚词、且该行必须有读数。
# 最后一条最关键：说明行/页脚行不会有读数，数据行一定有（哪怕是初始值）。
_STOP_POINT_WORDS = (
    "共", "合计", "小计", "总计", "说明", "备注", "日期", "时间",
    "页码", "第", "签名", "监测", "报表", "统计", "数量", "个数",
    "结论", "建议", "附表", "详见",
)
_CJK_POINT_PATTERN = re.compile(r"^[\w\u4e00-\u9fff\-/]+$")


def _is_cjk_point_like(text: str | None) -> bool:
    """中文点号的候选条件（比 is_valid_point_no 宽，但不含纯数值/页脚词/说明行）。"""
    if not text or len(text) > 24:
        return False
    if "：" in text or ":" in text:
        return False
    if _NUMERIC_LIKE.match(text):
        return False
    if not any(ch.isdigit() for ch in text):
        return False
    if any(w in text for w in _STOP_POINT_WORDS):
        return False
    return bool(_CJK_POINT_PATTERN.match(text))


def _accept_point_no(text: str, has_reading: bool) -> bool:
    """该单元格能否作为测点号。

    - ASCII 编号：沿用既有规则（放行，含缺测行）；
    - 中文编号：额外要求该行有读数——说明行/页脚行没有读数，因此不会被误当成测点。
    """
    if is_valid_point_no(text):
        return True
    return has_reading and _is_cjk_point_like(text)




@dataclass
class ParsedRow:
    """解析出的单条观测记录行（单位已统一为 mm）。"""

    point_no: str
    series_type: str            # settlement / displacement
    component: str              # value / X / Y
    current_mm: float | None    # 本次测值 × 1000；None = 缺测
    initial_mm: float | None    # 初始值 × 1000（V0 候选：仅初始值期有值）
    remark: str | None = None
    warning_level: str | None = None
    control_cum_mm: float | None = None   # 控制值·累计变化量（可提取生成预警规则）
    control_rate_mm_d: float | None = None  # 控制值·变化速率

    def to_record_dict(self) -> dict:
        return {
            "point_no": self.point_no,
            "series_type": self.series_type,
            "component": self.component,
            "value_mm": self.current_mm,
            "initial_mm": self.initial_mm,
            "remark": self.remark,
            "warning_level": self.warning_level,
            "control_cum_mm": self.control_cum_mm,
            "control_rate_mm_d": self.control_rate_mm_d,
        }


# 文件形态。两个上传入口要用它互相防呆：系统里有两个地方能传 Excel ——
#   ① 项目内「Excel 导入」：逐期观测台账（哪一天、本次测值多少）；
#   ② 项目列表「批量建项目」：「点号 + 坐标/高程」点位清单（用来建项目/测点）。
# 传错入口不会报错，只会**悄悄算出错的东西**（台账的本次测值被当成高程写成初始值，
# 或者一份点位清单被当成某一天的观测数据存档）—— 所以解析时就把它标出来。
KIND_LEDGER = "ledger"          # 逐期观测台账
KIND_POINT_LIST = "point_list"  # 点位清单（「点号 + 坐标/高程」，无表头也认）


@dataclass
class ParsedSheet:
    """单个数据 Sheet 的解析结果。"""

    sheet_name: str
    series_type: str
    obs_date: date | None
    column_map: dict[str, int] = field(default_factory=dict)   # 字段名 → 列号(1-based)
    rows: list[ParsedRow] = field(default_factory=list)
    skipped_sheets: list[str] = field(default_factory=list)
    # 这张表是怎么被认出来的（见上面的 KIND_* 说明）
    kind: str = KIND_LEDGER
    # 台账里的 V0 期（整表只填了「初始值」列、没有本次测值）——
    # 它的初始值就是坐标/高程，因此可以当点位清单用（见 is_period_ledger）。
    is_baseline: bool = False
    # 表头区里读到的「工程名称」候选（跨项目导入靠它认这张表属于哪个基坑）。
    # 空列表 = 表里没写名称，只能退回按文件名认。
    name_texts: list[str] = field(default_factory=list)

    @property
    def total_rows(self) -> int:
        return len(self.rows)


def is_period_ledger(sheet: ParsedSheet) -> bool:
    """是「某一天的观测台账」而不是点位清单。

    判据 = 台账 + 不是 V0 期（表中填了本次测值）。这种表喂给「批量建项目」
    会被当成坐标清单：本次测值只有几毫米，却会被当高程写成测点初始值，
    而且**全程不报错**，等发现时累计变化量早已错乱。
    """
    return sheet.kind == KIND_LEDGER and not sheet.is_baseline and sheet.total_rows > 0


def _type_from_line(line: str) -> str | None:
    """单行文字 → 观测类型。注意「竖向位移」必须先于「位移」判定（否则会被判成水平位移）。"""
    if any(k in line for k in ("沉降", "高程", "竖向位移", "垂直位移", "垂直沉降")):
        return "settlement"
    if any(k in line for k in ("水平位移", "立柱结构", "坐标", "位移")):
        return "displacement"
    return None


def _detect_series_type(sheet_name: str, context_lines: list[str] | None = None,
                        header_text: str = "") -> str | None:
    """判断观测类型：① Sheet 名关键词 → ② 表头区各行文字（标题行优先）。

    ② 支持"Sheet 名是 Sheet1、表头里也没有沉降/位移字样"的台账——
    这类文件靠表格标题行（如"沉降监测报表"）就能判定类型，否则整表会被丢弃。
    """
    for stype, kws in SHEET_KEYWORDS.items():
        if any(k in sheet_name for k in kws):
            return stype
    lines = list(context_lines or [])
    if header_text:
        lines.append(header_text)
    for line in lines:
        t = _type_from_line(line)
        if t:
            return t
    return None


def _detect_column_map(header_cells: list[str]) -> dict[str, int]:
    """表头单元格序列 → 字段列号映射。第一个匹配优先。"""
    result: dict[str, int] = {}
    for idx, cell in enumerate(header_cells):
        t = _norm(cell)
        if not t:
            continue
        for field_name, kws in COLUMN_KEYWORDS:
            if field_name in result:
                continue
            if any(k in t for k in kws):
                result[field_name] = idx + 1  # 1-based
                break
    return result


def _refine_control_columns(ws, header_row: int, col_map: dict[str, int]) -> None:
    """双层表头：把「控制值」合并组下方的第二层列补进列映射。

    真实台账结构（竖向位移 / 水平位移表）：
        R5   I=「控制值」(合并 I:J)      J=空
        R6   I=「累计变化量（mm)」        J=「变化速率（mm/d）」
        R7   I=30（累计控制值）           J=10（速率控制值）
    表头行定位停在第 5 行，原先只取到 I 列的「控制值」，J 列的速率控制值被丢弃，
    导致"变化速率"这一档预警阈值无来源。

    口径与 reports/master_learn.py::_detect_data_columns 保持一致（top 行认控制值组，
    第二层行认累计/速率）。第二层行必须是表头文字，否则不动（防止把首行数据当表头）。
    """
    ctrl_col = col_map.get("control_cum_mm")
    if not ctrl_col:
        return
    sub_row = header_row + 1
    max_col = ws.max_column or 0
    if sub_row > (ws.max_row or 0) or max_col == 0:
        return
    # 第二层若已经是数据行（点号列是合法编号）→ 说明是单层表头，直接返回
    if is_valid_point_no(_norm(ws.cell(row=sub_row, column=col_map.get("point_no") or 1).value)):
        return
    sub_texts = [_norm(ws.cell(row=sub_row, column=c).value) for c in range(1, max_col + 1)]
    if not any(sub_texts):
        return
    cum_col = next((i + 1 for i, t in enumerate(sub_texts) if t and "累计" in t), None)
    rate_col = next((i + 1 for i, t in enumerate(sub_texts) if t and "速率" in t), None)
    # 第二层必须落在「控制值」组之内（≥ 控制值列），避免误抓左侧的变化量/速率列
    if cum_col and cum_col >= ctrl_col:
        col_map["control_cum_mm"] = cum_col
    if rate_col and rate_col >= ctrl_col:
        col_map["control_rate_mm_d"] = rate_col


def _header_looks_like_table(header_cells: list, col_map: dict[str, int]) -> bool:
    """判断被选中的"表头行"到底是真表格表头，还是一段说明文字。

    放宽表头扫描范围后，概况/说明页里"……监测点号共 8 个，本次测值……"这类整段文字
    也可能同时含「点号」和「测值」。用两条表格特征把它挡掉：
    ① 点号列表头是短标签（≤12 字），段落文字通常很长；
    ② 至少认到 2 个不同的数据列（段落文字只有 1 格，所有字段都会落在同一列）。
    阈值取 2 而不是 3：最小可用台账就是「点号 + 本次测值」两列，不能误伤。
    """
    if "point_no" not in col_map:
        return False
    if "current_m" not in col_map and "initial_m" not in col_map:
        return False
    pc = col_map["point_no"]
    label = _norm(header_cells[pc - 1]) if 0 < pc <= len(header_cells) else ""
    if len(label) > 12:
        return False
    return len(set(col_map.values())) >= 2




def _mk_date(m: re.Match) -> date | None:
    try:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None


def _extract_obs_date(meta_rows: list[str], bare_dates: list[date] | None = None) -> date | None:
    """从表头元信息提取观测日期。

    ① 优先带标签的写法（本次监测时间 / 观测日期 / 测量日期 …）；
    ② 退一步找含"日期/时间"字样的行里的第一个日期（兼容"日期：2026.8.1"）；
    ③ 最后用表头区里"真正的日期单元格"兜底（模板把日期做成日期格式、没写标签）。
       调用方已排除含「检定/校准/仪器」字样的行，避免抓到检定日期。
    """
    for line in meta_rows:
        m = DATE_PATTERN.search(str(line or ""))
        if m:
            d = _mk_date(m)
            if d:
                return d
    for line in meta_rows:
        s = str(line or "")
        # 宽松匹配容易误伤"检定日期：2025.10.20"（仪器检定日期 ≠ 观测日期）——排掉
        if "检定" in s or "校准" in s:
            continue
        m = DATE_LOOSE_PATTERN.search(s)
        if m:
            d = _mk_date(m)
            if d:
                return d
    if bare_dates:
        return bare_dates[0]
    return None


def _head_name_texts(head_rows: list[list]) -> list[str]:
    """表头区里的「工程名称」候选（跨项目导入靠它认这张表属于哪个基坑）。

    标签与取法和 `master_report._extract_master_project_name`、
    `export_reader._name_candidates` 是**同一套**（标签后按空白切分、最后一段是名字，
    前面「兴善公路」这类地名前缀也一并留着）—— 三处各写一套的话，早晚有一处漏改，
    认项目就会和导出的口径对不上。

    ⚠ 必须**逐格**看，不能用表头区整行拼接的文字：拼起来会把隔壁那格
    「本次监测时间：2026-07-31」粘到名字后面，认项目当场认歪。
    ⚠ 也不能先过 `_norm`：它会把空格全删掉，「兴善公路 八号基坑」变成一整坨，
    就再也切不出"最后一段"了。
    """
    from backend.reports.master_report import NAME_LABELS

    out: list[str] = []
    for cells in head_rows:
        texts = [("" if v is None else str(v)).strip() for v in cells]
        for i, t in enumerate(texts):
            if not t:
                continue
            lb = next((x for x in NAME_LABELS if x in t), None)
            if lb is None:
                continue
            rest = t[t.find(lb) + len(lb):].strip()
            if not rest:
                # 「监测工程名称：」自己一格、名字写在右边那格（表单常见排版）
                rest = next((x for x in texts[i + 1:] if x), "")
            parts = [p.strip("（）()【】[]") for p in re.split(r"[\s\u3000]+", rest) if p.strip()]
            for cand in [rest.strip("（）()【】[]")] + parts:
                if cand and cand not in out:
                    out.append(cand)
            break  # 一行只认一处：「监测工程名称：」本身也含「工程名称：」
    return out


def _to_float(v) -> float | None:
    """宽松转 float；空/'-'/'—'/None → None（缺测）。"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("", "-", "—", "/", "无"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _mm_factor(header_text) -> float:
    """读值列换算到 mm 的系数。

    列头标注 (mm) → 数值本身即毫米，系数 1；
    其余情况（标注 m、或无单位标注，沿用用户台账惯例）→ ×1000。
    """
    t = _norm(header_text).lower()
    if _UNIT_MM_PATTERN.search(t) or t.endswith("mm"):
        return 1.0
    return 1000.0


# 「名字 + X/Y/Z」简易基准表（初始值清单）的列名识别
_SIMPLE_NAME_KEYS = ("名字", "点号", "编号", "测点")


def _xyz_of(text: str | None) -> str | None:
    """列头 → X/Y/Z 分量（支持 "Y"、"X坐标"、"高程" 等写法）。"""
    if not text:
        return None
    raw = _norm(text)
    u = raw.upper()
    cleaned = (
        u.replace("坐标", "")
        .replace("值", "")
        .replace("M", "")
        .replace("（", "")
        .replace("）", "")
        .replace("(", "")
        .replace(")", "")
    )
    if cleaned in ("X", "Y", "Z"):
        return cleaned
    if raw in ("东", "东坐标"):
        return "X"
    if raw in ("北", "北坐标"):
        return "Y"
    if raw in ("高程", "标高"):
        return "Z"
    return None


def _try_parse_simple_points(ws) -> ParsedSheet | None:
    """识别「名字 + X/Y/Z」简易基准表（初始值清单），它不是标准台账模板。

    表头形如：名字 | Y | X | Z（列顺序不限、按文字匹配）
    行数据形如：JJC1 | 510199.345 | 3413030.175 | 4.411
    语义：Z = 沉降初始读数(m)；X/Y = 水平位移初始坐标(m)；均 ×1000 入 mm。
    表中没有观测日期 → obs_date=None，由前端补填日期后入库（作为该期读数/基准）。
    """
    for r in range(1, min(15, ws.max_row or 0) + 1):
        texts = [_norm(ws.cell(row=r, column=c).value) for c in range(1, (ws.max_column or 0) + 1)]
        if not any(texts):
            continue
        name_col = next(
            (i + 1 for i, t in enumerate(texts) if t and any(k in t for k in _SIMPLE_NAME_KEYS)), None
        )
        if name_col is None:
            continue
        xyz_cols: dict[str, int] = {}
        for i, t in enumerate(texts):
            k = _xyz_of(t)
            if k and k not in xyz_cols:
                xyz_cols[k] = i + 1
        if not xyz_cols:
            continue

        rows: list[ParsedRow] = []
        for rr in range(r + 1, (ws.max_row or 0) + 1):
            pno = _norm(ws.cell(row=rr, column=name_col).value)
            if not pno or not _accept_point_no(pno, has_reading=True):
                continue
            for k, col in xyz_cols.items():
                v = _to_float(ws.cell(row=rr, column=col).value)
                if v is None:
                    continue
                v_mm = v * _mm_factor(ws.cell(row=r, column=col).value)
                if k == "Z":
                    rows.append(
                        ParsedRow(
                            point_no=pno, series_type="settlement", component="value",
                            current_mm=v_mm, initial_mm=v_mm,
                        )
                    )
                else:
                    rows.append(
                        ParsedRow(
                            point_no=pno, series_type="displacement", component=k,
                            current_mm=v_mm, initial_mm=v_mm,
                        )
                    )
        if not rows:
            continue
        return ParsedSheet(
            sheet_name=ws.title,
            series_type="both",
            obs_date=None,
            column_map={"point_no": name_col, **{k.lower(): v for k, v in xyz_cols.items()}},
            rows=rows,
            kind=KIND_POINT_LIST,
        )
    return None


def _try_parse_headerless_points(ws) -> ParsedSheet | None:
    """无表头的「名字 + 坐标 + 高程」清单（现场仪器导出常见，整表没有表头行和日期）。

    形态：
        A: 6号右左边基坑1   B: 3412997.533   C: 510197.577   D: 4.619
    列语义按数值量级推断：坐标列中位数 ≥1e5（本项目 X≈3413xxx、Y≈510xxx），
    其中绝对值大的当 X、小的当 Y；其余小数列（几米）当 Z（高程 → 沉降）。
    单位按米 ×1000 入库。

    ⚠ X/Y 是靠量级推出来的，存在互换风险 → 解析结果会在导入预览里逐列展示
      （点号=A列 · X=B列 · Y=C列 · Z=D列），由用户核对后再入库。
    """
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    if max_row < 2 or max_col < 3 or max_col > 8:
        return None

    names: list[str] = []
    grid: list[list[float | None]] = []
    for r in range(1, max_row + 1):
        name = _norm(ws.cell(row=r, column=1).value)
        if not name:
            # 空行 → **跳过，不要放弃整表**（2026-10-06 用户实测报障）。
            # 现场导出的清单常在数据下方留若干空行、或删过内容但保留行记录，
            # openpyxl 的 max_row 会一直报到那些行；早先写成「一行不合法就 return None」，
            # 结果是**这种文件整份解析不出来**，界面只说"没有读到数据行"，用户无从下手。
            continue
        if not _accept_point_no(name, has_reading=True):
            break  # 出现"不像点号"的一行 → 认为表到此结束（已收集的行照常使用）
        vals = [_to_float(ws.cell(row=r, column=c).value) for c in range(2, max_col + 1)]
        if sum(1 for v in vals if v is not None) < 2:
            break  # 这一行没有坐标 → 表结束（同上，不放弃前面的数据）
        names.append(name)
        grid.append(vals)
    if len(names) < 2:
        return None

    def median(ci: int) -> float | None:
        xs = sorted(v[ci] for v in grid if v[ci] is not None)
        return xs[len(xs) // 2] if xs else None

    # vals 从第 2 列开始取，故真实列号 = ci + 2
    cols = [(ci + 2, median(ci)) for ci in range(len(grid[0]))]
    cols = [(c, m) for c, m in cols if m is not None]
    # 坐标列按数值绝对值从大到小 → 大的当 X（本项目 X≈3413xxx），小的当 Y（≈510xxx）
    coords = sorted(
        ((c, m) for c, m in cols if abs(m) >= 1e5), key=lambda t: -abs(t[1])
    )
    smalls = [c for c, m in cols if abs(m) < 1e4]
    if not coords:
        return None
    x_col = coords[0][0]
    y_col = coords[1][0] if len(coords) >= 2 else None
    z_col = smalls[0] if smalls else None
    if y_col is None and z_col is None:
        return None

    rows: list[ParsedRow] = []
    for name, vals in zip(names, grid):
        def val(c: int | None) -> float | None:
            return vals[c - 2] if c else None

        z = val(z_col)
        if z is not None:
            rows.append(ParsedRow(point_no=name, series_type="settlement", component="value",
                                  current_mm=z * 1000, initial_mm=z * 1000))
        for comp, c in (("X", x_col), ("Y", y_col)):
            v = val(c)
            if v is not None:
                rows.append(ParsedRow(point_no=name, series_type="displacement", component=comp,
                                      current_mm=v * 1000, initial_mm=v * 1000))
    if not rows:
        return None

    column_map: dict[str, int] = {"point_no": 1, "x": x_col}
    if y_col:
        column_map["y"] = y_col
    if z_col:
        column_map["z"] = z_col
    return ParsedSheet(
        sheet_name=ws.title, series_type="both", obs_date=None,
        column_map=column_map, rows=rows, kind=KIND_POINT_LIST,
    )


def _parse_sheet(ws, header_row: int) -> tuple[dict[str, int], list[list], list]:
    """识别列映射并返回（列映射, 数据行, 表头单元格原文）。"""
    header_cells = [ws.cell(row=header_row, column=c).value for c in range(1, ws.max_column + 1)]
    col_map = _detect_column_map(header_cells)
    raw_rows: list[list] = []
    for r in range(header_row + 1, ws.max_row + 1):
        cells = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        # 整行空 → 跳过
        if all(c is None or _norm(c) == "" for c in cells):
            continue
        raw_rows.append(cells)
    return col_map, raw_rows, header_cells


def parse_workbook(path: str | Path) -> list[ParsedSheet]:
    """解析工作簿，返回各数据 Sheet 的结构化结果（单位统一 mm）。

    兼容三类文件：
    1. 老版 .xls（OLE2）——自动转 xlsx 后解析；
    2. 标准台账模板（双层表头 + 竖向位移/水平位移）；
    3. 「名字 + X/Y/Z」简易基准表（初始值清单，无日期）。
    """
    src = Path(path)
    converted: Path | None = None
    if is_legacy_xls(src):
        converted = xls_to_xlsx(src)
        src = converted
    wb = openpyxl.load_workbook(src, data_only=True)
    parsed: list[ParsedSheet] = []
    try:
        for ws in wb.worksheets:
            sheet_name = ws.title

            # ---- 1) 定位表头行（前 HEADER_SCAN_MAX 行内找同时含「点号/测点」+「测值/读数/累计」的行）----
            header_row: int | None = None
            header_text = ""
            meta_lines: list[str] = []
            head_rows: list[list] = []  # 表头区各行的**原始单元格**（抽工程名称用，见 _head_name_texts）
            bare_dates: list[date] = []  # 表头区里的"真日期单元格"（无标签时的兜底）
            for r in range(1, min(HEADER_SCAN_MAX, ws.max_row) + 1):
                cells = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
                joined = "".join(_norm(c) for c in cells)
                meta_lines.append(joined)
                head_rows.append(cells)
                # 日期单元格兜底：按单元格判断，跳过"检定日期/校准日期"这类邻格标签，
                # （不能整行跳过——观测日期与检定日期可能在同一行）
                for ci, v in enumerate(cells, start=1):
                    if isinstance(v, _dt.datetime):
                        d_cell = v.date()
                    elif isinstance(v, _dt.date):
                        d_cell = v
                    else:
                        continue
                    neigh = "".join(
                        _norm(cells[j - 1]) for j in (ci - 1, ci + 1) if 1 <= j <= len(cells)
                    )
                    if any(k in neigh for k in ("检定", "校准", "仪器")):
                        continue
                    bare_dates.append(d_cell)
                if ("点号" in joined or "测点" in joined) and ("测值" in joined or "读数" in joined or "累计" in joined):
                    header_row = r
                    header_text = joined
                    break

            # 类型判定：Sheet 名优先，其次表头区文字（含标题行，如"沉降监测报表"）
            series_type = (
                _detect_series_type(sheet_name, meta_lines[: max(1, header_row - 1)], header_text)
                if header_row
                else None
            )
            col_map: dict[str, int] = {}
            raw_rows: list[list] = []
            header_cells: list = []
            if header_row is not None:
                col_map, raw_rows, header_cells = _parse_sheet(ws, header_row)
                _refine_control_columns(ws, header_row, col_map)
            # 表头有效性：必须能认到「点号」+「本次测值/初始值」，且整体像一张真表格，
            # 否则视为说明页 → 落到下面的简易基准表兜底。
            has_valid_header = _header_looks_like_table(header_cells, col_map)
            if series_type is None or not has_valid_header:
                # 不是标准台账：尝试「名字 + X/Y/Z」简易基准表（有表头 / 无表头两种形态）
                simple = _try_parse_simple_points(ws) or _try_parse_headerless_points(ws)
                if simple is not None:
                    parsed.append(simple)
                continue  # 其余非数据 Sheet（封面/签字/概况等）跳过

            obs_date = _extract_obs_date(meta_lines[: max(1, header_row - 1)], bare_dates)

            ps = ParsedSheet(
                sheet_name=sheet_name,
                series_type=series_type,
                obs_date=obs_date,
                name_texts=_head_name_texts(head_rows[: max(0, header_row - 1)]),
            )
            ps.column_map = col_map
            # 读值列单位：表头标注 (mm) → 已是毫米不换算；否则按米 ×1000（沿用现有台账惯例）
            hdr_text = {
                f: (header_cells[c - 1] if 0 <= c - 1 < len(header_cells) else None) for f, c in col_map.items()
            }
            factor_current = _mm_factor(hdr_text.get("current_m"))
            factor_initial = _mm_factor(hdr_text.get("initial_m"))

            # ---- 2) V0 期检测：整 Sheet 无本次测值但有初始值 → 初始值即本期读数 ----
            # 用户模板语义：初始值文件（如 2026.7.30）只填初始值列，代表 V0 基准日；
            # 后续期文件每次都会重复初始值列，但因为有本次测值，不触发本逻辑。
            current_col_idx = col_map.get("current_m")
            initial_col_idx = col_map.get("initial_m")
            has_any_current = any(
                _to_float(cells[current_col_idx - 1]) is not None
                for cells in raw_rows
                if current_col_idx is not None and current_col_idx <= len(cells)
            )
            has_any_initial = any(
                _to_float(cells[initial_col_idx - 1]) is not None
                for cells in raw_rows
                if initial_col_idx is not None and initial_col_idx <= len(cells)
            )
            is_v0_sheet = (not has_any_current) and has_any_initial
            ps.is_baseline = is_v0_sheet

            # ---- 3) 逐行解析（含位移 X/Y 两行的继承规则）----
            pending_point_no: str | None = None  # 位移表 Y 行继承
            for cells in raw_rows:
                def col(field: str):
                    idx = col_map.get(field)
                    if idx is None or idx > len(cells):
                        return None
                    return cells[idx - 1]

                point_no = _norm(col("point_no"))
                current_m = _to_float(col("current_m"))
                # 行首有点号 → 新测点（位移 X 行）；点号空但测值有 → 继承上行（位移 Y 行）
                if point_no:
                    # 中文编号需该行确有读数，否则视为说明/页脚行（见 _accept_point_no）
                    has_reading = any(
                        _to_float(col(f)) is not None for f in ("current_m", "initial_m", "prev_m")
                    )
                    if not _accept_point_no(point_no, has_reading):
                        continue
                    pending_point_no = point_no
                elif current_m is not None and pending_point_no:
                    point_no = pending_point_no
                elif is_v0_sheet and pending_point_no and _to_float(col("initial_m")) is not None:
                    # V0 期：位移 Y 行同样只有初始值（点号空，继承上行）
                    point_no = pending_point_no
                else:
                    continue  # 无效行

                component = "value" if series_type == "settlement" else ("X" if _norm(col("point_no")) else "Y")

                initial_m = _to_float(col("initial_m"))
                prev_cum = _to_float(col("prev_cum_mm"))
                cum = _to_float(col("cum_mm"))
                delta = _to_float(col("delta_mm"))
                rate = _to_float(col("rate_mm_d"))
                warning_level = _norm(col("warning_level")) or None
                remark = _norm(col("remark")) or None

                # 控制值：「控制值」为合并表头，第二层拆成 累计控制值 / 速率控制值 两列
                # （_refine_control_columns 已把第二层列补进 col_map）。
                control_cum = _to_float(col("control_cum_mm"))
                control_rate = _to_float(col("control_rate_mm_d"))

                # V0 期：本次测值缺省时用初始值作为该期读数（V0 基准）
                eff_current_m = current_m if current_m is not None else (initial_m if is_v0_sheet else None)
                # 换算系数：本次测值有值按其列单位；V0 期取的是初始值，按初始值列单位
                eff_factor = factor_current if current_m is not None else factor_initial

                ps.rows.append(
                    ParsedRow(
                        point_no=point_no,
                        series_type=series_type,
                        component=component,
                        current_mm=eff_current_m * eff_factor if eff_current_m is not None else None,
                        initial_mm=initial_m * factor_initial if initial_m is not None else None,
                        remark=remark,
                        warning_level=warning_level,
                        control_cum_mm=control_cum,
                        control_rate_mm_d=control_rate,
                    )
                )
            parsed.append(ps)
    finally:
        wb.close()
        if converted is not None:
            converted.unlink(missing_ok=True)
    return parsed
