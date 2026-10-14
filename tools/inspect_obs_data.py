# -*- coding: utf-8 -*-
"""深入分析观测数据 Sheet 的完整结构与三期差异（只读）。

用法：
    python tools/inspect_obs_data.py                  # 目录下前 3 个 xlsx
    python tools/inspect_obs_data.py a.xlsx b.xlsx    # 显式指定文件
    python tools/inspect_obs_data.py --dir "D:\\某目录"

默认台账目录可用环境变量 MONITOR_LEDGER_DIR 覆盖，避免写死本机绝对路径。
"""
import argparse
import os
import sys
from pathlib import Path

import openpyxl

DEFAULT_LEDGER_DIR = os.environ.get(
    "MONITOR_LEDGER_DIR",
    r"D:\文件\测量台账\7、测量资料\基坑观测\八号（已填）",
)

SHEETS = ["竖向位移", "r.立柱结构水平位移"]


ap = argparse.ArgumentParser(description="只读深挖观测 Excel 的 Sheet 结构")
ap.add_argument("files", nargs="*", help="要分析的 xlsx（留空则取目录下前 N 个）")
ap.add_argument("--dir", dest="dir_", default=None, help="台账目录（默认 DEFAULT_LEDGER_DIR）")
ap.add_argument("--limit", type=int, default=3, help="目录模式下最多分析几个（默认 3）")
_args = ap.parse_args()

if _args.files:
    FILES = [Path(f) for f in _args.files]
else:
    _d = Path(_args.dir_ or DEFAULT_LEDGER_DIR)
    if not _d.is_dir():
        print("[错误] 台账目录不存在：%s" % _d)
        print("       请用 --dir 指定目录，或设置环境变量 MONITOR_LEDGER_DIR。")
        sys.exit(2)
    FILES = sorted(p for p in _d.glob("*.xlsx") if not p.name.startswith("~$"))[: _args.limit]

if not FILES:
    print("[错误] 没有找到任何 .xlsx 文件。")
    sys.exit(2)

for path in FILES:
    print("=" * 90)
    print("文件:", Path(path).name)
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=False)
    for ws in wb.worksheets:
        if ws.title not in SHEETS:
            continue
        print("-" * 90)
        print(f"[{ws.title}] 行={ws.max_row} 列={ws.max_column}")
        # 打印所有行（数据行重点）
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row), start=1):
            vals = []
            for c in row:
                v = c.value
                if v is None:
                    continue
                if isinstance(v, float) and v == int(v):
                    v = int(v)
                vals.append(f"{c.coordinate}={str(v)[:22]}")
            if vals:
                print(f"  R{row_idx:>2}: " + " | ".join(vals))
    wb.close()
    print()
