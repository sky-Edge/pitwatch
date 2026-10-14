# -*- coding: utf-8 -*-
"""只读分析观测 Excel 模板结构（不修改原文件）。

用法：
    python tools/inspect_obs_excel.py                  # 目录下前 3 个 xlsx
    python tools/inspect_obs_excel.py a.xlsx b.xlsx    # 显式指定文件
    python tools/inspect_obs_excel.py --dir "D:\\某目录" --limit 5

默认台账目录可用环境变量 MONITOR_LEDGER_DIR 覆盖，避免写死本机绝对路径。
"""
import argparse
import os
import sys
from pathlib import Path

import openpyxl

DEFAULT_LEDGER_DIR = os.environ.get(
    "MONITOR_LEDGER_DIR",
    r"D:\文件\测量台账\7、测量资料\基坑观测\八号（已填）",
)


def _resolve_files(argv):
    ap = argparse.ArgumentParser(description="只读分析基坑观测 Excel 的表结构")
    ap.add_argument("files", nargs="*", help="要分析的 xlsx（留空则取目录下前 N 个）")
    ap.add_argument("--dir", dest="dir_", default=None, help="台账目录（默认 DEFAULT_LEDGER_DIR）")
    ap.add_argument("--limit", type=int, default=3, help="目录模式下最多分析几个（默认 3）")
    args = ap.parse_args(argv)

    if args.files:
        return [Path(f) for f in args.files], None

    d = Path(args.dir_ or DEFAULT_LEDGER_DIR)
    if not d.is_dir():
        print("[错误] 台账目录不存在：%s" % d)
        print("       请用 --dir 指定目录，或设置环境变量 MONITOR_LEDGER_DIR。")
        sys.exit(2)
    found = sorted(p for p in d.glob("*.xlsx") if not p.name.startswith("~$"))
    return found[: args.limit], d


FILES, _USED_DIR = _resolve_files(sys.argv[1:])
if not FILES:
    print("[错误] 没有找到任何 .xlsx 文件。")
    sys.exit(2)

for path in FILES:
    print("=" * 80)
    print("文件:", Path(path).name)
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=False)
    print("Sheet 列表:", wb.sheetnames)
    for ws in wb.worksheets:
        print("-" * 60)
        print(f"Sheet: {ws.title} | 维度: {ws.dimensions} | max_row={ws.max_row} max_col={ws.max_column}")
        # 合并单元格
        merges = list(ws.merged_cells.ranges)[:20]
        print(f"合并单元格(前20): {merges}")
        # 打印前 8 行内容
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(8, ws.max_row)), start=1):
            vals = []
            for c in row:
                v = c.value
                if v is None:
                    vals.append("")
                elif isinstance(v, float) and v == int(v):
                    vals.append(str(int(v)))
                else:
                    vals.append(str(v)[:18])
            print(f"  R{row_idx}: " + " | ".join(vals))
    wb.close()
