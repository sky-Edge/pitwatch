# -*- coding: utf-8 -*-
"""批量建项目（一份「点号 + X/Y/Z」清单 → 拆多个基坑项目）测试。

样本取自用户真实清单：一行一个点，列序是「点号 | Y | X | Z」，
X/Y 由解析器按数值量级判定（X≈3413xxx、Y≈510xxx），所以列序反了也能认。
"""
import io
from pathlib import Path

import openpyxl
import pytest

from backend.services.batch_project_service import _suggest_groups, parse_point_sheet
from tests.test_import_parser import _make_wb

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SRC = Path(__file__).resolve().parent.parent / "frontend" / "src"
SERVICE = Path(__file__).resolve().parent.parent / "backend" / "services" / "batch_project_service.py"

# 真表前几行（点号 | Y | X | Z），单位 m
ROWS = [
    ("15号", 510023.8, 3413192.929, 3.119),
    ("15号1", 510026.156, 3413189.664, 3.079),
    ("15号2", 510029.077, 3413188.018, 3.095),
    ("14号", 510045.008, 3413174.821, 3.22),
    ("14号1", 510049.548, 3413170.911, 3.176),
    ("14号2", 510052.219, 3413171.143, 3.187),
]


def _wb_bytes(rows=None, name="兴 善_2026-09-28-15-13-40.xlsx") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "清单"
    for i, r in enumerate(rows if rows is not None else ROWS, start=1):
        for c, v in enumerate(r, start=1):
            ws.cell(row=i, column=c, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def _write(tmp_path: Path, rows=None) -> Path:
    p = tmp_path / "list.xlsx"
    p.write_bytes(_wb_bytes(rows))
    return p


def _preview(client):
    return client.post(
        "/api/v1/batch-import/preview",
        files={"file": ("清单.xlsx", _wb_bytes(), XLSX_MIME)},
    ).json()


def _commit(client, groups):
    return client.post("/api/v1/batch-import/commit", json={"groups": groups})


def _group(code, date, point_nos, **kw):
    pts = []
    for no in point_nos:
        row = next((r for r in ROWS if r[0] == no), None)
        assert row, no
        pts.append({"point_no": row[0], "x": row[2] * 1000, "y": row[1] * 1000, "z": row[3] * 1000})
    return {"group_name": code, "project_code": code, "project_name": code,
            "obs_date": date, "points": pts, **kw}


# ---------- 解析 ----------


def test_parse_point_list(tmp_path):
    """解析出全部点，X/Y 按量级判对（列序是 Y 在前也能认）。"""
    r = parse_point_sheet(_write(tmp_path))
    assert len(r["points"]) == 6
    assert r["column_map"] == {"point_no": 1, "x": 3, "y": 2, "z": 4}
    first = r["points"][0]
    assert first["point_no"] == "15号"
    assert first["x"] == pytest.approx(3413192929.0)   # 3413192.929 m × 1000
    assert first["y"] == pytest.approx(510023800.0)
    assert first["z"] == pytest.approx(3119.0)
    assert r["has_date"] is False


def test_parse_suggests_groups_by_prefix(tmp_path):
    r = parse_point_sheet(_write(tmp_path))
    got = {g["name"]: g["point_count"] for g in r["suggested_groups"]}
    assert got == {"15号": 3, "14号": 3}


def test_suggest_groups_silent_when_meaningless():
    """点号没规律时不要给一堆"一组一个点"的假分组。"""
    pts = [{"point_no": n} for n in ("甲", "乙", "丙")]
    assert _suggest_groups(pts) == []
    # 只有一个组也不给（等于没分）
    assert _suggest_groups([{"point_no": n} for n in ("JJ1", "JJ2")]) == []


def test_point_sheet_tolerates_trailing_blank_rows(tmp_path):
    """清单下方留了空行（或删过内容）也要能解析。

    踩坑（2026-10-06 用户实测报障）：openpyxl 的 `max_row` 会把"有行记录但没内容"的行
    一起报进来（在 Excel 里拖过格式、删过内容都会这样）。解析器早先写成
    **"遇到一行不合法就整表放弃（return None）"** —— 于是这类文件整份解析不出来，
    界面只说"没有读到数据行"。用户换了 3 个文件、还改存成 .xlsm 反复试了 7 次都不成功。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, r in enumerate(ROWS, start=1):
        for c, v in enumerate(r, start=1):
            ws.cell(row=i, column=c, value=v)
    # 数据下方"留空"，但让 openpyxl 仍认为有更多行（模拟格式残留 / 删过内容）
    for rr in range(len(ROWS) + 1, len(ROWS) + 12):
        ws.cell(row=rr, column=8).value = None
    assert ws.max_row > len(ROWS), "没有造出空行，这条测试等于没写"
    path = tmp_path / "带空行清单.xlsx"
    wb.save(path)
    wb.close()

    data = parse_point_sheet(path)
    assert [p["point_no"] for p in data["points"]] == [r[0] for r in ROWS]


def test_numeric_only_point_no_gets_actionable_message(tmp_path):
    """点号全是纯数字时，报错要能指导下一步，而不是甩一句通用提示。

    纯数字点号是**有意拒绝**的（真实台账下方常附一块「差值／坐标小表」，
    那里的数字会被误当成新测点）。但用户不知道这条规则 ——
    2026-10-06 实测：他只看到"没有读到数据行"，不知道该改什么。
    """
    rows = [(i, 510110.5872, 3413111.7272, 2.7986) for i in range(1, 11)]
    path = tmp_path / "纯数字点号.xlsx"
    path.write_bytes(_wb_bytes(rows))

    with pytest.raises(ValueError) as err:
        parse_point_sheet(path)
    msg = str(err.value)
    assert "纯数字" in msg, f"没有说清失败原因：{msg}"
    assert "钢便桥1" in msg or "JJC1" in msg, f"没有给出可以照做的改法：{msg}"


def test_preview_api(client):
    r = _preview(client)
    assert len(r["points"]) == 6
    assert len(r["suggested_groups"]) == 2


def test_preview_rejects_non_excel(client):
    resp = client.post(
        "/api/v1/batch-import/preview",
        files={"file": ("x.txt", b"hello", "text/plain")},
    )
    assert resp.status_code == 400


# ---------- 入口防呆：观测台账 ≠ 点位清单 ----------


def _ledger_bytes(baseline_only: bool = False) -> bytes:
    """造一份「逐期观测台账」（与真实模板同构）。

    baseline_only=True 时清掉「本次测值」列 —— 那就是**初始值清单**（V0 期），
    这张表里的初始值就是坐标/高程，可以当点位清单用。
    """
    wb = _make_wb()
    if baseline_only:
        for name in ("竖向位移", "r.立柱结构水平位移"):
            ws = wb[name]
            for r in range(7, 12):
                ws.cell(row=r, column=4).value = None
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def test_period_ledger_rejected_with_pointer(tmp_path):
    """把「某一天的观测台账」传到批量建项目：必须挡下，并说清该去哪。

    不挡的后果是**静默错**：台账里的「本次测值」只有几毫米（-1.0、-3.0），
    却被当成高程写成测点初始值，一路不报错；等发现时累计变化量早就错乱了。
    提示还必须能照着做 —— 只说"格式不对"用户只会反复换文件（已经发生过一次）。
    """
    path = tmp_path / "八号基坑（第一期）2026.7.31.xlsx"
    path.write_bytes(_ledger_bytes())

    with pytest.raises(ValueError) as err:
        parse_point_sheet(path)
    msg = str(err.value)
    assert "逐期观测台账" in msg, f"没说清这份文件是什么：{msg}"
    assert "Excel 导入" in msg, f"没告诉用户该去哪个入口：{msg}"
    assert "初始值" in msg, f"没给出补救办法（整理初始值列）：{msg}"


def test_period_ledger_rejected_via_api(client):
    resp = client.post(
        "/api/v1/batch-import/preview",
        files={"file": ("八号基坑（第一期）2026.7.31.xlsx", _ledger_bytes(), XLSX_MIME)},
    )
    assert resp.status_code == 422
    assert "Excel 导入" in resp.json()["detail"]


def test_baseline_ledger_still_usable(tmp_path):
    """只填了「初始值」列的 V0 期台账 = 初始值清单 —— 这条老路不能因为防呆被堵死。"""
    path = tmp_path / "初始值清单.xlsx"
    path.write_bytes(_ledger_bytes(baseline_only=True))

    r = parse_point_sheet(path)
    by_no = {p["point_no"]: p for p in r["points"]}
    assert set(by_no) == {"JJC1", "JJC2"}
    assert by_no["JJC1"]["z"] == pytest.approx(4907.0)            # 4.907 m × 1000
    assert by_no["JJC1"]["x"] == pytest.approx(3413055120.0)      # 3413055.12 m × 1000
    assert by_no["JJC1"]["y"] == pytest.approx(510174166.0)


def test_import_preview_marks_sheet_kind(client):
    """预览里带 `kind`：前端据此判断"是不是传错了入口"，两个页面互相指路。"""
    pid = client.post("/api/v1/projects", json={"code": "PRJ001", "name": "八号基坑"}).json()["id"]

    ledger = client.post(
        f"/api/v1/projects/{pid}/import/upload",
        files={"file": ("台账.xlsx", _ledger_bytes(), XLSX_MIME)},
    ).json()
    assert ledger["sheets"] and all(s["kind"] == "ledger" for s in ledger["sheets"])

    listed = client.post(
        f"/api/v1/projects/{pid}/import/upload",
        files={"file": ("清单.xlsx", _wb_bytes(), XLSX_MIME)},
    ).json()
    assert listed["sheets"] and all(s["kind"] == "point_list" for s in listed["sheets"])


def test_frontend_points_both_entries_at_each_other() -> None:
    """两个上传入口各写一句"传错了去哪"（防呆的可见部分，防回归）。"""
    imp = (SRC / "pages" / "ImportPage.tsx").read_text(encoding="utf-8")
    batch = (SRC / "pages" / "BatchProjectsPage.tsx").read_text(encoding="utf-8")
    assert "point_list" in imp, "导入页没有识别点位清单（提醒不会出现）"
    assert "批量建项目" in imp, "导入页没有告诉用户点位清单该去哪"
    assert "Excel 导入" in batch, "批量建项目页没有告诉用户观测台账该去哪"


# ---------- 落库 ----------


def test_commit_creates_projects_points_and_records(client):
    resp = _commit(
        client,
        [
            _group("15号", "2026-09-22", ["15号", "15号1", "15号2"]),
            _group("14号", "2026-09-22", ["14号", "14号1", "14号2"]),
        ],
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["totals"] == {
        "projects_created": 2,
        "projects_reused": 0,
        "points_created": 6,
        "points_skipped": 0,
        "records_created": 18,  # 每个点 3 条：沉降 value + 位移 X/Y
    }
    assert all(g["created_project"] for g in body["groups"])

    projects = {p["code"]: p for p in client.get("/api/v1/projects").json()}
    assert "15号" in projects and "14号" in projects


def test_commit_reuses_existing_project(client):
    """编号已存在 → 复用，不重复建、不报错（用户口径：已有的直接用）。"""
    client.post("/api/v1/projects", json={"code": "15号", "name": "十五号基坑"})
    body = _commit(client, [_group("15号", "2026-09-22", ["15号", "15号1"])]).json()
    assert body["totals"]["projects_created"] == 0
    assert body["totals"]["projects_reused"] == 1
    assert body["groups"][0]["project_name"] == "十五号基坑"  # 复用已有项目的名称
    assert len(client.get("/api/v1/projects").json()) == 1


def test_commit_chinese_point_no(client):
    """现场点号就是「15号1」这种，必须能建（point_service 那套校验只放行字母数字，不能用）。"""
    body = _commit(client, [_group("15号", "2026-09-22", ["15号1"])]).json()
    assert body["groups"][0]["points_created"] == 1
    pid = body["groups"][0]["project_id"]
    nos = [p["point_no"] for p in client.get(f"/api/v1/projects/{pid}/points").json()]
    assert nos == ["15号1"]


def test_commit_skips_existing_point_without_touching_it(client):
    """项目里已有同名测点 → 跳过并报出，**绝不覆盖**已有数据。"""
    pid = client.post("/api/v1/projects", json={"code": "15号", "name": "十五号"}).json()["id"]
    pt = client.post(f"/api/v1/projects/{pid}/points", json={"point_no": "15号1", "point_type": "both"}).json()
    client.post(
        f"/api/v1/points/{pt['id']}/records",
        json={"series_type": "settlement", "component": "value", "obs_date": "2026-09-01", "value_mm": 3079.0},
    )

    body = _commit(client, [_group("15号", "2026-09-22", ["15号1"])]).json()
    g = body["groups"][0]
    assert g["points_created"] == 0
    assert g["records_created"] == 0
    assert g["skipped"] and "已有同名测点" in g["skipped"][0]["reason"]
    # 原有记录分毫未动
    recs = client.get(f"/api/v1/points/{pt['id']}/records").json()
    assert len(recs) == 1 and recs[0]["obs_date"] == "2026-09-01"


def test_commit_group_without_date_is_rejected(client):
    """没给初始值日期 → 直接 422。

    日期必须由用户明确选定（现场"初始值是哪天测的"只有他知道），
    所以不给默认值、也不让它悄悄落在"今天"。服务层另有一层兜底（整组跳过并报错），
    这里断言的是 API 层的拦截。
    """
    resp = _commit(client, [_group("15号", None, ["15号1"])])
    assert resp.status_code == 422


def test_commit_skips_point_without_readings(client):
    body = _commit(
        client,
        [{"group_name": "15号", "project_code": "15号", "project_name": "15号",
          "obs_date": "2026-09-22", "points": [{"point_no": "15号1"}]}],
    ).json()
    g = body["groups"][0]
    assert g["points_created"] == 0
    assert "没有任何读数" in g["skipped"][0]["reason"]


def test_commit_is_idempotent(client):
    """重复提交同一批：第二次不该造出重复数据。"""
    payload = [_group("15号", "2026-09-22", ["15号1", "15号2"])]
    first = _commit(client, payload).json()
    second = _commit(client, payload).json()
    assert first["totals"]["points_created"] == 2
    assert second["totals"]["points_created"] == 0
    assert second["totals"]["records_created"] == 0
    assert len(second["groups"][0]["skipped"]) == 2


# ---------- 两条关键约束（写错会让数据悄悄歪掉） ----------


def test_commit_does_not_set_initial_value_mm(client):
    """绝不写 `initial_value_mm`。

    它是**测点级单值**，而 calc 会把同一个值传给沉降、位移 X、位移 Y 三条序列当 V0。
    高程 3.119m → 3119 写进去，位移 X 的 V0 就是 3119，
    而 X 读数是 3413192929 —— 累计变化量会变成天文数字，且看不出是哪一步错的。
    """
    body = _commit(client, [_group("15号", "2026-09-22", ["15号1"])]).json()
    pid = body["groups"][0]["project_id"]
    points = client.get(f"/api/v1/projects/{pid}/points").json()
    assert all(p["initial_value_mm"] is None for p in points)


def test_displacement_v0_is_first_reading(client):
    """位移 X 的首期累计必须为 0（V0=首期读数自动成立，未被污染）。"""
    body = _commit(client, [_group("15号", "2026-09-22", ["15号1"])]).json()
    pid = body["groups"][0]["project_id"]
    pt = client.get(f"/api/v1/projects/{pid}/points").json()[0]

    for series, comp, expected in (
        ("settlement", "value", 3079.0),
        ("displacement", "X", 3413189664.0),
        ("displacement", "Y", 510026156.0),
    ):
        series_rows = client.get(
            f"/api/v1/points/{pt['id']}/computed", params={"series_type": series, "component": comp}
        ).json()
        assert len(series_rows) == 1, f"{series}/{comp} 应有 1 期"
        row = series_rows[0]
        assert row["obs_date"] == "2026-09-22"
        assert row["cumulative_mm"] == pytest.approx(0.0), f"{series}/{comp} 首期累计应为 0"
        assert row["value_mm"] == pytest.approx(expected)


def test_service_never_writes_initial_value_mm() -> None:
    """源码守卫：这个坑太隐蔽，用一个断言焊死。"""
    src = SERVICE.read_text(encoding="utf-8")
    # 允许在文档/注释里提到它（解释为什么不能写），但代码里不能出现赋值
    code_lines = [ln for ln in src.splitlines() if not ln.strip().startswith("#")]
    offending = [ln for ln in code_lines if "initial_value_mm=" in ln]
    assert not offending, f"批量建点写了 initial_value_mm，会让位移 V0 错乱：{offending}"


def test_frontend_has_batch_entry() -> None:
    """入口与页面都在（防回归：功能被误删）。"""
    app = (SRC / "App.tsx").read_text(encoding="utf-8")
    projects = (SRC / "pages" / "ProjectsPage.tsx").read_text(encoding="utf-8")
    page = (SRC / "pages" / "BatchProjectsPage.tsx").read_text(encoding="utf-8")
    assert "batch-projects" in app and "BatchProjectsPage" in app
    assert "onBatchProjects" in projects
    assert "batchImportApi.preview" in page and "batchImportApi.commit" in page


def test_frontend_generates_point_no_by_rule() -> None:
    """点号按**规则**生成（前缀 + 起始序号，按清单顺序往下编），不是照搬表里的名字。

    现场清单里的点名（如「15号1」）往往不能直接当测点编号用 ——
    用户要的是 JJC1、JJC2… 这种规整编号。
    """
    page = (SRC / "pages" / "BatchProjectsPage.tsx").read_text(encoding="utf-8")
    assert "pointNoOf" in page, "没有按规则生成点号的逻辑"
    assert "cfg.prefix" in page and "cfg.start" in page, "缺少前缀/起始序号的配置"
    assert "point_no: pointNoOf(r)" in page, "提交时没用规则生成的点号（回传了表里的原名）"


def test_commit_accepts_rule_generated_point_nos(client):
    """点号来自规则（JJC1…JJC6）时，落库的就是这些编号 —— 与表里的名字无关。"""
    pts = [
        {"point_no": f"JJC{i}", "x": r[2] * 1000, "y": r[1] * 1000, "z": r[3] * 1000}
        for i, r in enumerate(ROWS, start=1)
    ]
    body = _commit(
        client,
        [{"group_name": "十五号", "project_code": "15号", "project_name": "十五号基坑",
          "obs_date": "2026-09-22", "points": pts}],
    ).json()
    assert body["groups"][0]["points_created"] == len(ROWS)
    pid = body["groups"][0]["project_id"]
    nos = sorted(
        (p["point_no"] for p in client.get(f"/api/v1/projects/{pid}/points").json()),
        key=lambda s: int(s[3:]),
    )
    assert nos == [f"JJC{i}" for i in range(1, len(ROWS) + 1)]
