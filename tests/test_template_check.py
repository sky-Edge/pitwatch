# -*- coding: utf-8 -*-
"""母版体检（导出**前**查母版本身）。

用户 2026-10-16 提的需求：「做个模版检查吧」。起因是同一天报的障 ——
导出件里 I/J「控制值」上半是规则里的 35/5、下半还是母版上印的 30/10：
母版预留的空白点位行**没有点号**，导出按点号认行、扫到第一个空点号就停，那几行够不着。
导出端已修（`fill_data_sheets` 把控制值写到说明行为止），但**光修一次不够**：
下次换一份形状不同的母版（说明块下面还有一张表、表名被复制成两份…）还会再犯，
所以要在点导出之前就把这些问题说给人听。

体检的判据和导出端**用的是同一批函数**，否则会出现"体检说没事、导出却出事"。
"""
from pathlib import Path

import openpyxl

from tests.conftest import put_control_rules
from tests.test_control_values import _download_master, _master_with_control, _upload_master
from tests.test_master_report import _make_master, _seed_project

BACKEND = Path(__file__).resolve().parent.parent / "backend"


# ---------- 工装 ----------


def _new_project(client, name="体检测试", code="PRJ-CHK"):
    return client.post("/api/v1/projects", json={"code": code, "name": name}).json()["id"]


def _check(client, pid: int, tid: int) -> dict:
    r = client.get(f"/api/v1/projects/{pid}/export-templates/{tid}/check")
    assert r.status_code == 200, r.text
    return r.json()


def _codes(data: dict) -> list[str]:
    return [i["code"] for i in data["issues"]]


def _one(data: dict, code: str) -> dict:
    got = [i for i in data["issues"] if i["code"] == code]
    assert got, f"没查出 {code}：{data['issues']}"
    return got[0]


def _master_below_note_row(tmp_path: Path) -> Path:
    """竖向位移表：7~8 点位行 → 第 13 行是说明 → **14~15 行还印着控制值**。

    14~15 行在数据区界线（说明行）之外，导出够不着 → 这正是"同一列两个数"的另一形态。
    """
    wb = openpyxl.load_workbook(_make_master())
    ws = wb["竖向位移"]
    for r in (7, 8):
        ws.cell(row=r, column=9, value=30)
        ws.cell(row=r, column=10, value=10)
    ws.cell(row=13, column=1, value="变化量正负号说明：“-”为上升“+”为下沉")
    for r in (14, 15):
        ws.cell(row=r, column=9, value=30)
        ws.cell(row=r, column=10, value=10)
    out = tmp_path / "master_below_note.xlsx"
    wb.save(out)
    wb.close()
    return out


# ---------- ① 拦得住的：导出会失败 ----------


def test_missing_control_rules_reported_as_error(client, tmp_path):
    """项目一类观测没配控制值 → 导出必被 422 拦；体检要提前用 error 说清楚。"""
    pid = _new_project(client)
    tid = _upload_master(client, pid, _master_with_control(tmp_path))

    data = _check(client, pid, tid)
    assert data["errors"] == 2, data["issues"]  # 沉降 + 位移各一条
    assert _codes(data).count("no_control_rule") == 2
    for issue in data["issues"]:
        if issue["code"] == "no_control_rule":
            assert "预警规则" in issue["detail"], "得告诉用户去哪儿配"


def test_half_filled_rule_says_which_cell_stays_blank(client, tmp_path):
    """只填了累计、没填速率 → 不算 error，但要说明速率那一格留空、不兜底。"""
    pid = _new_project(client)
    put_control_rules(client, pid, [
        {"indicator": "settlement", "level": "control", "threshold_cumulative_mm": 30,
         "threshold_rate_mm_per_day": 10, "is_active": True},
        {"indicator": "displacement", "level": "control", "threshold_cumulative_mm": 20,
         "threshold_rate_mm_per_day": None, "is_active": True},
    ])
    tid = _upload_master(client, pid, _master_with_control(tmp_path))

    data = _check(client, pid, tid)
    issue = _one(data, "half_control_rule")
    assert issue["level"] == "info" and "变化速率" in issue["title"], issue
    assert data["errors"] == 0


def test_missing_master_file_is_error(client, tmp_path):
    """原件不在盘上（历史上模板文件被测试删过）→ 直接 error，别说得含糊。"""
    pid = _new_project(client)
    put_control_rules(client, pid)
    tid = _upload_master(client, pid, _master_with_control(tmp_path))
    r = client.get(f"/api/v1/projects/{pid}/export-templates")
    fp = next(t["file_path"] for t in r.json() if t["id"] == tid)
    Path(fp).unlink()

    data = _check(client, pid, tid)
    assert _one(data, "file_missing")["level"] == "error"


# ---------- ② 够不着的旧值：这次报障的那一类 ----------


def test_stale_control_cells_below_the_note_row_are_flagged(client, tmp_path):
    """控制值印在**数据区界线之外**（说明行下面）→ warn，并给出具体格子。

    这就是"一半对得上一半对不上"的形态之一：导出只重写数据区，下面的原样发出去。
    """
    pid = _new_project(client)
    put_control_rules(client, pid)
    tid = _upload_master(client, pid, _master_below_note_row(tmp_path))

    data = _check(client, pid, tid)
    issue = _one(data, "stale_control_out_of_range")
    assert issue["level"] == "warn"
    assert any(c.startswith("竖向位移!I14") for c in issue["cells"]), issue["cells"]
    assert "导出只重写" in issue["detail"]


def test_reserved_blank_rows_inside_data_area_are_not_flagged(client, tmp_path):
    """母版点位行下面的**预留空白行**（没点号、印着旧控制值）现在会被导出重写 → 不该再报。

    这条和上一条只差说明行的位置：界线之内导出够得着（2026-10-16 修的就是它），
    界线之外才是真残留。体检必须和导出端同口径，不能一律报警。
    """
    wb = openpyxl.load_workbook(_make_master())
    ws = wb["竖向位移"]
    for r in range(7, 13):  # 7~8 点位行 + 9~12 空白预留行
        ws.cell(row=r, column=9, value=30)
        ws.cell(row=r, column=10, value=10)
    ws.cell(row=13, column=1, value="变化量正负号说明：“-”为上升“+”为下沉")
    p = tmp_path / "master_reserved.xlsx"
    wb.save(p)
    wb.close()

    pid = _new_project(client)
    put_control_rules(client, pid)
    tid = _upload_master(client, pid, p)
    assert "stale_control_out_of_range" not in _codes(_check(client, pid, tid))


def test_leftover_sheet_with_control_header_is_flagged(client, tmp_path):
    """母版里留着一张没被学到的旧页面（里面有「控制值」表头和旧数）→ warn。

    实测：整张复制出来的「竖向位移(2)」会被学习认成第二张数据表，导出会写它，没问题；
    有问题的是那种**结构不全、学习跳过、但上面印着限值**的遗留页 —— 导出根本不碰它。
    """
    wb = openpyxl.load_workbook(_make_master())
    ws = wb["竖向位移"]
    for r in (7, 8):
        ws.cell(row=r, column=9, value=30)
        ws.cell(row=r, column=10, value=10)
    old = wb.create_sheet("上期报表存档")
    old["I5"] = "控制值"
    old["I14"] = 30
    old["J14"] = 10
    p = tmp_path / "master_leftover.xlsx"
    wb.save(p)
    wb.close()

    pid = _new_project(client)
    put_control_rules(client, pid)
    tid = _upload_master(client, pid, p)
    issue = _one(_check(client, pid, tid), "sheet_not_learned")
    assert issue["level"] == "warn" and "上期报表存档" in issue["title"], issue


def test_second_control_header_in_same_sheet_is_flagged(client, tmp_path):
    """同一张表里第二块「控制值」表头（表下面又接了半张表）→ 那块不会被刷新。"""
    wb = openpyxl.load_workbook(_make_master())
    ws = wb["竖向位移"]
    ws.cell(row=18, column=9, value="控制值")
    ws.cell(row=19, column=9, value=30)
    p = tmp_path / "master_two_blocks.xlsx"
    wb.save(p)
    wb.close()

    pid = _new_project(client)
    put_control_rules(client, pid)
    tid = _upload_master(client, pid, p)
    assert _one(_check(client, pid, tid), "second_block_in_sheet")["level"] == "warn"


# ---------- ③ 工程名 / 点位行数：把导出会做什么说清楚 ----------


def test_project_name_replacement_is_announced(client, tmp_path):
    """母版自称「8号基坑」（标签与名字在同一格才认得出）、项目叫「体检测试」→ 说明导出会替换。"""
    wb = openpyxl.load_workbook(_make_master())
    wb["封面"]["E11"] = "工程名称：8号基坑"  # 名字单独放 F11 时标签格里没有名字，认不出来
    p = tmp_path / "master_named.xlsx"
    wb.save(p)
    wb.close()

    pid = _new_project(client, name="体检测试")
    put_control_rules(client, pid)
    tid = _upload_master(client, pid, p)

    issue = _one(_check(client, pid, tid), "name_replaced")
    assert "8号基坑" in issue["title"] and "体检测试" in issue["title"], issue
    assert "处" in issue["title"], "要把替换处数说清楚"


def test_unrecognized_name_slot_tells_the_user_to_check(client, tmp_path):
    """认不出母版自称的基坑名 → 工程名称会原样保留，这条得提示人工核对（不是默默放过）。"""
    pid = _new_project(client, name="体检测试")
    put_control_rules(client, pid)
    tid = _upload_master(client, pid, _master_with_control(tmp_path))

    issue = _one(_check(client, pid, tid), "name_not_detected")
    assert issue["level"] == "info" and "工程名称" in issue["detail"], issue


def test_point_row_gap_is_announced(client, tmp_path):
    """母版只有 2 个点位行、项目 8 个测点 → 导出会补行并把说明块下移，提前说。"""
    pid = _seed_project(client, pno="PRJ-CHK-GAP")
    tid = _upload_master(client, pid, _master_with_control(tmp_path))

    data = _check(client, pid, tid)
    added = _one(data, "rows_will_be_added")
    assert added["level"] == "info" and "签字" in added["detail"], added


def test_extra_master_rows_are_announced(client, tmp_path):
    """反过来（母版点位行比本项目多）：多出来的点号会被连标签一起清掉，也要说出来。"""
    pid = _new_project(client, name="只有一个点", code="PRJ-CHK-ONE")
    put_control_rules(client, pid)
    client.post(f"/api/v1/projects/{pid}/points", json={"point_no": "JJC1", "point_type": "both"})
    tid = _upload_master(client, pid, _master_with_control(tmp_path))

    cleared = _one(_check(client, pid, tid), "rows_will_be_cleared")
    assert cleared["level"] == "info" and "JJC2" in cleared["detail"], cleared


def test_bound_check_refuses_other_projects_master(client, tmp_path):
    """没绑定本项目的母版：导出会被 400 拦，体检同样拦（别拿别人的版式查出问题算在本项目头上）。"""
    other = _new_project(client, name="别人家", code="PRJ-OTHER")
    tid = _upload_master(client, other, _master_with_control(tmp_path))
    mine = _new_project(client, name="我这边", code="PRJ-MINE")

    r = client.get(f"/api/v1/projects/{mine}/export-templates/{tid}/check")
    assert r.status_code == 400, r.text
    assert "没有绑定本项目" in r.json()["detail"]


# ---------- ④ 与导出端同源 ----------


def test_check_uses_the_export_side_functions() -> None:
    """体检不许另写一套「哪些行会被重写」的判断。

    自己数一遍行 → 体检和导出迟早各说各话（用户看到的就成了「查不出问题的检查」）。
    """
    src = (BACKEND / "services" / "template_check_service.py").read_text(encoding="utf-8")
    assert "_master_point_blocks" in src and "tail_start_row" in src, "数据区界线必须复用导出端判据"
    assert "report_control_values" in src, "控制值得从唯一出处取"
    assert "_sync_project_name" in src, "工程名替换结果要照导出端的实际行为报"
    for forbidden in ("MONITOR_DB_URL", ".save(", "wb.save"):
        assert forbidden not in src, f"体检是只读的，不该出现 {forbidden}"


def test_clean_master_and_full_rules_have_no_error(client, tmp_path):
    """规则配齐 + 母版正常 → errors==0（体检不能一上来就一片红，那样没人看）。"""
    pid = _seed_project(client, pno="PRJ-CHK-OK")
    tid = _upload_master(client, pid, _master_with_control(tmp_path))

    data = _check(client, pid, tid)
    assert data["errors"] == 0, data["issues"]
    assert all(i["level"] in ("warn", "info") for i in data["issues"]), data["issues"]
    # 顺带确认真的能导出（体检过了 ≠ 导出一定过，两边都跑一遍才算闭环）
    assert _download_master(client, pid, tid)["竖向位移"].cell(row=7, column=9).value == 30
