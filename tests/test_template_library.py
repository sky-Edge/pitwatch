# -*- coding: utf-8 -*-
"""测试：模板库 —— 模板独立于项目，由模板声明"作用于哪些项目"。

用户 2026-10-08：「每个项目单独维护一个模板太繁琐了，能否把模板从项目中独立提取出来，
改由模板来定义其作用于哪个项目？」

改之前：`export_templates.project_id` 把模板钉死在一个项目上（18 份模板里 15 份
是同一个版式族反复上传）。改之后：
- 可见性 = 内置 ∪ 通用(`is_global`) ∪ 绑定到该项目(`template_projects`)；
- 上传/改版时只维护**一份**母版，改一次全体跟随（同名上传自动继承绑定）；
- 导出处补上归属校验：拿没绑本项目的母版导出，直接拦下并指路。
"""
import io
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import text

from backend.core.db import engine, migrate_template_library

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ─────────────────────────── 脚手架 ───────────────────────────


def _tpl_bytes(v0: float = 3.1) -> bytes:
    """一份最小可用的母版（竖向位移表，双层表头，数据从 R7 起）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "竖向位移"
    ws["A5"] = "监测\n点号"
    ws["B5"] = "初始值（m)"
    ws["C5"] = "上次测值（m）"
    ws["D5"] = "本次测值（m）"
    ws["F5"] = "本次累计（mm)"
    ws["A7"] = "JJC1"
    ws["B7"] = v0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _mk_project(client, code: str, name: str) -> int:
    return client.post("/api/v1/projects", json={"code": code, "name": name}).json()["id"]


def _mk_project_with_data(client, code: str, name: str) -> int:
    pid = _mk_project(client, code, name)
    pt = client.post(
        f"/api/v1/projects/{pid}/points", json={"point_no": "JJC1", "point_type": "settlement"}
    ).json()
    for d, v in (("2026-07-30", 4907.0), ("2026-07-31", 4906.0)):
        client.post(
            f"/api/v1/points/{pt['id']}/records",
            json={"series_type": "settlement", "component": "value", "obs_date": d, "value_mm": v},
        )
    return pid


def _upload(client, pid: int, name: str, *, scope: list[int] | None = None, is_global: bool = False,
            via_library: bool = False):
    """上传母版。默认走**项目内入口**（兼容薄壳）；via_library=True 走新的模板库入口。"""
    data = {"name": name}
    if via_library:
        data["project_ids"] = ",".join(str(x) for x in (scope or []))
        if is_global:
            data["is_global"] = "true"
        url = "/api/v1/export-templates/upload"
    else:
        url = f"/api/v1/projects/{pid}/export-templates/upload"
    return client.post(url, data=data, files={"file": ("t.xlsx", _tpl_bytes(), XLSX_MIME)})


def _visible_in(client, pid: int) -> list[str]:
    lst = client.get(f"/api/v1/projects/{pid}/export-templates").json()
    return sorted(t["name"] for t in lst if not t["is_builtin"])


def _tpl_of(client, pid: int, name: str) -> dict:
    lst = client.get(f"/api/v1/projects/{pid}/export-templates").json()
    return next(t for t in lst if t["name"] == name)


def _export(client, pid: int, tid: int):
    return client.get(
        f"/api/v1/projects/{pid}/reports/excel",
        params={"template_id": tid, "obs_date": "2026-07-31"},
    )


# ─────────────────────── 1. 适用范围（多对多 + 通用） ───────────────────────


def test_upload_binds_only_current_project(client):
    """项目内上传（兼容入口）→ 只绑本项目，别的项目看不到。"""
    a, b = _mk_project(client, "A", "甲基坑"), _mk_project(client, "B", "乙基坑")
    assert _upload(client, a, "甲母版").status_code == 200

    assert _visible_in(client, a) == ["甲母版"]
    assert _visible_in(client, b) == [], "别人的母版不该出现在我的下拉里"


def test_new_upload_lands_in_library_dir(client):
    """原件落在模板库目录（模板不再归属某个项目），且模板库清单能一次看到它。"""
    a = _mk_project(client, "A", "甲基坑")
    _upload(client, a, "甲母版")

    lib = client.get("/api/v1/export-templates").json()
    mine = [t for t in lib if t["name"] == "甲母版"]
    assert len(mine) == 1
    assert "_library" in mine[0]["file_path"], mine[0]["file_path"]
    assert mine[0]["project_ids"] == [a]
    assert mine[0]["project_names"] == ["甲基坑"]


def test_scope_can_cover_several_projects(client):
    """一份母版绑给多个项目 → 每个绑定的项目都能看到、都能导出。

    这就是"不用每个项目各传一遍"的核心。
    """
    a = _mk_project_with_data(client, "A", "甲基坑")
    b = _mk_project_with_data(client, "B", "乙基坑")
    tid = _upload(client, a, "基坑通用母版").json()["id"]

    r = client.put(f"/api/v1/export-templates/{tid}/projects", json={"project_ids": [a, b]})
    assert r.status_code == 200, r.text
    assert r.json()["project_ids"] == [a, b]

    assert _visible_in(client, a) == ["基坑通用母版"]
    assert _visible_in(client, b) == ["基坑通用母版"]
    assert _export(client, b, tid).status_code == 200, "绑定过的项目必须能导出"


def test_global_template_is_visible_everywhere(client):
    """通用母版：所有项目都能用（不必逐个勾）。"""
    a = _mk_project_with_data(client, "A", "甲基坑")
    b = _mk_project_with_data(client, "B", "乙基坑")
    tid = _upload(client, a, "全项目通用").json()["id"]

    r = client.put(f"/api/v1/export-templates/{tid}/projects", json={"project_ids": [], "is_global": True})
    assert r.status_code == 200, r.text
    assert r.json()["is_global"] is True and r.json()["project_ids"] == []

    assert _visible_in(client, a) == ["全项目通用"]
    assert _visible_in(client, b) == ["全项目通用"]
    assert _export(client, b, tid).status_code == 200


def test_unbind_makes_template_invisible(client):
    """解绑后该项目立刻看不到它（可见性真的按绑定走，不靠历史字段兜着）。"""
    a = _mk_project_with_data(client, "A", "甲基坑")
    b = _mk_project_with_data(client, "B", "乙基坑")
    tid = _upload(client, a, "甲母版").json()["id"]
    client.put(f"/api/v1/export-templates/{tid}/projects", json={"project_ids": [a, b]})
    assert _visible_in(client, b) == ["甲母版"]

    client.put(f"/api/v1/export-templates/{tid}/projects", json={"project_ids": [a]})
    assert _visible_in(client, b) == [], "解绑了还出现在下拉里"


def test_scope_cannot_be_emptied(client):
    """既不通用的、又不绑任何项目 → 拦下（现在没有母版库页面，它会在哪儿都选不到）。"""
    a = _mk_project(client, "A", "甲基坑")
    tid = _upload(client, a, "甲母版").json()["id"]

    r = client.put(f"/api/v1/export-templates/{tid}/projects", json={"project_ids": []})
    assert r.status_code == 400
    detail = r.json()["detail"]
    assert "通用" in detail, detail
    assert _visible_in(client, a) == ["甲母版"], "被拦下就不该改动原有绑定"


def test_scope_rejects_unknown_project(client):
    a = _mk_project(client, "A", "甲基坑")
    tid = _upload(client, a, "甲母版").json()["id"]
    r = client.put(f"/api/v1/export-templates/{tid}/projects", json={"project_ids": [9999]})
    assert r.status_code == 404


def test_builtin_scope_cannot_be_changed(client):
    a = _mk_project(client, "A", "甲基坑")
    builtin = next(t for t in client.get(f"/api/v1/projects/{a}/export-templates").json() if t["is_builtin"])
    r = client.put(f"/api/v1/export-templates/{builtin['id']}/projects", json={"project_ids": [a]})
    assert r.status_code == 400


def test_library_upload_can_bind_without_any_project_context(client):
    """新入口：不经过任何项目也能建模板（只绑指定项目 / 直接设为通用）。"""
    a, b = _mk_project(client, "A", "甲基坑"), _mk_project(client, "B", "乙基坑")
    r = _upload(client, 0, "库里建的母版", scope=[a, b], via_library=True)
    assert r.status_code == 200, r.text
    assert r.json()["project_ids"] == [a, b]
    assert _visible_in(client, a) == ["库里建的母版"]
    assert _visible_in(client, b) == ["库里建的母版"]


def test_library_upload_without_scope_is_rejected(client):
    """什么都不绑又不是通用 → 400 且说清怎么办（工具要傻瓜式，不能建个用不了的模板）。"""
    r = _upload(client, 0, "孤儿母版", scope=[], via_library=True)
    assert r.status_code == 400
    assert "通用" in r.json()["detail"]


# ─────────────────────── 2. 导出处必须校验归属 ───────────────────────


def test_export_rejects_template_not_bound_to_project(client):
    """拿别的项目的母版导出 → 拦下并指路（此前这条校验是缺的，后端谁都能套）。"""
    a = _mk_project_with_data(client, "A", "甲基坑")
    b = _mk_project_with_data(client, "B", "乙基坑")
    tid = _upload(client, a, "甲母版").json()["id"]

    r = _export(client, b, tid)
    assert r.status_code == 400, r.text
    detail = r.json()["detail"]
    assert "没有绑定本项目" in detail and "设置适用项目" in detail, detail


def test_export_batch_also_rejects_unbound_template(client):
    """批量导出同样在前置检查里拦下（不然会产出 N 行"每期都失败"的误导清单）。"""
    a = _mk_project_with_data(client, "A", "甲基坑")
    b = _mk_project_with_data(client, "B", "乙基坑")
    tid = _upload(client, a, "甲母版").json()["id"]

    r = client.get(
        f"/api/v1/projects/{b}/reports/excel/batch",
        params={"template_id": tid, "obs_dates": "2026-07-30,2026-07-31"},
    )
    assert r.status_code == 400
    assert "没有绑定本项目" in r.json()["detail"]


def test_list_and_export_agree_on_usability(client):
    """下拉里能看到的，导出就必须能用 —— 两条判据出自同一个 `usable_clause`。

    （各写一份判断，迟早出现"下拉里有、导出却被拦"的自相矛盾）
    """
    a = _mk_project_with_data(client, "A", "甲基坑")
    b = _mk_project_with_data(client, "B", "乙基坑")
    tid = _upload(client, a, "甲母版").json()["id"]

    for pid in (a, b):
        listed = client.get(f"/api/v1/projects/{pid}/export-templates").json()
        if any(t["id"] == tid for t in listed):
            assert _export(client, pid, tid).status_code == 200, f"项目{pid}：下拉有但导出被拦"
        else:
            assert _export(client, pid, tid).status_code == 400, f"项目{pid}：下拉没有却能导出"


# ─────────────────────── 3. 同名上传：改一次全体跟随 ───────────────────────


def test_same_name_upload_inherits_bindings(client):
    """重传同名母版 = 改版：原先绑定的项目自动继承（不必一个个重新勾）。"""
    a, b, c = (
        _mk_project(client, "A", "甲基坑"),
        _mk_project(client, "B", "乙基坑"),
        _mk_project(client, "C", "丙基坑"),
    )
    tid = _upload(client, a, "基坑母版").json()["id"]
    client.put(f"/api/v1/export-templates/{tid}/projects", json={"project_ids": [a, b]})

    r = _upload(client, c, "基坑母版")  # 从丙项目重传同一个名字
    assert r.status_code == 200, r.text
    assert "自动继承" in r.json()["message"], r.json()["message"]
    assert sorted(r.json()["project_ids"]) == [a, b, c]

    for pid in (a, b, c):
        assert "基坑母版" in _visible_in(client, pid)
    # 旧记录不该留活跃残留：同名只剩一条
    lib = client.get("/api/v1/export-templates").json()
    assert len([t for t in lib if t["name"] == "基坑母版"]) == 1


def test_global_template_stays_global_after_reupload(client):
    a = _mk_project(client, "A", "甲基坑")
    tid = _upload(client, a, "通用母版").json()["id"]
    client.put(f"/api/v1/export-templates/{tid}/projects", json={"project_ids": [], "is_global": True})

    b = _mk_project(client, "B", "乙基坑")
    r = _upload(client, a, "通用母版")
    assert r.json()["is_global"] is True, "重传之后不该掉成「只绑本项目」"
    assert "通用" in r.json()["message"]
    assert _visible_in(client, b) == ["通用母版"], "通用母版重传后仍应全项目可见"


def test_rename_conflict_is_global_now(client):
    """名字唯一性从"项目内唯一"升级为**全局唯一** —— 同名母版在下拉框里分不出来。"""
    a, b = _mk_project(client, "A", "甲基坑"), _mk_project(client, "B", "乙基坑")
    _upload(client, a, "母版A")
    tid_b = _upload(client, b, "母版B").json()["id"]

    r = client.put(f"/api/v1/export-templates/{tid_b}", json={"name": "母版A"})
    assert r.status_code == 400
    assert "换个名字" in r.json()["detail"]


def test_any_project_can_edit_shared_template(client):
    """模板已独立于项目：绑定项目之外也能改它（不再有"不属于该项目"这种拦截）。"""
    a, b = _mk_project(client, "A", "甲基坑"), _mk_project(client, "B", "乙基坑")
    tid = _upload(client, a, "共享母版").json()["id"]
    assert _visible_in(client, b) == [], "没绑定的项目本来就不该看到它"

    assert client.put(f"/api/v1/export-templates/{tid}", json={"name": "改名了"}).status_code == 200
    assert _visible_in(client, a) == ["改名了"]


# ─────────────────────── 4. 迁移与兼容 ───────────────────────


def test_migration_backfills_links_and_is_idempotent(client):
    """直接写库造一条"老形态"模板（只有 project_id、没有绑定）→ 迁移应回填关联，且可重复执行。"""
    a = _mk_project(client, "A", "甲基坑")
    with engine.begin() as conn:
        conn.execute(
            text(
                "INSERT INTO export_templates (project_id, name, template_type, is_builtin, layout_json) "
                "VALUES (:pid, '老记录母版', 'excel', 0, NULL)"
            ),
            {"pid": a},
        )
    first = migrate_template_library()
    second = migrate_template_library()
    assert first["backfilled"] >= 1
    assert second["backfilled"] == 0, "迁移不幂等（重复执行又插了一遍）"

    with engine.connect() as conn:
        n = conn.execute(
            text("SELECT COUNT(*) FROM template_projects WHERE project_id = :pid"), {"pid": a}
        ).scalar()
    assert n == 1, f"关联表里有 {n} 条，应为 1"


def test_legacy_row_without_link_still_visible(client):
    """兜底：一条绑定都没有、也不通用的"孤儿"老行，仍按老字段 project_id 归属（代码回退友好）。

    ⚠ 这条兜底**不会**让解绑失效 —— 服务层禁止把模板配成"0 绑定且不通用"，
    所以手工解绑后不存在这种状态（见 test_unbind_makes_template_invisible）。
    """
    a, b = _mk_project(client, "A", "甲基坑"), _mk_project(client, "B", "乙基坑")
    with engine.begin() as conn:
        conn.execute(
            text(
                "INSERT INTO export_templates (project_id, name, template_type, is_builtin, layout_json) "
                "VALUES (:pid, '未迁移母版', 'excel', 0, NULL)"
            ),
            {"pid": a},
        )
    assert _visible_in(client, a) == ["未迁移母版"]
    assert _visible_in(client, b) == []


# ─────────────────────── 5. 硬删项目：不误删共用母版 ───────────────────────


def test_hard_delete_project_keeps_shared_template(client):
    """母版被两个项目共用 → 删掉其中一个项目只是解绑，母版与原件都保留。"""
    a = _mk_project_with_data(client, "A", "甲基坑")
    b = _mk_project_with_data(client, "B", "乙基坑")
    tid = _upload(client, a, "共用母版").json()["id"]
    client.put(f"/api/v1/export-templates/{tid}/projects", json={"project_ids": [a, b]})
    path = _tpl_of(client, a, "共用母版")["file_path"]
    assert Path(path).exists()

    r = client.delete(f"/api/v1/projects/{a}/hard", params={"force": "A"})
    assert r.status_code == 200, r.text
    assert r.json()["templates_kept"] == ["共用母版"], "应报告保留的共用母版（界面要告诉用户）"

    assert Path(path).exists(), "共用母版的原件被删了 —— 乙项目下次导出就会失败"
    assert _visible_in(client, b) == ["共用母版"]
    assert _export(client, b, tid).status_code == 200


def test_hard_delete_project_removes_exclusive_template(client):
    """只被本项目用的母版 → 跟着删（不留孤儿记录）。"""
    a = _mk_project_with_data(client, "A", "甲基坑")
    tid = _upload(client, a, "独占母版").json()["id"]

    r = client.delete(f"/api/v1/projects/{a}/hard", params={"force": "A"})
    assert r.status_code == 200
    assert r.json()["templates_kept"] == []
    lib = client.get("/api/v1/export-templates").json()
    assert all(t["id"] != tid for t in lib), "独占母版应随项目一起删掉"


# ─────────────────────── 6. 源码守卫（前后端不脱节） ───────────────────────

SRC = Path(__file__).resolve().parent.parent / "frontend" / "src"


def _code_only(src: str) -> str:
    """剥掉注释行 —— 解释性注释里出现的旧文案不该把守卫自己绊倒。"""
    return "\n".join(
        line for line in src.splitlines() if not line.strip().startswith(("//", "/*", "*", "{/*"))
    )


def test_frontend_exposes_scope_editing() -> None:
    """界面上必须有"设置适用项目"的入口，并且上传时能选适用范围。"""
    page = _code_only((SRC / "pages" / "ReportsPage.tsx").read_text(encoding="utf-8"))
    assert "ScopePicker" in page, "没有适用范围选择器"
    assert "设置适用项目" in page, "没有事后修改适用范围的入口"
    assert "所有项目通用" in page, "没有通用开关"
    assert "setTemplateScope" in page and "project_names" in page, "没把适用范围显示出来"


def test_client_has_scope_api() -> None:
    client_src = _code_only((SRC / "api" / "client.ts").read_text(encoding="utf-8"))
    assert "setTemplateScope" in client_src
    assert "templates/${templateId}/projects`" in client_src, "适用范围接口的 URL 不对"
    # 适用范围选择器要用到项目清单（client 里是 projectApi.list）
    page = _code_only((SRC / "pages" / "ReportsPage.tsx").read_text(encoding="utf-8"))
    assert "projectApi.list" in page, "页面没拉项目清单，适用项目勾不了"


def test_backend_and_frontend_declare_same_feature_flag() -> None:
    """前后端的能力标记必须同时登记 —— 漏一边会每次启动都弹"后端是旧版本"。"""
    from backend.main import BACKEND_FEATURES

    app = (SRC / "App.tsx").read_text(encoding="utf-8")
    assert "template_library" in BACKEND_FEATURES
    assert "template_library" in app
