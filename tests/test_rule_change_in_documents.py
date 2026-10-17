# -*- coding: utf-8 -*-
"""改预警规则 → **生成的文档**必须一一对应（用户 2026-10-14）。

用户原话：「预警规则更改条件 生成的文档中 也要一一对应 不只是对系统中更改」

这条不能靠读代码确认 —— 报表有三条生成路径（标准版式 / 母版套打 / 旧模板版式）
外加 Word 日报，每条都有"哪些格子来自规则、哪些来自母版"的各自的口径。
所以本文件用的是**实测**：同一批数据，规则改前改后各导一次，把整份文档逐格 diff，
凡是和阈值有关的地方都必须跟着变；文档里残留旧限值 = 这条要求没做到。

（探针原样跑出来的结果：Excel 只有 I/J 两列随规则变，Word 的「预警状态」列和
「预警与异常」清单随规则变 —— 下面就是把这些事实钉成断言。）
"""
import io
from datetime import date
from pathlib import Path

import openpyxl
from docx import Document

from tests.test_control_values import (
    _ctl,
    _download_master,
    _master_with_control,
    _put_rules,
    _rules,
    _upload_master,
)
from tests.test_master_report import _seed_project

BACKEND = Path(__file__).resolve().parent.parent / "backend"
SRC = BACKEND.parent / "frontend" / "src"

LOOSE = [_rules("settlement", 30, 10), _rules("displacement", 20, 3)]
TIGHT = [_rules("settlement", 12, 1.5), _rules("displacement", 8, 1)]


# ─────────────────────────── 读文档 ───────────────────────────


def _cells(wb: openpyxl.Workbook) -> dict[str, str]:
    """整卷所有非空单元格 → {Sheet!坐标: 文本}（公式按算出来的值读）。"""
    out: dict[str, str] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and str(c.value).strip() != "":
                    out[f"{ws.title}!{c.coordinate}"] = str(c.value)
    return out


def _docx_text(client, pid: int, obs: str = "2026-08-21") -> tuple[dict, dict]:
    """Word 日报：返回 (表格单元格, 段落)。"""
    r = client.get(f"/api/v1/projects/{pid}/reports/daily", params={"obs_date": obs})
    assert r.status_code == 200, r.text
    doc = Document(io.BytesIO(r.content))
    paras = {f"para{i}": p.text for i, p in enumerate(doc.paragraphs) if p.text.strip()}
    cells: dict[str, str] = {}
    for ti, t in enumerate(doc.tables):
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                if cell.text.strip():
                    cells[f"tbl{ti}!{ri},{ci}"] = cell.text
    return cells, paras


def _std(client, pid: int, obs: str = "2026-08-21") -> dict:
    r = client.get(f"/api/v1/projects/{pid}/reports/excel", params={"obs_date": obs})
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
    try:
        return _cells(wb)
    finally:
        wb.close()


# ─────────────── ① 控制值两列：三条路径都按新规则重写 ───────────────


def test_only_control_cells_move_when_the_rule_moves(client, tmp_path):
    """改限值后重新导出：**只有**「控制值」两列该变，别的一个字都不许多变。

    反过来说也成立 —— 该变的没变就是"文档没跟规则走"。所以这里两头都断：
    ① 前后差异必须全部落在数据区的 I/J 上（说明没有别的地方偷偷抄了限值）；
    ② I/J 的每一行都真的出现在差异里（说明补出来的点位行也没抄母版旧值）。
    """
    pid = _seed_project(client)
    tid = _upload_master(client, pid, _master_with_control(tmp_path))

    _put_rules(client, pid, LOOSE)
    before = {"std": _std(client, pid), "master": _cells(_download_master(client, pid, tid))}
    _put_rules(client, pid, TIGHT)
    after = {"std": _std(client, pid), "master": _cells(_download_master(client, pid, tid))}

    for path in ("std", "master"):
        diff = {c for c in before[path] if before[path][c] != after[path].get(c, "")}
        assert diff, f"{path}：改了规则，导出的文档却一格没变"
        assert all(c in after[path] for c in diff), f"{path}：有格子被抹掉了"
        for coord in diff:
            sheet, _, addr = coord.partition("!")
            col = "".join(ch for ch in addr if ch.isalpha())
            row = int("".join(ch for ch in addr if ch.isdigit()))
            assert sheet in ("竖向位移", "r.立柱结构水平位移"), f"{path}：{coord} 也跟着变了"
            assert col in ("I", "J"), f"{path}：{coord} 不该随规则变"
            assert row >= 7, f"{path}：表头行 {coord} 被写坏了"
        # 8 个点：沉降 8 行 × 2 列 + 位移 16 行 × 2 列
        assert len(diff) == (8 + 16) * 2, f"{path}：只有 {len(diff)} 格变了，点位行漏写"


def test_control_values_equal_the_current_rule_row_by_row(client, tmp_path):
    """逐行核对：I 列 = 规则的累计控制值，J 列 = 规则的速率控制值（两条导出路径同口径）。"""
    pid = _seed_project(client)
    tid = _upload_master(client, pid, _master_with_control(tmp_path))

    for rules, want_s, want_d in ((LOOSE, (30, 10), (20, 3)), (TIGHT, (12, 1.5), (8, 1))):
        _put_rules(client, pid, rules)
        for wb_get in (
            lambda: openpyxl.load_workbook(io.BytesIO(client.get(
                f"/api/v1/projects/{pid}/reports/excel", params={"obs_date": "2026-08-21"}
            ).content), data_only=True),
            lambda: _download_master(client, pid, tid),
        ):
            wb = wb_get()
            try:
                ws, wsd = wb["竖向位移"], wb["r.立柱结构水平位移"]
                for r in range(7, 15):  # 8 个点（母版 2 行 + 自动补 6 行）
                    assert (_ctl(ws, r, 9), _ctl(ws, r, 10)) == want_s, f"{r}行 竖向位移 {want_s}"
                for r in range(7, 23):  # 8 个点 × X/Y 两行
                    assert (_ctl(wsd, r, 9), _ctl(wsd, r, 10)) == want_d, f"{r}行 水平位移 {want_d}"
            finally:
                wb.close()


# ─────────────── ② Word 日报的判定文字：按新规则重算 ───────────────


def test_docx_warning_column_and_list_follow_the_rule_change(client):
    """同一批读数，宽松规则下报表"全部正常"，收紧后同一行必须变成"预警"。

    注意是**水平位移**先变成预警（速率限值 1.0，本期速率正好 1.0），沉降仍是正常
    （限值 1.5）—— 各指标各按自己的规则判，这也是一一对应的一部分。
    """
    pid = _seed_project(client)

    _put_rules(client, pid, LOOSE)
    cells_loose, paras_loose = _docx_text(client, pid)
    assert all(v == "正常" for v in cells_loose.values() if v in ("正常", "预警", "报警"))
    assert any("未触发预警" in t for t in paras_loose.values()), paras_loose

    _put_rules(client, pid, TIGHT)
    cells_tight, paras_tight = _docx_text(client, pid)
    levels = [v for v in cells_tight.values() if v in ("正常", "预警", "报警")]
    assert "预警" in levels and "报警" not in levels
    hits = [t for t in paras_tight.values() if t.startswith("【预警】")]
    assert hits, paras_tight
    assert all("位移" in t for t in hits), hits          # 只有水平位移超了
    assert all("达到控制值" in t for t in hits), hits     # 措辞是"控制值"，不是"报警值"


def test_deactivated_alarm_rule_stops_judging(client):
    """规则**停用**（is_active=False）= 系统不再判这一级，文档必须跟着不判。

    拿「报警」级做实验最干净：停掉报警、控制值还在 → 文档照样出
    （拦不拦只看控制值有没有，见 `ensure_control_rules`），
    但同一行从【报警】降回【预警】。
    """
    pid = _seed_project(client)
    on = [
        _rules("settlement", 0.5, 0.5), _rules("settlement", 1.0, 1.0, level="alarm"),
        _rules("displacement", 0.5, 0.5), _rules("displacement", 1.0, 1.0, level="alarm"),
    ]
    _put_rules(client, pid, on)
    _, paras = _docx_text(client, pid)
    assert any(t.startswith("【报警】") for t in paras.values()), paras

    _put_rules(client, pid, [dict(r, is_active=False) if r["level"] == "alarm" else r for r in on])
    _, paras_off = _docx_text(client, pid)
    assert not any(t.startswith("【报警】") for t in paras_off.values()), paras_off
    assert any(t.startswith("【预警】") for t in paras_off.values()), paras_off


# ─────────────── ③ 「本次变化量 ΔV」：只进 Word 的判定，不进表格 ───────────────


def _delta_rules(delta_control: float) -> list[dict]:
    """按现有 LOOSE 加一档 ΔV（累计/速率照旧，只多这一维）。"""
    return [dict(r, threshold_delta_mm=delta_control) for r in LOOSE]


def test_delta_limit_does_not_touch_the_excel(client, tmp_path):
    """填不填 ΔV，导出的 Excel 必须**一格不差**。

    版式里没有「本次变化量控制值」这一列。要是哪天有人把它印上去，表上就多出一个
    在系统「控制值」两列里查不到出处的数 —— 2026-10-08 那句"表里写 30、系统按 20 判"
    就是这么来的。用户 2026-10-16 定的口径是**表里印的数 = 系统判预警的数**，
    所以新加的判定维度宁可不上表，也不能上表之后和判定脱钩。
    """
    pid = _seed_project(client)
    tid = _upload_master(client, pid, _master_with_control(tmp_path))

    _put_rules(client, pid, LOOSE)
    std_without = _std(client, pid)
    master_without = _cells(_download_master(client, pid, tid))

    _put_rules(client, pid, _delta_rules(8))
    assert _std(client, pid) == std_without, "ΔV 上了标准版式的表格"
    assert _cells(_download_master(client, pid, tid)) == master_without, "ΔV 上了母版套打"


def test_docx_names_the_dimension_that_triggered(client):
    """ΔV 触发的预警，Word 里必须点名「本次变化量」并把那个数一起印出来。

    数据：相隔 29 天累计 +30mm → 速率只有约 1mm/d。限值 累计 40、速率 3 都不超，
    ΔV 控制值 20 → 判预警。若文案照旧只印累计与速率，读者看到的是
    "累计 30、速率 1.03 …达到控制值" —— 两个数都没超，等于文档在说谎。
    """
    from tests.test_warning_rule_library import _mk_project_with_data

    pid = _mk_project_with_data(client, "D", "丁基坑")
    _put_rules(client, pid, [
        {
            "indicator": "settlement",
            "level": "control",
            "threshold_cumulative_mm": 40,
            "threshold_delta_mm": 20,
            "threshold_rate_mm_per_day": 3,
        },
    ])

    cells, paras = _docx_text(client, pid, obs="2026-09-30")
    hits = [t for t in paras.values() if t.startswith("【预警】")]
    assert hits, f"ΔV 超了却文档里没预警：{paras}"
    assert all("本次变化量达到控制值" in t for t in hits), hits
    assert all("本次 30" in t for t in hits), f"点名了 ΔV 就得把这个数一起印出来：{hits}"
    assert not any("累计变化量达到" in t or "速率达到" in t for t in hits), hits
    assert any(v == "预警" for v in cells.values()), "「预警状态」列也要跟着变"


# ─────────────── ④ 已导出的旧文件：改规则后要能查出来 ───────────────


def test_tightened_cumulative_limit_flags_old_exports(client, tmp_path):
    """文件检查原来只判速率超限。把**累计**控制值调小之后，旧报表里早已超新限值的行
    一条都不报 —— 用户改的正是这条规则，检查却对它无感（2026-10-14 补）。
    """
    from tests.test_export_audit import _kinds, _mk_project_with_data, _mk_rule_set, _run, _save_export

    pid = _mk_project_with_data(client, "PRJ-G", "庚基坑")  # 07-31：累计 -1.0mm、速率 -1.0mm/d
    _mk_rule_set(client, "庚规则", [pid], cum=30.0, rate=10.0)
    _save_export(client, pid, "2026-07-31", tmp_path)
    assert _run(client, tmp_path)["problem_rows"] == 0

    # 只把**累计**限值收到 1.0（速率仍宽松）→ 必须只报累计超限，不报速率超限
    _mk_rule_set(client, "庚规则2", [pid], cum=1.0, rate=10.0)
    res = _run(client, tmp_path)
    kinds = _kinds(res)
    assert "cum_over" in kinds, res["items"]
    assert "rate_over" not in kinds, "只改了累计，速率不该跟着报"
    it = next(i for i in res["items"] if any(p["kind"] == "cum_over" for p in i["problems"]))
    assert it["levels"]["report"] == "control", it


def _old_style_template(tmp_path):
    """造一份**旧式模板**（只认数据表版式，走 `build_from_template` 那条路）。

    I/J 上预先印着 30/10 —— 模拟"母版是从别的基坑那期另存的"。
    """
    from openpyxl import Workbook

    from backend.models.export_template import ExportTemplate

    wb = Workbook()
    ws = wb.active
    ws.title = "竖向位移"
    ws["A5"] = "监测\n点号"
    for col, title in ((5, "上次累计（mm)"), (6, "本次累计（mm)"), (7, "本次变化（mm)"), (8, "速率（mm/d）")):
        ws.cell(row=5, column=col, value=title)
    ws["I5"] = "控制值"
    ws["I6"] = "累计变化量（mm)"
    ws["J6"] = "变化速率（mm/d）"
    for r in (7, 8):
        ws.cell(row=r, column=9, value=30)
        ws.cell(row=r, column=10, value=10)
    path = tmp_path / "old_style.xlsx"
    wb.save(path)
    wb.close()
    return ExportTemplate(
        name="旧式模板",
        template_type="excel",
        file_path=str(path),
        layout_json={
            "sheets": [{
                "sheet_name": "竖向位移",
                "kind": "settlement",
                "mode": "per_row",
                "layout": {
                    "data_start_row": 7, "point_col": 1, "v0_m": 2, "prev_m": 3, "current_m": 4,
                    "prev_cum_mm": 5, "cum_mm": 6, "delta_mm": 7, "rate": 8,
                    "control_cum": 9, "control_rate": 10, "warning_level": 11,
                },
            }]
        },
    )


def test_every_path_agrees_when_rules_are_set_missing_or_half(client, tmp_path):
    """四条生成路径对**同一种情况**必须给同一个结果 —— 三种情况都验。

    用户 2026-10-14 定：空规则统一"拦住生成"，与生成数据同一个口径。
    实测出来过的不一致，都在这条里钉死：
    · 没配规则时母版套打保留底模上印的 30/10、旧模板版式把格子写空 —— 换个模板换个口径；
    · 数据生成被 422 拦下、日报却照样能出 —— 换个入口换个口径。
    半配（只填累计）不拦，但**缺的那一维一律留空**：写 0 会让所有点变成超限，
    留住底模原值则是在成果文件上印一个系统里查不到的数。
    """
    from backend.core.db import SessionLocal
    from backend.reports.template_excel import build_from_template

    tpl = _old_style_template(tmp_path)

    def _old_values(pid: int) -> dict:
        with SessionLocal() as db:
            buf, _, _ = build_from_template(db, pid, tpl, date(2026, 8, 21))
        wb = openpyxl.load_workbook(io.BytesIO(buf.getvalue()), data_only=True)
        try:
            return _cells(wb)
        finally:
            wb.close()

    def _ij(cells: dict, row: int) -> tuple:
        return _ctl_from_text(cells.get(f"竖向位移!I{row}")), _ctl_from_text(cells.get(f"竖向位移!J{row}"))

    def _excel_detail(client, pid: int, **params) -> str:
        r = client.get(f"/api/v1/projects/{pid}/reports/excel", params={"obs_date": "2026-08-21", **params})
        assert r.status_code == 422, f"没拦住，出了 200：{r.text[:120]}"
        return r.json()["detail"]

    # ① 配了规则 → 三条 Excel 路印的是同一个数
    with_rules = _seed_project(client, pno="PRJ-R")
    _put_rules(client, with_rules, [_rules("settlement", 25, 8), _rules("displacement", 6, 2)])
    tid = _upload_master(client, with_rules, _master_with_control(tmp_path))
    old, std, mas = _old_values(with_rules), _std(client, with_rules), _cells(_download_master(client, with_rules, tid))
    for r in range(7, 15):  # 8 个沉降测点
        assert _ij(old, r) == (25, 8), f"旧模板 r{r} 没按规则写：{_ij(old, r)}"
        assert _ij(std, r) == (25, 8), f"标准版式 r{r}：{_ij(std, r)}"
        assert _ij(mas, r) == (25, 8), f"母版套打 r{r}：{_ij(mas, r)}"

    # ② 一条控制值都没配 → 四条路**全部拦住**，且话里指到「预警规则」
    bare = _seed_project(client, pno="PRJ-N", rules=False)
    bare_tid = _upload_master(client, bare, _master_with_control(tmp_path))
    for detail in (
        _excel_detail(client, bare),                                  # 标准版式
        _excel_detail(client, bare, template_id=bare_tid),            # 母版套打
        _blocked_old(bare, tpl),                                      # 旧模板版式（进程内直调）
    ):
        assert "预警规则" in detail, detail
    r = client.get(f"/api/v1/projects/{bare}/reports/daily", params={"obs_date": "2026-08-21"})
    assert r.status_code == 422 and "预警规则" in r.json()["detail"], r.text

    # ③ 只填累计、没填速率 → I 列照写，J 列**留空**（三条 Excel 路一致）
    half = _seed_project(client, pno="PRJ-H")
    _put_rules(client, half, [_rules("settlement", 25, None), _rules("displacement", 6, None)])
    half_tid = _upload_master(client, half, _master_with_control(tmp_path))
    for name, cells in (
        ("旧模板", _old_values(half)),
        ("标准版式", _std(client, half)),
        ("母版套打", _cells(_download_master(client, half, half_tid))),
    ):
        for r in range(7, 15):
            assert _ij(cells, r) == (25, None), f"{name} r{r}：{_ij(cells, r)}（没填的速率既不许写 0，也不许留底模旧值）"


def _blocked_old(pid: int, tpl) -> str:
    """旧模板那条路是进程内直接调的 → 断它抛 `ValueError`（API 层再转成 422）。"""
    from backend.core.db import SessionLocal
    from backend.reports.template_excel import build_from_template

    with SessionLocal() as db:
        try:
            build_from_template(db, pid, tpl, date(2026, 8, 21))
        except ValueError as e:
            return str(e)
    raise AssertionError("旧模板版式没拦：项目一条控制值都没有还能出日报")


def _ctl_from_text(text: str | None) -> float | None:
    return None if text is None else round(float(text), 3)


# ─────────────── ⑤ 阈值取数只有一个出处（源码守卫） ───────────────
def test_limits_are_read_from_one_place_only() -> None:
    """`RuleLimits` 只允许在 `warning_service.rule_limits` 里构造。

    以前 calc / comparison / warning_event / export_audit 各抄了一遍同样的四行取值。
    规则加一个条件、或某处漏读一个字段，代码上看不出任何异常，结果就是
    "页面按新规则判、文档按旧口径判" —— 用户报的正是这类对不上。
    """
    sites = {
        py.relative_to(BACKEND).as_posix()
        for py in BACKEND.rglob("*.py")
        if "RuleLimits(" in py.read_text(encoding="utf-8")
    }
    # （`warnings/engine.py` 里是 `class RuleLimits:` 定义，不带括号，所以不在sites里）
    assert sites == {"services/warning_service.py"}, sites


def test_audit_service_does_not_reread_rules_itself() -> None:
    """文件检查不许自己拼规则，只能用 `rule_limits` / `report_control_values`。"""
    text = (BACKEND / "services" / "export_audit_service.py").read_text(encoding="utf-8")
    assert "effective_rules(" not in text, "又自己读了一遍生效规则 → 判定口径会出现第二套"
    assert "rule_limits" in text and "report_control_values" in text


def test_every_export_path_checks_rules_before_writing() -> None:
    """四条生成路径**都**要先过 `ensure_control_rules` —— 漏一条就等于留了个后门。

    "空规则拦住"这条口径最容易被绕过去的地方正是路径数量：Excel 有三条生成路
    （标准版式 / 母版套打 / 旧模板版式）加一份 Word 日报，只拦其中两条的话，
    用户换一种导出方式照样能拿到那份印着系统里查不到的控制值的报表。
    """
    for rel in (
        "reports/excel_report.py",
        "reports/master_report.py",
        "reports/template_excel.py",
        "reports/docx_report.py",
    ):
        text = (BACKEND / rel).read_text(encoding="utf-8")
        assert "ensure_control_rules(" in text, f"{rel} 没在生成前拦空规则"
    # 数据生成走的是同一条判据（不是又写一个"缺不缺"的判断）
    gen = (BACKEND / "services" / "data_generator_service.py").read_text(encoding="utf-8")
    assert "has_control_rule(" in gen, "生成数据又自己写了一套'规则缺不缺'的判据"


# ─────────────── ⑥ 界面上把"文档会怎么变"说清楚（源码守卫） ───────────────
def test_modal_says_which_docs_follow_the_rule() -> None:
    """改规则时，用户必须能看到"新生成的文档跟着变、老文件不变"。

    「一一对应」这条要求最容易漏的不是代码，而是**期望**：
    改完限值的人默认手头那份日报也一起变了，实际不是 —— 于是拿着旧文件来报"没同步"。
    """
    page = (SRC / "components" / "WarningRulesModal.tsx").read_text(encoding="utf-8")
    assert "Word 日报" in page and "控制值" in page, "没说清楚哪些文档跟着这条规则走"
    assert "之前已经导出的老文件" in page and "重新导出" in page, "没说老文件不会自动更新"
    # 保存后的提示：改了限值必须当场提醒"老文件还是旧值"，光靠页面下方的小字没人看
    assert "limitsChanged" in page and "文件检查" in page, "保存时没有针对'改了限值'的提示"
    assert "cum_over" in (SRC / "pages" / "AuditPage.tsx").read_text(encoding="utf-8"), (
        "检查页不认「累计超限」这个问题类型，改累计限值就翻不出旧报表"
    )
