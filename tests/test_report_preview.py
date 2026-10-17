# -*- coding: utf-8 -*-
"""导出 Excel 的界面预览（`GET /projects/{pid}/reports/preview`）。

用户 2026-10-16：「把导出Excel可视化 做进系统中吧」，路线定为**后端出预览数据 + 前端 antd 表格**
（不为此引 SheetJS/Univer：前端现在只有 antd/axios/dayjs/echarts，免安装包要多背约 1MB，
而且前端自己解析文件不知道系统规则，没法核对「控制值」）。

因此这个功能值钱的地方有两处，测试就盯这两处：
1. **同源** —— 预览必须走 `_build_excel_one`，界面画的格子 = 下载文件的格子；
2. **核对** —— 预览顺带逐格比对「控制值」是不是当前预警规则里的那对数，
   并把导出够不着的母版遗留格子（界线之外还印着数）单独点名。
"""
import hashlib

import openpyxl

from backend.core.config import settings
from tests.conftest import put_control_rules
from tests.test_control_values import (
    _download_master,
    _master_with_control,
    _master_with_reserved_rows,
    _upload_master,
)
from tests.test_master_report import _make_master, _seed_project


def _preview(client, pid, **params):
    r = client.get(f"/api/v1/projects/{pid}/reports/preview", params=params)
    assert r.status_code == 200, r.text
    return r.json()


def _cells(doc, sheet_name):
    return {(c["r"], c["c"]): c for c in next(s for s in doc["sheets"] if s["name"] == sheet_name)["cells"]}


def _template_files() -> dict[str, str]:
    """母版原件目录的指纹 —— 预览只许读，一个字节都不许改。"""
    return {
        str(p.relative_to(settings.templates_dir)): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in settings.templates_dir.rglob("*") if p.is_file()
    }


# ─────────────────────── ① 内置标准版式 ───────────────────────


def test_builtin_preview_draws_the_standard_report(client):
    pid = _seed_project(client)
    doc = _preview(client, pid, obs_date="2026-08-21")
    assert doc["filename"].endswith(".xlsx")
    names = [s["name"] for s in doc["sheets"]]
    assert "竖向位移" in names, names
    assert doc["truncated"] is False
    # 表头 + 8 个点位行都画得出来（界面拿到的是非空网格，不是一句"暂无数据"）
    assert len(_cells(doc, "竖向位移")) > 20
    # 标准版式没有母版原件可核对 → 不做控制值核对（每行本就取自同一个规则，没有第二个出处）
    assert doc["control_checks"] == []


# ─────────────────────── ② 预览 = 下载的文件 ───────────────────────


def test_preview_grid_is_the_downloaded_file(client, tmp_path):
    """同一条生成路径：把下载回来的 xlsx 再转一遍，必须与预览 JSON 逐格相同。

    这条断言就是"预览所见即下载所得"的保证书 —— 哪天有人给预览另写一套取数，这里会红。
    """
    from backend.reports.preview_model import bytes_to_preview

    pid = _seed_project(client)
    tid = _upload_master(client, pid, _make_master())
    params = {"template_id": tid, "obs_date": "2026-08-21", "weather": "晴"}
    doc = _preview(client, pid, **params)
    resp = client.get(f"/api/v1/projects/{pid}/reports/excel", params=params)
    assert resp.status_code == 200, resp.text
    assert doc["sheets"] == bytes_to_preview(resp.content)["sheets"]


def test_preview_shows_dates_not_serial_numbers(client, tmp_path):
    """封面日期格在文件里是 46255、Excel 按日期格式显示 2026-08-21 —— 预览也得是后者。"""
    src = _make_master()
    wb = openpyxl.load_workbook(src)
    wb["封面"]["F16"].number_format = "yyyy-mm-dd"  # 真实母版上这一格是日期格式
    p = tmp_path / "master_cover_date.xlsx"
    wb.save(p)
    wb.close()

    pid = _seed_project(client)
    tid = _upload_master(client, pid, p)
    doc = _preview(client, pid, template_id=tid, obs_date="2026-08-21")
    cover = _cells(doc, "封面")
    assert cover[(16, 6)]["v"] == "2026-08-21", cover[(16, 6)]


# ─────────────────────── ③ 控制值逐格核对 ───────────────────────


def test_preview_control_values_equal_the_current_rules(client, tmp_path):
    """母版上印着 30/10、20/3，规则改成 25/8、18/2.5 → 预览里每一格都得是新数。"""
    p = _master_with_control(tmp_path, s=(30, 10), d=(20, 3))
    pid = _seed_project(client, rules=False)
    put_control_rules(client, pid, [
        {"indicator": "settlement", "level": "control", "threshold_cumulative_mm": 25.0,
         "threshold_rate_mm_per_day": 8.0, "is_active": True},
        {"indicator": "displacement", "level": "control", "threshold_cumulative_mm": 18.0,
         "threshold_rate_mm_per_day": 2.5, "is_active": True},
    ])
    tid = _upload_master(client, pid, p)
    doc = _preview(client, pid, template_id=tid, obs_date="2026-08-21")

    checks = {c["sheet"]: c for c in doc["control_checks"]}
    assert set(checks) == {"竖向位移", "r.立柱结构水平位移"}, list(checks)
    for sheet, kind, cum, rate in [("竖向位移", "settlement", 25.0, 8.0),
                                   ("r.立柱结构水平位移", "displacement", 18.0, 2.5)]:
        c = checks[sheet]
        assert c["kind"] == kind and c["ok"] is True, c
        assert c["expected"] == {"control_cum": cum, "control_rate": rate}, c
        assert c["checked_rows"] >= 1 and c["mismatch"] == [] and c["residual"] == []
    # 界面上画出来的那一格也真是新数（不是只在核对结论里改了口）
    assert _cells(doc, "竖向位移")[(7, 9)]["v"] == "25"
    assert _cells(doc, "竖向位移")[(7, 10)]["v"] == "8"


def test_preview_names_stale_values_the_export_cannot_reach(client, tmp_path):
    """数据区界线**之外**还印着控制值 → 导出不会碰它，预览必须点名（residual）。

    这正是"报表上半对得上一半对不上"那类问题的形态：界线以下那一格是母版遗留。
    """
    p = _master_with_reserved_rows(tmp_path, s=(30, 10))
    wb = openpyxl.load_workbook(p)
    wb["竖向位移"].cell(row=20, column=9, value=30)  # 说明行之后的孤立旧值
    wb.save(p)
    wb.close()

    # 一个测点（= 母版点位行数）：不让"自动补行"参与，数据区界线才是稳定的 7~12
    pid = client.post("/api/v1/projects", json={"code": "PRJ-RESID", "name": "遗留值测试"}).json()["id"]
    pt = client.post(f"/api/v1/projects/{pid}/points", json={"point_no": "JJC1", "point_type": "both"}).json()
    for d in ("2026-08-20", "2026-08-21"):
        for st, comp in (("settlement", "value"), ("displacement", "X"), ("displacement", "Y")):
            client.post(f"/api/v1/points/{pt['id']}/records",
                        json={"series_type": st, "component": comp, "obs_date": d, "value_mm": 4909.0})
    put_control_rules(client, pid, [
        {"indicator": "settlement", "level": "control", "threshold_cumulative_mm": 25.0,
         "threshold_rate_mm_per_day": 8.0, "is_active": True},
        {"indicator": "displacement", "level": "control", "threshold_cumulative_mm": 18.0,
         "threshold_rate_mm_per_day": 2.5, "is_active": True},
    ])
    tid = _upload_master(client, pid, p)
    doc = _preview(client, pid, template_id=tid, obs_date="2026-08-21")
    c = next(x for x in doc["control_checks"] if x["sheet"] == "竖向位移")
    # 数据区（含没点号的预留行）已全部改成 25/8
    assert c["area"] == [7, 12], c
    assert c["ok"] is False and c["residual"] == ["竖向位移!I20=30"], c
    # 而导出文件里那一格确实还是 30 —— 预览没有冤枉它
    assert _download_master(client, pid, tid)["竖向位移"]["I20"].value == 30


def test_half_configured_rule_leaves_the_missing_cell_blank(client, tmp_path):
    """规则只配了累计没配速率 → 速率那列留空是**对的**，不许报问题（缺的不兜底）。"""
    pid = _seed_project(client, rules=False)
    put_control_rules(client, pid, [
        # 沉降只填累计（速率那格就该空着）；位移两项都填，否则整份文档被 `ensure_control_rules` 拦下
        {"indicator": "settlement", "level": "control", "threshold_cumulative_mm": 25.0,
         "is_active": True},
        {"indicator": "displacement", "level": "control", "threshold_cumulative_mm": 18.0,
         "threshold_rate_mm_per_day": 2.5, "is_active": True},
    ])
    tid = _upload_master(client, pid, _master_with_control(tmp_path))
    doc = _preview(client, pid, template_id=tid, obs_date="2026-08-21")
    c = next(x for x in doc["control_checks"] if x["sheet"] == "竖向位移")
    assert c["expected"] == {"control_cum": 25.0, "control_rate": None}, c
    assert c["ok"] is True and c["blank"] == [], c
    assert (7, 10) not in _cells(doc, "竖向位移")  # 速率列整列留空 = 预览上就是空白格


def test_preview_does_not_blame_merged_note_rows(client, tmp_path):
    """母版的「变化量正负号说明」是**横跨控制值列**的合并大块，不许当成空白控制值点名。

    2026-10-16 十号基坑真机预览误报了 6 格空白，成因是这样：导出会把母版上多余的点位行
    （本项目 10 个点、母版 12 个）连同紧跟其后的说明块一起清掉，于是"点号列第一次出现文字"
    的行往后漂了 3 行，按行数出来的数据区界线穿进了 `A31:L33` 这块合并区。
    那几格在 Excel 里连自己的值都放不进去（从属格），报成"控制值空白"是冤枉。
    """
    src = _master_with_control(tmp_path, s=(30, 10), d=(20, 3))
    p = tmp_path / "master_merged_note.xlsx"
    wb = openpyxl.load_workbook(src)
    ws = wb["竖向位移"]
    ws.cell(row=9, column=1, value="变化量正负号说明：“-”为上升“+”为下沉")
    ws.merge_cells("A9:J10")  # 与真实母版同形状：说明块横跨到控制值两列
    wb.save(p)
    wb.close()

    # 项目只有 1 个点（母版有 2 行）→ 导出会把第 2 行和说明块的文字一起清掉
    pid = client.post("/api/v1/projects", json={"code": "PRJ-MERGED", "name": "合并说明块"}).json()["id"]
    pt = client.post(f"/api/v1/projects/{pid}/points", json={"point_no": "JJC1", "point_type": "both"}).json()
    for d in ("2026-08-20", "2026-08-21"):
        for st, comp in (("settlement", "value"), ("displacement", "X"), ("displacement", "Y")):
            client.post(f"/api/v1/points/{pt['id']}/records",
                        json={"series_type": st, "component": comp, "obs_date": d, "value_mm": 4909.0})
    put_control_rules(client, pid, [
        {"indicator": "settlement", "level": "control", "threshold_cumulative_mm": 25.0,
         "threshold_rate_mm_per_day": 8.0, "is_active": True},
        {"indicator": "displacement", "level": "control", "threshold_cumulative_mm": 18.0,
         "threshold_rate_mm_per_day": 2.5, "is_active": True},
    ])
    tid = _upload_master(client, pid, p)
    doc = _preview(client, pid, template_id=tid, obs_date="2026-08-21")
    c = next(x for x in doc["control_checks"] if x["sheet"] == "竖向位移")
    assert c["ok"] is True, c
    assert c["blank"] == [] and c["mismatch"] == [], c
    # 界面上的着色也别盖住说明块：核对过的行数不含合并区的从属格
    assert c["checked_rows"] == len(range(c["area"][0], c["area"][1] + 1)) - 2, c


def test_general_format_numbers_are_not_shown_in_scientific_notation():
    """General 格式的大数要按 Excel 的口径写全（最多 11 位有效数字），不许缩成 `3.41309e+06`。

    坐标类数据（初始值 3413090.672 m）在母版上就是 General 格式，
    预览里写成科学计数法会让人以为数据被算坏了。
    """
    from backend.reports.preview_model import _display

    assert _display(3413090.672, "General")[0] == "3413090.672"
    assert _display(0.1 + 0.2, "General")[0] == "0.3"  # Excel 也不显那一长串浮点尾巴
    assert _display(2.0, "0.0")[0] == "2.0"  # 自带格式的仍按格式补位


# ─────────────────────── ④ 只读 + 与导出同样的拦截 ───────────────────────
def test_preview_never_touches_the_master_file(client, tmp_path):
    pid = _seed_project(client)
    tid = _upload_master(client, pid, _master_with_control(tmp_path))
    before = _template_files()
    _preview(client, pid, template_id=tid, obs_date="2026-08-21")
    assert _template_files() == before


def test_preview_refuses_without_rules_like_excel_does(client):
    """没配控制值规则时导出被拦 422（用户 2026-10-14 定）—— 预览同一条路，也得拦。"""
    pid = _seed_project(client, rules=False)
    r = client.get(f"/api/v1/projects/{pid}/reports/preview", params={"obs_date": "2026-08-21"})
    assert r.status_code == 422, r.text


def test_preview_reuses_the_export_dispatcher():
    """守卫：预览不许自己另起一套生成/判据。

    这个功能的价值全在"预览 = 下载的那份"。哪天有人把 `/preview` 改成直接调
    `build_master_report` 或自己写一遍数据区界线，界面就会和成果文件各说各话。
    """
    from pathlib import Path

    api = (Path(__file__).resolve().parent.parent / "backend" / "api" / "v1" / "reports.py").read_text(encoding="utf-8")
    body = api.split("def preview_excel", 1)[1].split("@reports_router", 1)[0]
    assert "_build_excel_one(" in body, "预览必须与 /excel 共用 _build_excel_one"
    assert "build_master_report" not in body and "build_period_excel" not in body

    pm = (Path(__file__).resolve().parent.parent / "backend" / "reports" / "preview_model.py").read_text(encoding="utf-8")
    assert "from backend.services.template_check_service import data_area_rows" in pm, (
        "控制值核对的数据区界线必须与体检/导出同一个判据"
    )
    assert ".save(" not in pm, "预览只读，不许写文件"
