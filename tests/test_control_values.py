# -*- coding: utf-8 -*-
"""「报表里的控制值」与「系统里的预警规则」必须一一对应（2026-10-08 用户要求）。

用户原话：「表格里的控制值 变化速率和累计变化 这个数值的 和 系统里设置的预警规则一一对应上」

**改之前两条路径各说各话**：
1. 母版套打导出：I/J「控制值」直接**照抄母版那一行**（母版常是别的基坑的表，写的是那一期的旧值）
   —— 系统里改了限值，报表纹丝不动；
2. 导入：台账里的控制值读是读出来了（预览里也显示），但**从不进预警规则** ——
   用「批量建项目」建出来的项目根本没有规则，报表里却印着 30/10、20/3。

**现在的规则**：报表「控制值」两列的**唯一出处 = 项目预警规则的 control 级**
（`warning_service.report_control_values`，两条导出路径共用）；
台账里的控制值可以一键同步进规则（用户勾选才做，默认不勾）。
"""
import io
from pathlib import Path

import openpyxl
import pytest

from tests.test_import_parser import _make_wb  # 与真实台账同构（含控制值 30/10、20/3）
from tests.test_master_report import _make_master, _seed_project

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SRC = Path(__file__).resolve().parent.parent / "frontend" / "src"
BACKEND = Path(__file__).resolve().parent.parent / "backend"


# ---------- 工装 ----------


def _rules(indicator: str, cum, rate, level: str = "control") -> dict:
    return {
        "indicator": indicator,
        "level": level,
        "threshold_cumulative_mm": cum,
        "threshold_rate_mm_per_day": rate,
        "is_active": True,
    }


def _put_rules(client, pid: int, rules: list[dict]) -> None:
    r = client.put(f"/api/v1/projects/{pid}/warning-rules", json={"rules": rules})
    assert r.status_code == 200, r.text


def _download_master(client, pid: int, tid: int, obs_date: str = "2026-08-21") -> openpyxl.Workbook:
    resp = client.get(
        f"/api/v1/projects/{pid}/reports/excel",
        params={"template_id": tid, "obs_date": obs_date, "weather": "晴"},
    )
    assert resp.status_code == 200, resp.text
    return openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)


def _upload_master(client, pid: int, path: Path, name: str = "报送母版") -> int:
    with open(path, "rb") as fh:
        r = client.post(
            f"/api/v1/projects/{pid}/export-templates/upload",
            data={"name": name},
            files={"file": (path.name, fh, XLSX_MIME)},
        )
    assert r.status_code == 200, r.text
    return r.json()["id"]


def _master_with_control(tmp_path: Path, s=(30, 10), d=(20, 3)) -> Path:
    """同构母版 + I/J 填上控制值（模拟"母版里本来写着 30/10、20/3"的真实情况）。"""
    wb = openpyxl.load_workbook(_make_master())
    ws = wb["竖向位移"]
    for r in range(7, 9):
        ws.cell(row=r, column=9, value=s[0])
        ws.cell(row=r, column=10, value=s[1])
    wsd = wb["r.立柱结构水平位移"]
    for r in (7, 8):
        wsd.cell(row=r, column=9, value=d[0])
        wsd.cell(row=r, column=10, value=d[1])
    out = tmp_path / "master_with_control.xlsx"
    wb.save(out)
    wb.close()
    return out


def _master_with_reserved_rows(tmp_path: Path, s=(30, 10)) -> Path:
    """十号母版那种形状：点位行下面还有几行**没点号的预留空行**，可 I/J 印着旧的 30/10。

    再在下面放一行「变化量正负号说明」——它是数据区的界线，本身没有控制值。
    """
    wb = openpyxl.load_workbook(_make_master())
    ws = wb["竖向位移"]
    for r in range(7, 13):  # 7~8 是点位行，9~12 是空白预留行
        ws.cell(row=r, column=9, value=s[0])
        ws.cell(row=r, column=10, value=s[1])
    ws.cell(row=13, column=1, value="变化量正负号说明：“-”为上升“+”为下沉")
    out = tmp_path / "master_with_reserved_rows.xlsx"
    wb.save(out)
    wb.close()
    return out


def _ctl(ws, row: int, col: int):
    """读出控制值（统一成 float 比较，能同时覆盖整数与非整数限值）。"""
    v = ws.cell(row=row, column=col).value
    return None if v is None else round(float(v), 3)


def _setup_points(client, pno="PRJ001"):
    """项目 + JJC1/JJC2（台账里就这两个点）。"""
    pid = client.post("/api/v1/projects", json={"code": pno, "name": "控制值测试"}).json()["id"]
    for no in ("JJC1", "JJC2"):
        client.post(f"/api/v1/projects/{pid}/points", json={"point_no": no, "point_type": "both"})
    return pid


def _upload_ledger(client, pid: int):
    wb = _make_wb()
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return client.post(
        f"/api/v1/projects/{pid}/import/upload",
        files={"file": ("八号基坑（第一期）2026.7.31.xlsx", buf.getvalue(), XLSX_MIME)},
    )


def _import_rows(data: dict) -> list[dict]:
    """把预览里的有效行转成确认接口要的形状（照前端 `withDate` 的口径补日期）。"""
    rows = []
    for r in data["valid_rows"]:
        rows.append(
            {
                "point_no": r["point_no"],
                "series_type": r["series_type"],
                "component": r["component"],
                "obs_date": r["obs_date"] or "2026-07-31",
                "value_mm": r["value_mm"],
                "initial_mm": r.get("initial_mm"),
                "remark": r.get("remark"),
            }
        )
    return rows


# ---------- ① 导出：控制值按预警规则写 ----------


def test_master_export_writes_control_from_rules(client, tmp_path):
    """母版里写 30/10、20/3，但项目规则是 25/8、18/2.5 → 导出的必须是**规则**的数。

    包括**自动补出来的点位行**（项目 8 个点、母版只有 2 个）——控制值是整表共同的设计限值，
    与该点当天有没有读数无关，每一行都要有。
    """
    pid = _seed_project(client)
    _put_rules(client, pid, [_rules("settlement", 25, 8), _rules("displacement", 18, 2.5)])
    tid = _upload_master(client, pid, _master_with_control(tmp_path))

    wb = _download_master(client, pid, tid)
    ws = wb["竖向位移"]
    for r in range(7, 15):  # 8 个测点（母版 2 行 + 自动补 6 行）
        assert _ctl(ws, r, 9) == 25, f"竖向位移 r{r} 累计控制值应取规则值 25"
        assert _ctl(ws, r, 10) == 8, f"竖向位移 r{r} 速率控制值应取规则值 8"

    wsd = wb["r.立柱结构水平位移"]
    for r in range(7, 23):  # 8 个测点 × 2 行（X/Y）
        assert _ctl(wsd, r, 9) == 18, f"水平位移 r{r} 累计控制值应取规则值 18"
        assert _ctl(wsd, r, 10) == 2.5, f"水平位移 r{r} 速率控制值应取规则值 2.5"
    wb.close()


def test_master_export_writes_rule_values_even_on_added_rows(client, tmp_path):
    """母版上印的 30/10 与规则里的 30/10 数值相同 —— 但**出处已经换了**，逐行按规则写。

    这条原来是断"没配规则时保留母版原值"；用户 2026-10-14 把口径改成
    「空规则和生成数据一样，一律拦住」（拦不住地留一个系统里查不到的数在成果文件上），
    拦不住的那个 case 见 `test_rule_change_in_documents`。
    这里改成验"逐行都按规则"：改一次限值，母版固有行和**自动补出来的行**必须一起变 ——
    补出来的行当初是抄母版那一行的 I/J 来的，那是同一张表里的第二个数值出处。
    """
    pid = _seed_project(client)  # 默认规则 30/10、20/3，与母版上印的一致
    tid = _upload_master(client, pid, _master_with_control(tmp_path))

    wb = _download_master(client, pid, tid)
    assert (_ctl(wb["竖向位移"], 7, 9), _ctl(wb["竖向位移"], 7, 10)) == (30, 10)
    assert (_ctl(wb["竖向位移"], 14, 9), _ctl(wb["竖向位移"], 14, 10)) == (30, 10)  # 补出来的行
    assert (_ctl(wb["r.立柱结构水平位移"], 7, 9), _ctl(wb["r.立柱结构水平位移"], 7, 10)) == (20, 3)
    wb.close()

    _put_rules(client, pid, [_rules("settlement", 22, 7), _rules("displacement", 15, 2)])
    wb2 = _download_master(client, pid, tid)
    for row in (7, 14):  # 固有行与补出来的行一起跟着改 = 没有第二个出处
        assert (_ctl(wb2["竖向位移"], row, 9), _ctl(wb2["竖向位移"], row, 10)) == (22, 7), row
    assert (_ctl(wb2["r.立柱结构水平位移"], 7, 9), _ctl(wb2["r.立柱结构水平位移"], 7, 10)) == (15, 2)
    wb2.close()


def test_master_reserved_blank_rows_also_follow_the_rule(client, tmp_path):
    """母版点位行下面**没点号的预留空行**也印着控制值 —— 必须一起按规则重写。

    用户 2026-10-16 报障：项目叫十号基坑，导出的表里 I/J 上半是规则里的 35/5、
    下半还是原来的 30/10。母版实测：竖向位移表 7~18 行有点号，19~26 行是空着的预留行，
    可**这 8 行每一行都印着 30/10**。`_master_point_blocks` 靠点号认行、扫到第一个空点号就停，
    所以那 8 行从来没被写过 —— 同一列出现两个数值出处，正是"表里写 30、系统按 35 判"的老毛病。
    """
    pid = client.post("/api/v1/projects", json={"code": "PRJ-BLANK", "name": "预留行测试"}).json()["id"]
    pt = client.post(f"/api/v1/projects/{pid}/points", json={"point_no": "JJC1", "point_type": "both"}).json()
    for d, v in (("2026-08-20", 4909.0), ("2026-08-21", 4910.0)):
        for st, comp in (("settlement", "value"), ("displacement", "X"), ("displacement", "Y")):
            client.post(f"/api/v1/points/{pt['id']}/records",
                        json={"series_type": st, "component": comp, "obs_date": d, "value_mm": v})
    _put_rules(client, pid, [_rules("settlement", 25, 8), _rules("displacement", 18, 2.5)])

    tid = _upload_master(client, pid, _master_with_reserved_rows(tmp_path))
    wb = _download_master(client, pid, tid)
    ws = wb["竖向位移"]
    for r in range(7, 13):  # 7~8 点位行 + 9~12 母版预留的空白点位行
        assert (_ctl(ws, r, 9), _ctl(ws, r, 10)) == (25, 8), f"竖向位移 r{r} 还是母版上印的 30/10"
    # 界线：数据区结束在「变化量正负号说明」那一行，它没有控制值、也不许被写上一个
    assert (ws.cell(row=13, column=9).value, ws.cell(row=13, column=10).value) == (None, None)
    wb.close()


def test_two_export_paths_use_the_same_control_values(client, tmp_path):
    """母版套打 与 系统标准版式，控制值必须是同一对数（两条路径同源）。"""
    pid = _seed_project(client)
    _put_rules(client, pid, [_rules("settlement", 27, 9), _rules("displacement", 19, 2.8)])
    tid = _upload_master(client, pid, _master_with_control(tmp_path))

    m = _download_master(client, pid, tid)
    assert (_ctl(m["竖向位移"], 7, 9), _ctl(m["竖向位移"], 7, 10)) == (27, 9)

    resp = client.get("/api/v1/projects/%d/reports/excel" % pid, params={"obs_date": "2026-08-21"})
    assert resp.status_code == 200, resp.text
    std = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)
    assert (_ctl(std["竖向位移"], 7, 9), _ctl(std["竖向位移"], 7, 10)) == (27, 9)
    m.close()
    std.close()


# ---------- ② 导入：台账里的控制值与规则对照，可选同步 ----------


def test_import_preview_reports_control_values_vs_rules(client):
    """台账里写着 30/10（沉降）、20/3（位移），项目还没配规则 → 预览要明说"对不上"。"""
    pid = _setup_points(client)
    resp = _upload_ledger(client, pid)
    assert resp.status_code == 200, resp.text
    data = resp.json()

    got = {s["indicator"]: s for s in data["control_suggest"]}
    assert set(got) == {"settlement", "displacement"}
    assert (got["settlement"]["control_cum_mm"], got["settlement"]["control_rate_mm_d"]) == (30, 10)
    assert (got["displacement"]["control_cum_mm"], got["displacement"]["control_rate_mm_d"]) == (20, 3)
    assert got["settlement"]["differs"] is True and got["settlement"]["has_rule"] is False


def test_confirm_without_sync_leaves_rules_untouched(client):
    """默认不同步：导数据就只是导数据，判定口径不能被顺手改掉。"""
    pid = _setup_points(client)
    data = _upload_ledger(client, pid).json()
    r = client.post(
        f"/api/v1/import/batches/{data['batch_id']}/confirm",
        json={"valid_rows": _import_rows(data)},
    )
    assert r.status_code == 200, r.text
    assert client.get(f"/api/v1/projects/{pid}/warning-rules").json() == []


def test_confirm_with_sync_writes_control_rules_end_to_end(client):
    """勾选同步后：台账的控制值 → 预警规则 → 导出的报表，三处是同一个数。

    这就是用户要的"一一对应"：从**台账里印着的数字**一路走到**报表里印着的数字**，
    中间没有任何一处是抄来的旧值。
    """
    pid = _setup_points(client)
    data = _upload_ledger(client, pid).json()
    r = client.post(
        f"/api/v1/import/batches/{data['batch_id']}/confirm",
        json={"valid_rows": _import_rows(data), "sync_control_values": True},
    )
    assert r.status_code == 200, r.text
    changes = r.json()["control_changes"]
    assert {c["indicator"] for c in changes} == {"settlement", "displacement"}
    assert all(c["created"] for c in changes), "原来没规则，应记成新建"

    rules = {(x["indicator"], x["level"]): x for x in client.get(f"/api/v1/projects/{pid}/warning-rules").json()}
    assert (rules[("settlement", "control")]["threshold_cumulative_mm"],
            rules[("settlement", "control")]["threshold_rate_mm_per_day"]) == (30, 10)
    assert (rules[("displacement", "control")]["threshold_cumulative_mm"],
            rules[("displacement", "control")]["threshold_rate_mm_per_day"]) == (20, 3)

    # 导出（标准版式）→ 表里的控制值就是台账里那两个数
    resp = client.get(f"/api/v1/projects/{pid}/reports/excel", params={"obs_date": "2026-07-31"})
    assert resp.status_code == 200, resp.text
    wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)
    assert (_ctl(wb["竖向位移"], 7, 9), _ctl(wb["竖向位移"], 7, 10)) == (30, 10)
    assert (_ctl(wb["r.立柱结构水平位移"], 7, 9), _ctl(wb["r.立柱结构水平位移"], 7, 10)) == (20, 3)
    wb.close()


def test_sync_only_touches_control_level(client):
    """同步**只动「控制值」**：报警值、别的指标一概不碰。

    台账里只有控制值这一对数字，报警值毫无依据 —— 凭它推算是编数据。

    ⚠ 这里的报警值（40/12、30/4）**必须比台账的控制值（30/10、20/3）大**：
    控制值涨到报警值之上 = 倒挂，同步会被整组拒收（见 `test_sync_refused_when_control_would_outrank_alarm`），
    这条用例要测的是"只动 control 级"，别让校验先把请求拦掉。
    """
    pid = _setup_points(client)
    _put_rules(
        client,
        pid,
        [
            _rules("settlement", 20, 2),
            _rules("settlement", 40, 12, level="alarm"),
            _rules("displacement", 18, 2),
            _rules("displacement", 30, 4, level="alarm"),
        ],
    )
    data = _upload_ledger(client, pid).json()
    client.post(
        f"/api/v1/import/batches/{data['batch_id']}/confirm",
        json={"valid_rows": _import_rows(data), "sync_control_values": True},
    )

    rules = {(x["indicator"], x["level"]): x for x in client.get(f"/api/v1/projects/{pid}/warning-rules").json()}
    assert (rules[("settlement", "control")]["threshold_cumulative_mm"],
            rules[("settlement", "control")]["threshold_rate_mm_per_day"]) == (30, 10)
    assert (rules[("displacement", "control")]["threshold_cumulative_mm"],
            rules[("displacement", "control")]["threshold_rate_mm_per_day"]) == (20, 3)
    # 报警值原样
    assert (rules[("settlement", "alarm")]["threshold_cumulative_mm"],
            rules[("settlement", "alarm")]["threshold_rate_mm_per_day"]) == (40, 12)
    assert (rules[("displacement", "alarm")]["threshold_cumulative_mm"],
            rules[("displacement", "alarm")]["threshold_rate_mm_per_day"]) == (30, 4)
    # 规则条数没变多（没被"整组重建"成只剩 control）
    assert len(rules) == 4


def test_sync_refused_when_control_would_outrank_alarm(client):
    """台账控制值比系统报警值还大 → **规则一条不改**，但数据照常入库、原因如实带回前端。

    倒挂的规则是废的（一超限直接落「报警」，「预警」这一级永远出不来），不能悄悄写进去。
    可数据在前一步已经 commit 了：接口不能因此报 500 —— 那会让人以为导入失败，
    而记录其实一条没丢（这正是本仓库踩过两次的"改了没反应 / 说有错其实没错"）。
    """
    pid = _setup_points(client)
    _put_rules(
        client,
        pid,
        [
            _rules("settlement", 20, 2),
            _rules("settlement", 40, 4, level="alarm"),  # 报警速率 4 < 台账控制速率 10
            _rules("displacement", 18, 2),
            _rules("displacement", 30, 3, level="alarm"),
        ],
    )
    data = _upload_ledger(client, pid).json()
    r = client.post(
        f"/api/v1/import/batches/{data['batch_id']}/confirm",
        json={"valid_rows": _import_rows(data), "sync_control_values": True},
    )
    assert r.status_code == 200, r.text
    assert r.json()["imported_rows"] > 0, "数据必须真的进了库"
    assert r.json()["control_changes"] == []
    body = r.json()["control_sync_error"]
    assert "报警值不能低于控制值" in body, body
    assert "变化速率" in body and "10" in body, f"要说清是哪一维、哪两个数：{body}"

    rules = {(x["indicator"], x["level"]): x for x in client.get(f"/api/v1/projects/{pid}/warning-rules").json()}
    assert (rules[("settlement", "control")]["threshold_rate_mm_per_day"]) == 2, "倒挂时一条规则都不许改"


def test_sync_is_noop_when_already_consistent(client):
    """台账值与现有规则一致 → 一条变更都不该报（别把"没变化"也说成改了）。"""
    pid = _setup_points(client)
    _put_rules(client, pid, [_rules("settlement", 30, 10), _rules("displacement", 20, 3)])
    data = _upload_ledger(client, pid).json()
    assert all(not s["differs"] for s in data["control_suggest"])

    r = client.post(
        f"/api/v1/import/batches/{data['batch_id']}/confirm",
        json={"valid_rows": _import_rows(data), "sync_control_values": True},
    )
    assert r.json()["control_changes"] == []


# ---------- 源码守卫 ----------


def test_report_control_values_has_a_single_source() -> None:
    """控制值取数只允许有一个出处，两条导出路径都得走它。

    否则又会回到"母版套打照抄母版、标准版式按规则"那种各说各话的局面。
    """
    excel = (BACKEND / "reports" / "excel_report.py").read_text(encoding="utf-8")
    master = (BACKEND / "reports" / "master_report.py").read_text(encoding="utf-8")
    assert "report_control_values" in excel and "effective_rules" not in excel, \
        "标准版式应改用唯一的取数函数，不再自己 merge 规则"
    assert "report_control_values" in master, "母版套打没有走控制值唯一出处"


def test_frontend_lets_user_choose_sync() -> None:
    """前端必须把"对不上"摆出来，并让用户自己决定要不要同步（默认不勾）。"""
    page = (SRC / "pages" / "ImportPage.tsx").read_text(encoding="utf-8")
    assert "control_suggest" in page, "预览没有把台账里的控制值拿出来对照"
    assert "syncControl" in page, "没有同步开关"
    assert "导入时把台账里的控制值同步到本项目预警规则" in page, "开关没说清是在干什么"
    modal = (SRC / "components" / "WarningRulesModal.tsx").read_text(encoding="utf-8")
    assert "报表" in modal and "控制值" in modal, "预警规则弹窗没说明它和报表「控制值」列是一回事"
