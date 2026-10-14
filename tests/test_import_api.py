# -*- coding: utf-8 -*-
"""导入引擎 API 集成测试（九步流程端到端）。"""
import io

import openpyxl
import pytest

from tests.test_import_parser import _make_wb


def _xlsx_bytes() -> bytes:
    wb = _make_wb()
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def _setup_project_with_points(client):
    """建项目 + 建 JJC1/JJC2 测点（both 类型，与模板一致）。"""
    p = client.post("/api/v1/projects", json={"code": "PRJ001", "name": "八号基坑"}).json()
    pid = p["id"]
    for no in ("JJC1", "JJC2"):
        client.post(
            f"/api/v1/projects/{pid}/points",
            json={"point_no": no, "point_type": "both", "location_desc": "桩顶"},
        )
    return pid


def test_import_full_flow(client):
    """九步流程端到端：上传 → 预览 → 确认 → 落库。"""
    pid = _setup_project_with_points(client)

    # Step 1-7：上传并预览
    resp = client.post(
        f"/api/v1/projects/{pid}/import/upload",
        files={"file": ("八号基坑（第一期）2026.7.31.xlsx", _xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["filename"].startswith("八号基坑")
    assert len(data["sheets"]) == 2  # 竖向位移 + 水平位移
    # 沉降 2 行 + 位移 4 行 = 6 行有效
    assert len(data["valid_rows"]) == 6
    assert len(data["conflict_rows"]) == 0
    assert len(data["error_rows"]) == 0
    batch_id = data["batch_id"]

    # Step 8-9：确认导入
    conf = client.post(f"/api/v1/import/batches/{batch_id}/confirm", json={"valid_rows": data["valid_rows"]})
    assert conf.status_code == 200
    assert conf.json()["imported_rows"] == 6

    # 落库校验：沉降 JJC1 已入 4.906m → 4906mm
    pts = client.get(f"/api/v1/projects/{pid}/points").json()
    jjc1 = next(x for x in pts if x["point_no"] == "JJC1")
    recs = client.get(f"/api/v1/points/{jjc1['id']}/records").json()
    assert len(recs) == 3  # 沉降 value + 位移 X + 位移 Y
    settle = next(r for r in recs if r["series_type"] == "settlement")
    assert settle["value_mm"] == pytest.approx(4906.0)
    assert settle["source"] == "excel"
    assert settle["import_batch_id"] == batch_id

    # 批次审计
    batches = client.get("/api/v1/import/batches").json()
    assert batches[0]["status"] == "confirmed"
    assert batches[0]["imported_rows"] == 6


def test_import_duplicate_conflict(client):
    """重复检查：同日同点已存在 → conflict 行（不覆盖）。"""
    pid = _setup_project_with_points(client)
    pts = client.get(f"/api/v1/projects/{pid}/points").json()
    jjc1 = next(x for x in pts if x["point_no"] == "JJC1")
    # 手工录入一条同日沉降记录
    client.post(
        f"/api/v1/points/{jjc1['id']}/records",
        json={"series_type": "settlement", "component": "value", "obs_date": "2026-07-31", "value_mm": 4906.0},
    )
    # 导入同文件 → 沉降 JJC1 行进入 conflict
    resp = client.post(
        f"/api/v1/projects/{pid}/import/upload",
        files={"file": ("dup.xlsx", _xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    data = resp.json()
    conflict_nos = {(r["point_no"], r["series_type"], r["component"]) for r in data["conflict_rows"]}
    assert ("JJC1", "settlement", "value") in conflict_nos
    assert len(data["valid_rows"]) == 5  # 6 - 1 冲突
    assert data["conflict_rows"][0]["existing_value_mm"] == 4906.0

    # 确认导入：只导入非冲突行，冲突行不被覆盖
    conf = client.post(f"/api/v1/import/batches/{data['batch_id']}/confirm", json={"valid_rows": data["valid_rows"]})
    assert conf.json()["imported_rows"] == 5
    recs = client.get(f"/api/v1/points/{jjc1['id']}/records", params={"series_type": "settlement"}).json()
    assert len(recs) == 1  # 仍是手工录入的那条


def test_import_unknown_point_error(client):
    """校验：测点不存在 → error 行。"""
    pid = _setup_project_with_points(client)
    # 在 Excel 中加一个模板外的测点 JJC9
    wb = _make_wb()
    ws = wb["竖向位移"]
    ws.cell(row=9, column=1, value="JJC9")
    ws.cell(row=9, column=2, value=5.0)
    ws.cell(row=9, column=4, value=4.9)
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    resp = client.post(
        f"/api/v1/projects/{pid}/import/upload",
        files={"file": ("unknown.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    data = resp.json()
    assert any(r["point_no"] == "JJC9" and "不存在" in r["error"] for r in data["error_rows"])


def test_import_reject_non_excel(client):
    """上传非 Excel → 400。"""
    pid = _setup_project_with_points(client)
    resp = client.post(
        f"/api/v1/projects/{pid}/import/upload",
        files={"file": ("a.txt", b"hello", "text/plain")},
    )
    assert resp.status_code == 400


def test_confirm_twice_forbidden(client):
    """同一批次二次确认 → 400（状态已 confirmed）。"""
    pid = _setup_project_with_points(client)
    resp = client.post(
        f"/api/v1/projects/{pid}/import/upload",
        files={"file": ("a.xlsx", _xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    data = resp.json()
    bid = data["batch_id"]
    assert client.post(f"/api/v1/import/batches/{bid}/confirm", json={"valid_rows": data["valid_rows"]}).status_code == 200
    assert client.post(f"/api/v1/import/batches/{bid}/confirm", json={"valid_rows": []}).status_code == 400


def _xlsx_bytes_with_date(d: str) -> bytes:
    """生成指定观测日期的模板 Excel（改 R3「本次监测时间」，用于多文件/多期导入测试）。"""
    wb = _make_wb()
    for sn in wb.sheetnames:
        ws = wb[sn]
        ws.cell(row=3, column=1, value=f"本次监测时间：{d}")
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_import_multiple_files_sequentially(client):
    """多文件一次导入（前端多选 → 逐个上传+确认）：各期数据分别入库，互不干扰。"""
    pid = _setup_project_with_points(client)
    files = [
        ("八号基坑（第十期）2026.08.01.xlsx", "2026-08-01"),
        ("八号基坑（第十一期）2026.08.02.xlsx", "2026-08-02"),
        ("八号基坑（第十二期）2026.08.03.xlsx", "2026-08-03"),
    ]
    for fn, d in files:
        up = client.post(f"/api/v1/projects/{pid}/import/upload", files={"file": (fn, _xlsx_bytes_with_date(d), _MIME)})
        assert up.status_code == 200, up.text
        data = up.json()
        assert data["conflict_rows"] == []
        assert len(data["valid_rows"]) == 6
        conf = client.post(f"/api/v1/import/batches/{data['batch_id']}/confirm", json={"valid_rows": data["valid_rows"]})
        assert conf.status_code == 200
        assert conf.json()["imported_rows"] == 6

    # 三个日期都落库（JJC1 沉降序列）
    pts = client.get(f"/api/v1/projects/{pid}/points").json()
    jjc1 = next(x for x in pts if x["point_no"] == "JJC1")
    recs = client.get(f"/api/v1/points/{jjc1['id']}/records").json()
    settle_dates = sorted(r["obs_date"] for r in recs if r["series_type"] == "settlement")
    assert settle_dates == ["2026-08-01", "2026-08-02", "2026-08-03"]


def test_import_same_period_twice_skips_no_overwrite(client):
    """同一期文件重复导入（多选里混入重复文件）：第二次整批冲突跳过，原值不被覆盖。"""
    pid = _setup_project_with_points(client)
    fn = "八号基坑（第十期）2026.08.01.xlsx"
    blob = _xlsx_bytes_with_date("2026-08-01")

    up1 = client.post(f"/api/v1/projects/{pid}/import/upload", files={"file": (fn, blob, _MIME)})
    d1 = up1.json()
    client.post(f"/api/v1/import/batches/{d1['batch_id']}/confirm", json={"valid_rows": d1["valid_rows"]})

    up2 = client.post(f"/api/v1/projects/{pid}/import/upload", files={"file": (fn, blob, _MIME)})
    assert up2.status_code == 200
    d2 = up2.json()
    assert len(d2["valid_rows"]) == 0
    assert len(d2["conflict_rows"]) == 6


def test_import_overrides_generated_data(client):
    """用户口径：以导入的数据为准，不以"系统编的"数据为准。

    同日已有 generated 记录 → 预览归入 override_rows；确认导入后以用户数据替换，
    旧生成记录标记 superseded（留痕、不再参与计算）。
    """
    import datetime

    from backend.core.db import SessionLocal
    from backend.models.observation_record import ObservationRecord

    pid = _setup_project_with_points(client)
    pts = client.get(f"/api/v1/projects/{pid}/points").json()
    jjc1 = next(x for x in pts if x["point_no"] == "JJC1")

    # 造一条「系统自动生成」的沉降记录（模拟"自动生成数据"功能产物）
    db = SessionLocal()
    try:
        db.add(
            ObservationRecord(
                point_id=jjc1["id"],
                series_type="settlement",
                component="value",
                obs_date=datetime.date(2026, 8, 1),
                value_mm=9999.0,
                status="normal",
                source="generated",
            )
        )
        db.commit()
    finally:
        db.close()

    # 导入同日真实台账
    up = client.post(
        f"/api/v1/projects/{pid}/import/upload",
        files={"file": ("真实台账.xlsx", _xlsx_bytes_with_date("2026-08-01"), _MIME)},
    )
    assert up.status_code == 200, up.text
    data = up.json()
    # JJC1 沉降与生成记录同日 → 可覆盖；其余 5 行为普通有效行
    assert len(data["override_rows"]) == 1
    assert data["override_rows"][0]["point_no"] == "JJC1"
    assert data["override_rows"][0]["existing_value_mm"] == 9999.0
    assert len(data["valid_rows"]) == 5
    assert len(data["conflict_rows"]) == 0

    # 确认导入（前端会把 override_rows 一并回传）
    conf = client.post(
        f"/api/v1/import/batches/{data['batch_id']}/confirm",
        json={"valid_rows": data["valid_rows"] + data["override_rows"]},
    )
    assert conf.status_code == 200
    body = conf.json()
    assert body["imported_rows"] == 6
    assert body["replaced_generated"] == 1

    # 生效记录 = 导入值(4906mm)、来源 excel；旧生成记录已 superseded 留痕
    db = SessionLocal()
    try:
        active = (
            db.query(ObservationRecord)
            .filter(
                ObservationRecord.point_id == jjc1["id"],
                ObservationRecord.series_type == "settlement",
                ObservationRecord.obs_date == datetime.date(2026, 8, 1),
                ObservationRecord.superseded_by_id.is_(None),
            )
            .all()
        )
        assert len(active) == 1
        assert active[0].source == "excel"
        assert active[0].value_mm == 4906.0
        old = (
            db.query(ObservationRecord)
            .filter(
                ObservationRecord.point_id == jjc1["id"],
                ObservationRecord.source == "generated",
            )
            .one()
        )
        assert old.superseded_by_id == active[0].id
    finally:
        db.close()


def test_import_still_protects_manual_records(client):
    """手工录入（manual）的记录仍受保护：导入撞上它只跳过，绝不覆盖。"""
    pid = _setup_project_with_points(client)
    pts = client.get(f"/api/v1/projects/{pid}/points").json()
    jjc1 = next(x for x in pts if x["point_no"] == "JJC1")
    client.post(
        f"/api/v1/points/{jjc1['id']}/records",
        json={"series_type": "settlement", "component": "value", "obs_date": "2026-08-01", "value_mm": 8888.0},
    )

    up = client.post(
        f"/api/v1/projects/{pid}/import/upload",
        files={"file": ("真实台账.xlsx", _xlsx_bytes_with_date("2026-08-01"), _MIME)},
    )
    data = up.json()
    assert len(data["override_rows"]) == 0          # 手工记录不进入"可覆盖"
    assert len(data["conflict_rows"]) == 1          # 而是冲突跳过
    assert data["conflict_rows"][0]["existing_value_mm"] == 8888.0


def _xyz_bytes() -> bytes:
    """「名字 + X/Y/Z」初始值清单（无日期）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Points_Manual"
    ws.append(["名字", "Y", "X", "Z"])
    ws.append(["XK1", 510182.646, 3413017.442, 4.144])
    ws.append(["XK2", 510183.327, 3413020.837, 4.113])
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def test_no_date_table_marks_rows_and_imports_after_date_supplied(client):
    """无日期表（初始值清单）：每行都带 needs_date 标记；用户指定日期后数据正常入库。

    回归：以前这类行的点号若在项目中不存在，会被归为"未知点"错误行且不带日期，
    前端补日期入口看不到它们 → 点被建出来、数据却被静默丢弃（imported_rows=0）。
    """
    p = client.post("/api/v1/projects", json={"code": "PRJ900", "name": "七号左边墩基坑"}).json()
    pid = p["id"]

    up = client.post(
        f"/api/v1/projects/{pid}/import/upload",
        files={"file": ("左边.xlsx", _xyz_bytes(), _MIME)},
    )
    assert up.status_code == 200, up.text
    data = up.json()
    assert data["sheets"][0]["needs_date"] is True
    assert len(data["valid_rows"]) == 0        # 无日期 → 不进有效行
    assert len(data["no_date_rows"]) == 0      # 点号不存在 → 先归为未知点
    assert len(data["error_rows"]) == 6        # 2 点 × (X/Y/Z)
    # 关键：错误行也带 needs_date，前端才能把日期补上
    assert all(r["needs_date"] for r in data["error_rows"])
    assert {r["point_no"] for r in data["error_rows"]} == {"XK1", "XK2"}

    # 前端补上初测日期 + 自动建点后一并回传
    rows = [{**r, "obs_date": "2026-07-28"} for r in data["error_rows"]]
    actions = [
        {"excel_no": "XK1", "auto_create": True, "infer_point_type": "both", "map_to": None},
        {"excel_no": "XK2", "auto_create": True, "infer_point_type": "both", "map_to": None},
    ]
    conf = client.post(
        f"/api/v1/import/batches/{data['batch_id']}/confirm",
        json={"valid_rows": rows, "point_actions": actions},
    )
    assert conf.status_code == 200, conf.text
    body = conf.json()
    assert body["imported_rows"] == 6
    assert body["skipped_no_date"] == 0
    assert body["points_auto_created"] == 2

    pts = client.get(f"/api/v1/projects/{pid}/points").json()
    xk1 = next(x for x in pts if x["point_no"] == "XK1")
    recs = client.get(f"/api/v1/points/{xk1['id']}/records").json()
    assert len(recs) == 3  # 沉降 Z + 位移 X + 位移 Y
    assert all(r["obs_date"] == "2026-07-28" for r in recs)


def test_confirm_without_date_counts_skipped_not_silent(client):
    """没给日期的行不会被静默丢弃：返回 skipped_no_date 计数（用户能看见）。"""
    pid = _setup_project_with_points(client)
    up = client.post(
        f"/api/v1/projects/{pid}/import/upload",
        files={"file": ("八号基坑（第十期）2026.08.01.xlsx", _xlsx_bytes_with_date("2026-08-01"), _MIME)},
    )
    data = up.json()
    rows = [{**r, "obs_date": None} for r in data["valid_rows"]]
    conf = client.post(f"/api/v1/import/batches/{data['batch_id']}/confirm", json={"valid_rows": rows})
    assert conf.status_code == 200
    assert conf.json()["imported_rows"] == 0
    assert conf.json()["skipped_no_date"] == 6
