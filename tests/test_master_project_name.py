# -*- coding: utf-8 -*-
"""导出的整卷日报里，工程名称必须是**当前项目**的名字（2026-10-08 用户报障）。

用户原话：「我项目名称是十号的基坑 我导出怎么是十一号的 ?」

**实测到的原因**（拿生产库里那份母版 + 真实项目跑一次导出，逐格扫"十一号"）：
那份母版是拿**十一号**那期的报表另存的（用户自己存的时候只改了文件名），
母版里有 **8 处固定文字**写着"十一号"：封面工程名称格、签字"（十一号）"、
两张数据表的"监测工程名称：十一号"与"工点名称：兴善公路 十一号"、以及两段
"监测结论及建议：本次监测工作为十一号围檩…"。

系统只替换它学到的**变量槽**（日期/期次/数据），这些固定文字原样留着 → 报表上全是别人的基坑名。

**原有的 `_sync_project_name` 为什么没拦住**：它只认 `数字号…(基坑|墩)` 这种**带后缀**的实体
（"5号右边墩基坑"），而这份母版自称的是光秃秃的**「十一号」**（后面直接跟空格）→ 匹配不上。

**现在**：先从名称标签格（监测工程名称/工点名称/工程名称…）读出母版"自称的旧名"，
再全簿整体替换成当前项目名（旧名后若紧跟"基坑/墩"连后缀一起吃掉，不会变成"十号基坑基坑"）。
"""
import openpyxl
import pytest

from tests.test_control_values import _download_master, _upload_master
from tests.test_master_report import _make_master, _seed_project

OLD = "十一号"
NEW = "十号基坑"


def _master_named(tmp_path, old: str = OLD, suffix: bool = False):
    """同构母版 + 母版"自称的基坑名"（照抄用户那份：封面/签字/两张表/结论段都写着旧名）。"""
    from pathlib import Path

    ent = old + ("基坑" if suffix else "")
    wb = openpyxl.load_workbook(_make_master())
    wb["封面"]["F11"] = ent
    wb["签字"]["A9"] = f"（{ent}）"
    for sn, tail_row in (("竖向位移", 31), ("r.立柱结构水平位移", 34)):
        ws = wb[sn]
        ws["A2"] = f"监测工程名称：{ent}"
        ws["G3"] = f"工点名称：兴善公路  {ent}"
        ws.cell(row=tail_row, column=1, value=f"监测结论及建议：本次监测工作为{ent}围檩位移监测工作。")
    out = Path(tmp_path) / "master_named.xlsx"
    wb.save(out)
    wb.close()
    return out


def _seed_named_project(client, name: str = NEW) -> int:
    pid = _seed_project(client)
    r = client.put(f"/api/v1/projects/{pid}", json={"name": name})
    assert r.status_code == 200, r.text
    return pid


def _texts(wb) -> dict[str, list[str]]:
    """全簿所有字符串单元格：{sheet: [文本…]}"""
    out: dict[str, list[str]] = {}
    for sn in wb.sheetnames:
        vals = []
        for row in wb[sn].iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.strip():
                    vals.append(v)
        out[sn] = vals
    return out


def _dump(path) -> dict[str, list[str]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    data = _texts(wb)
    wb.close()
    return data


# ---------- 主场景 ----------


def test_master_self_name_is_replaced_everywhere(client, tmp_path):
    """母版自称「十一号」→ 导出的整卷里不能再出现它，全部换成当前项目名。

    8 处都要覆盖（封面/签字/两张数据表的名称行/两段结论文字），
    少一处用户翻到那页就还是会问"怎么是十一号的"。
    """
    pid = _seed_named_project(client)
    tid = _upload_master(client, pid, _master_named(tmp_path))
    wb = _download_master(client, pid, tid)

    texts = _texts(wb)
    flat = "\n".join(t for v in texts.values() for t in v)
    assert OLD not in flat, f"报表里还残留旧基坑名：{[t for t in [x for v in texts.values() for x in v] if OLD in t]}"

    assert "十号基坑" in wb["封面"]["F11"].value
    assert "十号基坑" in wb["签字"]["A9"].value
    assert "监测工程名称：十号基坑" in wb["竖向位移"]["A2"].value
    assert "兴善公路  十号基坑" in wb["竖向位移"]["G3"].value
    assert "本次监测工作为十号基坑围檩" in str(wb["竖向位移"]["A31"].value)
    assert "本次监测工作为十号基坑围檩" in str(wb["r.立柱结构水平位移"]["A34"].value)
    wb.close()


def test_master_file_itself_is_not_touched(client, tmp_path):
    """只改导出的副本，母版原件一个字节都不许动（下次导出还得靠它）。"""
    path = _master_named(tmp_path)
    before = path.read_bytes()
    pid = _seed_named_project(client)
    tid = _upload_master(client, pid, path)
    wb = _download_master(client, pid, tid)
    wb.close()
    assert path.read_bytes() == before, "母版原件被改写了！"


def test_suffix_is_not_duplicated(client, tmp_path):
    """母版写「十一号基坑」+ 项目叫「十号基坑」→ 不能变成「十号基坑基坑」。"""
    pid = _seed_named_project(client)
    tid = _upload_master(client, pid, _master_named(tmp_path, suffix=True))
    wb = _download_master(client, pid, tid)
    flat = "\n".join(t for v in _texts(wb).values() for t in v)
    assert "十号基坑基坑" not in flat
    assert OLD not in flat
    wb.close()


def test_old_name_contained_in_new_is_left_alone(client, tmp_path):
    """旧名被新名包含（旧「十号」/ 新「十号基坑」）→ 一律不动，免得替成「十号基坑基坑」。"""
    pid = _seed_named_project(client)
    tid = _upload_master(client, pid, _master_named(tmp_path, old="十号"))
    wb = _download_master(client, pid, tid)
    flat = "\n".join(t for v in _texts(wb).values() for t in v)
    assert "十号基坑基坑" not in flat
    assert "监测工程名称：十号" in wb["竖向位移"]["A2"].value
    wb.close()


def test_data_and_point_numbers_are_untouched(client, tmp_path):
    """整簿替换不能误伤数据：点号、日期、累计值都还在。"""
    pid = _seed_named_project(client)
    tid = _upload_master(client, pid, _master_named(tmp_path))
    wb = _download_master(client, pid, tid)
    ws = wb["竖向位移"]
    assert ws.cell(row=7, column=1).value == "JJC1"
    assert ws.cell(row=7, column=2).value is not None      # 初始值
    assert ws.cell(row=7, column=6).value is not None      # 本次累计
    assert "2026-08-21" in str(ws.cell(row=3, column=1).value)
    wb.close()


# ---------- 安全兜底：认不出来就别乱改 ----------


def test_extraction_ignores_a_single_label_hit():
    """名称标签只出现一处 → 不认（宁可不换，也别拿一个来路不明的词到处替换）。

    `_make_master()` 那份母版只有竖向位移表的 A2 一处"监测工程名称：8号基坑"。
    """
    from backend.reports.master_report import _extract_master_project_name
    from tests.test_master_report import _make_master

    import openpyxl

    wb = openpyxl.load_workbook(_make_master())
    assert _extract_master_project_name(wb) is None
    wb.close()


def test_extraction_needs_two_hits():
    """源码级判据：单处命中不认（防误伤）。"""
    import inspect

    from backend.reports import master_report

    src = inspect.getsource(master_report._extract_master_project_name)
    assert "n >= 2" in src, "放宽成「认一处就替换」会把报表改乱"
    assert "_looks_like_project_name" in src


@pytest.mark.parametrize("bad", ["兴善公路", "基坑", "2026年8月20日", "（", "监测工程名称"])
def test_looks_like_project_name_rejects_junk(bad):
    from backend.reports.master_report import _looks_like_project_name

    assert not _looks_like_project_name(bad)


@pytest.mark.parametrize("good", ["十一号", "五号右边墩基坑", "10号基坑", "十四号基坑"])
def test_looks_like_project_name_accepts_real_names(good):
    from backend.reports.master_report import _looks_like_project_name

    assert _looks_like_project_name(good)
