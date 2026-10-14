# -*- coding: utf-8 -*-
"""导出报表**不显示预警**（2026-09-30 用户要求：「导出的 Excel 不要显示预警」）。

两处来源，都必须在导出文件里消失：
1. 数据表的「预警等级」列（K 列）—— 原实现按引擎结果写"报警/预警"；
2. 母版自带的「预警统计」sheet —— 那是**上个基坑**的点位与预警等级
   （实测："管线竖向位移监测 GGC02-03 黄色预警"、"地表沉降监测 DBC01-01"），
   导出时原样保留，于是你的报表里出现别人的点号、别人的预警。

⚠ 口径不是"统一写正常"：真有报警却报正常属于伪造监测数据，任何情况下都不做 ——
   这一列只是**不填**（母版里它本来也是空的）。
"""
from __future__ import annotations

import json

import pytest
from io import BytesIO
from urllib.parse import unquote
from pathlib import Path

from openpyxl import Workbook, load_workbook

from backend.core.config import settings
from backend.core.db import SessionLocal
from backend.models.export_template import ExportTemplate
from backend.reports import export_options

SETTLE = "竖向位移"
STATS = "预警统计"
TAIL_TEXT = "变化量正负号说明："
# 母版自带的两行"别人的"预警数据（照抄真实母版）
STALE_TITLE = "本期预警点统计表"
STALE_HEADER = ["测  项", "测 点", "本期预警"]
STALE_ROWS = [
    ["管线竖向位移监测", "GGC02-03", "黄色预警"],
    ["地表沉降监测", "DBC05-01", "橙色预警"],
]


def _layout(with_meta: bool = True) -> dict:
    """两种形态都要能"不显示预警"：

    · meta 非空 → 后端走**母版整卷路径**（用户现在用的）；
    · meta 为空 → 走**旧式模板路径**（`template_excel`），换了模板也不能漏。
    """
    cols = {"point_col": 1, "v0_m": 2, "prev_m": 3, "current_m": 4, "prev_cum_mm": 5,
            "cum_mm": 6, "delta_mm": 7, "rate": 8, "warning_level": 11,
            "control_cum": 9, "control_rate": 10, "data_start_row": 7}
    meta = {"sign_date_cell": "签字!E5"} if with_meta else {}
    return {"meta": meta, "sheets": [{"sheet_name": SETTLE, "kind": "settlement",
                                      "layout": dict(cols), "mode": "per_row"}]}


def _build_master(path: Path, template_points: int = 8) -> dict:
    """造一份带「预警统计」sheet 的样表：K 列写着母版自带的预警、统计表里是别的基坑的数据。"""
    wb = Workbook()
    ws = wb.active
    ws.title = SETTLE
    ws["A5"] = "监测点号"
    ws["K5"] = "预警\n等级"
    for i in range(1, template_points + 1):
        r = 6 + i
        ws.cell(row=r, column=1, value=f"JJC{i}")
        ws.cell(row=r, column=9, value=30)
        ws.cell(row=r, column=10, value=2)
        ws.cell(row=r, column=11, value="预警")     # 母版自带的旧预警，必须被清掉
    tail = 7 + template_points + 12
    ws.cell(row=tail, column=1, value=TAIL_TEXT)

    st = wb.create_sheet(STATS)
    st["B2"] = "最大变化量统计表"
    st.merge_cells("B2:L2")
    st["B3"], st["C3"] = "测  项", "本期最大变化量"
    st["B5"], st["C5"] = "地表沉降监测", "DBC01-01"      # 数据行
    st["B27"] = STALE_TITLE
    st.merge_cells("B27:L27")
    for c, v in enumerate(STALE_HEADER, start=2):
        st.cell(row=28, column=c, value=v)
    st.merge_cells("B28:B29")
    for i, row in enumerate(STALE_ROWS, start=30):
        for c, v in enumerate(row, start=2):
            st.cell(row=i, column=c, value=v)
    wb.save(path)


def _master_file(name: str) -> Path:
    path = Path(settings.templates_dir) / f"{name}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    _build_master(path)
    return path


def _mk_project(client, n_points: int = 8) -> int:
    """建项目 + n 个测点，第 1 个点造一条**真会报警**的数据（单期 10mm，控制速率 2mm/d）。"""
    p = client.post("/api/v1/projects", json={"code": "PRJ-WARN", "name": "不显示预警测试"}).json()
    for i in range(1, n_points + 1):
        pt = client.post(f"/api/v1/projects/{p['id']}/points",
                         json={"point_no": f"JJC{i}", "point_type": "settlement"}).json()
        for d, v in (("2026-09-01", 3500.0), ("2026-09-02", 3500.0 + (10.0 if i == 1 else 0.0))):
            client.post(f"/api/v1/points/{pt['id']}/records",
                        json={"series_type": "settlement", "component": "value",
                              "obs_date": d, "value_mm": v})
    client.put(f"/api/v1/projects/{p['id']}/warning-rules", json={
        "rules": [
            {"indicator": "settlement", "level": "control",
             "threshold_cumulative_mm": 30, "threshold_rate_mm_per_day": 2},
            {"indicator": "settlement", "level": "alarm",
             "threshold_cumulative_mm": 40, "threshold_rate_mm_per_day": 4},
        ],
        "source_note": "测试规范",
    })
    return p["id"]


def _mk_template(project_id: int, path: Path, layout: dict) -> int:
    with SessionLocal() as db:
        tpl = ExportTemplate(project_id=project_id, name="预警母版", template_type="excel",
                             is_builtin=False, file_path=str(path), layout_json=layout)
        db.add(tpl)
        db.commit()
        db.refresh(tpl)
        return tpl.id


def _export(client, project_id: int, template_id: int):
    r = client.get(f"/api/v1/projects/{project_id}/reports/excel",
                   params={"template_id": template_id, "obs_date": "2026-09-02"})
    assert r.status_code == 200, r.text
    summary = json.loads(unquote(r.headers.get("x-export-summary", "{}") or "{}"))
    return load_workbook(BytesIO(r.content)), summary


# ────────────────────── 1. 数据表的预警等级列 ──────────────────────


@pytest.mark.parametrize("with_meta", [True, False], ids=["母版路径", "旧模板路径"])
def test_warning_level_column_is_empty_even_when_alarm(client, with_meta):
    """真有报警的测点，导出的 K 列也必须空着（不是写"正常"，是不填）。"""
    pid = _mk_project(client)
    tid = _mk_template(pid, _master_file("m_warn"), _layout(with_meta))

    wb, _ = _export(client, pid, tid)
    ws = wb[SETTLE]
    col = [ws.cell(row=r, column=11).value for r in range(7, 15)]
    assert col == [None] * 8, f"K 列还写着预警等级：{col}"

    # 数据本身必须是**真的报警过**，否则这条测试等于没测
    status = client.get(f"/api/v1/projects/{pid}/warning-status",
                        params={"obs_date": "2026-09-02"}).json()
    assert any(x["warning_level"] == "alarm" for x in status), (
        f"造的样例根本没触发报警，这条测试测不到东西：{status[:2]}"
    )


@pytest.mark.parametrize("with_meta", [True, False], ids=["母版路径", "旧模板路径"])
def test_other_columns_are_unaffected(client, with_meta):
    """只是预警列不填，数值列照旧全填（别把整表清空了）。"""
    pid = _mk_project(client)
    tid = _mk_template(pid, _master_file("m_warn"), _layout(with_meta))

    wb, _ = _export(client, pid, tid)
    ws = wb[SETTLE]
    assert ws.cell(row=7, column=1).value == "JJC1"
    assert ws.cell(row=7, column=2).value is not None          # B 初始值
    assert ws.cell(row=7, column=4).value is not None          # D 本次测值
    assert ws.cell(row=7, column=6).value == 10.0              # F 本次累计 10mm
    assert ws.cell(row=7, column=7).value == 10.0              # G 本次变化


# ────────────────────── 2. 母版自带的「预警统计」表 ──────────────────────


@pytest.mark.parametrize("with_meta", [True, False], ids=["母版路径", "旧模板路径"])
def test_stale_stats_table_is_cleared(client, with_meta):
    """母版统计表里别的基坑的点位与预警等级必须被清掉，标题与表头保留。"""
    pid = _mk_project(client)
    tid = _mk_template(pid, _master_file("m_warn"), _layout(with_meta))

    wb, summary = _export(client, pid, tid)
    ws = wb[STATS]
    assert summary["cleared_stats"].get(STATS), "导出摘要里没说清了多少格"
    for r in (5,) + tuple(range(30, 32)):
        for c in range(2, 5):
            assert ws.cell(row=r, column=c).value is None, (
                f"统计表里还留着母版的旧数据 {ws.cell(row=r, column=c).coordinate}"
            )
    # 标题 / 表头 / 合并区都在 —— 表结构不能被破坏
    assert "最大变化量统计表" in str(ws["B2"].value or "")
    assert STALE_TITLE in str(ws["B27"].value or "")
    assert ws["B28"].value == "测  项" and ws["C28"].value == "测 点"
    assert "B27:L27" in {str(m) for m in ws.merged_cells.ranges}
    assert "B28:B29" in {str(m) for m in ws.merged_cells.ranges}


@pytest.mark.parametrize("with_meta", [True, False], ids=["母版路径", "旧模板路径"])
def test_whole_workbook_has_no_warning_words(client, with_meta):
    """整份导出文件里不该再出现"黄色预警/橙色预警"这类**母版残留**的等级字样。"""
    pid = _mk_project(client)
    tid = _mk_template(pid, _master_file("m_warn"), _layout(with_meta))

    wb, _ = _export(client, pid, tid)
    bad = []
    for sn in wb.sheetnames:
        for row in wb[sn].iter_rows():
            for c in row:
                v = str(c.value or "")
                if "黄色预警" in v or "橙色预警" in v or "红色预警" in v:
                    bad.append(f"{sn}!{c.coordinate}={v[:24]}")
    assert not bad, f"导出文件里还有预警等级字样：{bad[:5]}"


@pytest.mark.parametrize("with_meta", [True, False], ids=["母版路径", "旧模板路径"])
def test_other_sheets_are_untouched(client, with_meta):
    """没有"统计表"标题的 Sheet（封面/概况/分析…）一个格都不能动。"""
    pid = _mk_project(client)
    tid = _mk_template(pid, _master_file("m_warn"), _layout(with_meta))

    wb, summary = _export(client, pid, tid)
    assert SETTLE not in summary["cleared_stats"], "数据表被统计清理误伤了"
    assert list(summary["cleared_stats"]) == [STATS]


# ────────────────────── 3. 口径开关（源码守卫） ──────────────────────


def test_switch_restores_warning_column(client, monkeypatch):
    """开关打开时 K 列照旧写预警等级 —— 证明这是"开关关掉了"，不是"功能删了"。"""
    monkeypatch.setattr(export_options, "EXPORT_SHOW_WARNING_LEVEL", True)
    from backend.reports import master_report
    monkeypatch.setattr(master_report, "EXPORT_SHOW_WARNING_LEVEL", True)

    pid = _mk_project(client)
    tid = _mk_template(pid, _master_file("m_warn"), _layout())
    wb, _ = _export(client, pid, tid)
    assert wb[SETTLE].cell(row=7, column=11).value == "报警", "开关打开后 K 列应该写回预警等级"


def test_all_report_writers_respect_the_switch():
    """源码守卫：三条导出路径（母版/旧模板/标准版式）都必须接同一个开关。

    漏掉任何一条，换个模板导出又会把预警带出来。
    """
    root = Path(__file__).resolve().parent.parent / "backend" / "reports"
    for name in ("master_report.py", "template_excel.py", "excel_report.py"):
        src = (root / name).read_text(encoding="utf-8")
        assert "EXPORT_SHOW_WARNING_LEVEL" in src, f"{name} 没有接口径开关"
        assert "_warning_label(" in src and "EXPORT_SHOW_WARNING_LEVEL" in src, (
            f"{name} 的预警列没有受开关控制"
        )
    # 母版副本里的统计表旧数据：两条导出路径都要清
    for name in ("master_report.py", "template_excel.py"):
        src = (root / name).read_text(encoding="utf-8")
        assert "clear_stale_stat_tables" in src, f"{name} 没清母版统计表的旧数据"
    opts = (root / "export_options.py").read_text(encoding="utf-8")
    assert "EXPORT_SHOW_WARNING_LEVEL = False" in opts, "开关默认值被改成了 True"
