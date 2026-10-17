# -*- coding: utf-8 -*-
"""测试：主界面「集中管理」—— 全站默认值 + 项目批量操作。

用户 2026-10-11：
    「能不能把一些操作集中主界面上？每个项目都要调单独设置有点繁琐」

两件事各写成一组测试：
- **默认值**（`/settings`）：新建项目自动带上规则套 / 母版 —— 建完就能用，不用再配一遍；
- **批量操作**（`/bulk/projects/*`）：勾几个项目一次做完归档、换规则、换母版、导日报。

重点守的几条（都是"错了会静默出事"的地方）：
1. 批量归档**只改 is_active**，观测数据一条都不能动；
2. 批量套母版是"加上去"，**不能把项目原有的母版踢掉**；
3. 批量套规则是"改绑"，**一个项目只能有一套**，且返回里要能看出原来用的是哪套；
4. 批量导出**部分失败要报出来**（写进 ZIP 里的说明.txt），绝不静默少给一个文件；
5. 改全站默认值**不影响已有项目**。
"""
import io
import zipfile

from openpyxl import Workbook

from tests.conftest import put_control_rules

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ─────────────────────────── 脚手架 ───────────────────────────


def _tpl_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "竖向位移"
    ws["A5"] = "监测\n点号"
    ws["B5"] = "初始值（m)"
    ws["C5"] = "上次测值（m）"
    ws["D5"] = "本次测值（m）"
    ws["F5"] = "本次累计（mm)"
    ws["A7"] = "JJC1"
    ws["B7"] = 3.1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _mk_project(client, code: str, name: str) -> int:
    return client.post("/api/v1/projects", json={"code": code, "name": name}).json()["id"]


def _mk_project_with_data(client, code: str, name: str, dates=("2026-07-30", "2026-07-31")) -> int:
    pid = _mk_project(client, code, name)
    put_control_rules(client, pid)
    pt = client.post(
        f"/api/v1/projects/{pid}/points", json={"point_no": "JJC1", "point_type": "settlement"}
    ).json()
    for d in dates:
        client.post(
            f"/api/v1/points/{pt['id']}/records",
            json={"series_type": "settlement", "component": "value", "obs_date": d, "value_mm": 4907.0},
        )
    return pid


def _mk_rule_set(client, name: str, project_ids=None, is_global=False, cum=30.0) -> int:
    body = {
        "name": name,
        "is_global": is_global,
        "project_ids": project_ids or [],
        "rules": [
            {"indicator": "settlement", "level": "control", "threshold_cumulative_mm": cum,
             "threshold_rate_mm_per_day": 10.0, "is_active": True},
            {"indicator": "settlement", "level": "alarm", "threshold_cumulative_mm": 40.0,
             "threshold_rate_mm_per_day": 12.0, "is_active": True},
        ],
    }
    r = client.post("/api/v1/warning-rule-sets", json=body)
    assert r.status_code in (200, 201), r.text
    return r.json()["id"]


def _mk_template(client, name: str, scope: list[int] | None, is_global=False) -> int:
    data = {"name": name, "project_ids": ",".join(str(x) for x in (scope or []))}
    if is_global:
        data["is_global"] = "true"
    r = client.post(
        "/api/v1/export-templates/upload",
        data=data,
        files={"file": ("t.xlsx", _tpl_bytes(), XLSX_MIME)},
    )
    assert r.status_code in (200, 201), r.text
    return r.json()["id"]


def _record_count(client, pid: int) -> int:
    return client.get(f"/api/v1/projects/{pid}/data-count").json()["records"]


def _rules_of_project(client, pid: int) -> list[dict]:
    return client.get(f"/api/v1/projects/{pid}/warning-rules").json()


def _tpl_names_of_project(client, pid: int) -> set[str]:
    return {t["name"] for t in client.get(f"/api/v1/projects/{pid}/export-templates").json()}


# ─────────────────────────── 一、全站默认值 ───────────────────────────


def test_settings_defaults_then_update(client) -> None:
    """缺省是一份"什么都不绑"的设置；改了之后读回来还是它。"""
    cfg = client.get("/api/v1/settings").json()
    assert cfg == {
        "default_rule_set_id": None,
        "default_template_id": None,
        "default_weather": "晴",
        # 「文件检查」的两个键（2026-10-12 加的）—— 没配过就是空
        "audit_folder": None,
        "audit_last_run": None,
    }

    r = client.put("/api/v1/settings", json={"default_weather": "小雨"})
    assert r.status_code == 200, r.text
    assert r.json()["default_weather"] == "小雨"
    assert client.get("/api/v1/settings").json()["default_weather"] == "小雨"


def test_settings_reject_missing_references(client) -> None:
    """引用了不存在的规则套 / 母版要当场报错，不能默默存进去（建项目时才发现就晚了）。"""
    assert client.put("/api/v1/settings", json={"default_rule_set_id": 9999}).status_code == 404
    assert client.put("/api/v1/settings", json={"default_template_id": 9999}).status_code == 404


def test_settings_reject_word_template_as_default(client) -> None:
    """Word 日报走的是另一条固定版式，不能拿来当"默认导出母版"。"""
    word = next(
        t for t in client.get("/api/v1/export-templates").json() if t["template_type"] == "word"
    )
    r = client.put("/api/v1/settings", json={"default_template_id": word["id"]})
    assert r.status_code == 400, r.text
    assert "Excel" in r.json()["detail"]


def test_new_project_inherits_default_rule_set(client) -> None:
    """新建项目自动绑上默认规则套 —— 建完就能判预警。"""
    sid = _mk_rule_set(client, "全站默认规则", cum=28.0)
    client.put("/api/v1/settings", json={"default_rule_set_id": sid})

    pid = _mk_project(client, "PRJ-A", "甲基坑")
    rules = _rules_of_project(client, pid)
    assert any(r["threshold_cumulative_mm"] == 28.0 for r in rules), (
        "新建项目没有套上默认规则套：" + repr(rules)
    )
    # 规则套那边也要记上这个项目（双向一致）
    s = next(x for x in client.get("/api/v1/warning-rule-sets").json() if x["id"] == sid)
    assert pid in s["project_ids"]


def test_new_project_inherits_default_template(client) -> None:
    """新建项目自动加进默认母版的适用范围 —— 建完就能直接导日报。"""
    anchor = _mk_project(client, "PRJ-ANCHOR", "锚点项目")
    tid = _mk_template(client, "全站默认母版", [anchor])
    client.put("/api/v1/settings", json={"default_template_id": tid})

    pid = _mk_project(client, "PRJ-B", "乙基坑")
    assert "全站默认母版" in _tpl_names_of_project(client, pid)


def test_new_project_under_global_template_needs_no_binding(client) -> None:
    """默认母版若是「通用」，新项目不用逐项绑定也照样看得见（通用本来就覆盖所有项目）。"""
    tid = _mk_template(client, "通用母版", [], is_global=True)
    client.put("/api/v1/settings", json={"default_template_id": tid})

    pid = _mk_project(client, "PRJ-C", "丙基坑")
    assert "通用母版" in _tpl_names_of_project(client, pid)
    t = next(x for x in client.get("/api/v1/export-templates").json() if x["id"] == tid)
    assert pid not in t["project_ids"], "通用母版不该再逐个项目绑一遍"


def test_settings_change_does_not_touch_existing_projects(client) -> None:
    """改全站默认值**只影响以后新建的项目** —— 已有项目一个都不能动。"""
    old = _mk_project(client, "PRJ-OLD", "老项目")
    sid = _mk_rule_set(client, "新默认规则", cum=25.0)
    client.put("/api/v1/settings", json={"default_rule_set_id": sid})

    # 老项目没有任何规则（没配过）
    assert _rules_of_project(client, old) == []
    s = next(x for x in client.get("/api/v1/warning-rule-sets").json() if x["id"] == sid)
    assert old not in s["project_ids"], "改默认值把已有项目也改了 —— 这属于动历史配置"


def test_broken_default_does_not_block_project_creation(client) -> None:
    """默认值指向已被删掉的规则套时，建项目**照常成功**（正事不能被配角带崩）。"""
    sid = _mk_rule_set(client, "待删规则", cum=25.0)
    client.put("/api/v1/settings", json={"default_rule_set_id": sid})
    assert client.delete(f"/api/v1/warning-rule-sets/{sid}").status_code == 200

    r = client.post("/api/v1/projects", json={"code": "PRJ-D", "name": "丁基坑"})
    assert r.status_code == 201, r.text


# ─────────────────────────── 二、批量操作 ───────────────────────────


def test_bulk_archive_keeps_all_records(client) -> None:
    """批量归档只改"在用/归档"，**观测数据一条都不能动**。"""
    a = _mk_project_with_data(client, "PRJ-1", "一号")
    b = _mk_project_with_data(client, "PRJ-2", "二号")
    before = {a: _record_count(client, a), b: _record_count(client, b)}

    r = client.post("/api/v1/bulk/projects/archive", json={"project_ids": [a, b], "active": False})
    assert r.status_code == 200, r.text
    assert r.json()["changed"] == 2

    for pid in (a, b):
        assert client.get(f"/api/v1/projects/{pid}").json()["is_active"] is False
        assert _record_count(client, pid) == before[pid], "归档动了观测数据！"

    # 再归档一次：已经是的要如实说"没动"，而不是又算一次 changed
    r2 = client.post("/api/v1/bulk/projects/archive", json={"project_ids": [a, b], "active": False})
    assert r2.json()["changed"] == 0
    assert r2.json()["already"] == 2


def test_bulk_archive_can_restore(client) -> None:
    a = _mk_project(client, "PRJ-3", "三号")
    client.post("/api/v1/bulk/projects/archive", json={"project_ids": [a], "active": False})
    client.post("/api/v1/bulk/projects/archive", json={"project_ids": [a], "active": True})
    assert client.get(f"/api/v1/projects/{a}").json()["is_active"] is True


def test_bulk_rule_set_rebinds_multiple_projects(client) -> None:
    """一次把几个项目改绑到同一套规则，并说清它们原来用的是哪套。"""
    old_set = _mk_rule_set(client, "旧规则", cum=30.0)
    a = _mk_project_with_data(client, "PRJ-4", "四号")
    b = _mk_project(client, "PRJ-5", "五号")
    client.post("/api/v1/bulk/projects/rule-set", json={"project_ids": [a], "set_id": old_set})

    new_set = _mk_rule_set(client, "新规则", cum=22.0)
    r = client.post(
        "/api/v1/bulk/projects/rule-set", json={"project_ids": [a, b], "set_id": new_set}
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["set_name"] == "新规则"
    # 改绑前 a 用的是"旧规则"、b 一套都没有 —— 界面要靠这个如实告诉用户
    assert body["before"][str(a)] == "旧规则"
    assert body["before"][str(b)] == "（没有，不判预警）"

    for pid in (a, b):
        assert any(x["threshold_cumulative_mm"] == 22.0 for x in _rules_of_project(client, pid))


def test_bulk_rule_set_can_unbind(client) -> None:
    """set_id=null → 解绑，改回吃通用规则。"""
    sid = _mk_rule_set(client, "临时规则", cum=30.0)
    a = _mk_project(client, "PRJ-6", "六号")
    client.post("/api/v1/bulk/projects/rule-set", json={"project_ids": [a], "set_id": sid})
    assert _rules_of_project(client, a) != []

    client.post("/api/v1/bulk/projects/rule-set", json={"project_ids": [a], "set_id": None})
    assert _rules_of_project(client, a) == []


def test_bulk_template_adds_without_removing_existing(client) -> None:
    """批量套母版是「加上去」，两头都不能丢：

    ① 项目原来那份母版必须还在（一个项目可以有好几份可用母版）；
    ② **目标母版原来绑着的其它项目也不能被踢掉** —— 只勾了 B 就把 A 挤出去，
       是这条接口最容易写错的地方（把"追加"写成了"覆盖"）。
    """
    a = _mk_project(client, "PRJ-7", "七号")
    b = _mk_project(client, "PRJ-8", "八号")
    c = _mk_project(client, "PRJ-8C", "八号丙")
    _mk_template(client, "原有母版", [a])
    shared = _mk_template(client, "共用母版", [a, c])  # 注意：a 一开始就在这份母版里

    r = client.post("/api/v1/bulk/projects/template", json={"project_ids": [b], "template_id": shared})
    assert r.status_code == 200, r.text

    # ① 项目 a 自己那份「原有母版」还在，并且新套的「共用母版」也来了
    assert {"原有母版", "共用母版"} <= _tpl_names_of_project(client, a)
    # ② a、c 原来就在「共用母版」的适用范围里，不能被 b 的到来挤掉
    t = next(x for x in client.get("/api/v1/export-templates").json() if x["id"] == shared)
    assert {a, c, b} <= set(t["project_ids"]), f"批量套母版把原有的绑定覆盖掉了：{t['project_ids']}"


def test_bulk_template_rejects_builtin(client) -> None:
    """内置版式本来就所有项目能用，硬塞返回 400（并说清原因）。"""
    a = _mk_project(client, "PRJ-9", "九号")
    builtin = next(
        t for t in client.get("/api/v1/export-templates").json() if t["is_builtin"]
    )
    r = client.post(
        "/api/v1/bulk/projects/template", json={"project_ids": [a], "template_id": builtin["id"]}
    )
    assert r.status_code == 400, r.text


def test_bulk_export_one_file_per_project(client) -> None:
    """每个项目导最新一期，一个项目一个文件。"""
    a = _mk_project_with_data(client, "PRJ-10", "十号", dates=("2026-07-30", "2026-07-31"))
    b = _mk_project_with_data(client, "PRJ-11", "十一号", dates=("2026-08-01",))

    r = client.post("/api/v1/bulk/projects/export", json={"project_ids": [a, b]})
    assert r.status_code == 200, r.text
    zf = zipfile.ZipFile(io.BytesIO(r.content))
    xlsx = [n for n in zf.namelist() if n.endswith(".xlsx")]
    assert len(xlsx) == 2, zf.namelist()
    # 文件名里要能看出是哪个项目、哪一期
    assert any("PRJ-10" in n and "2026-07-31" in n for n in xlsx), xlsx
    assert any("PRJ-11" in n and "2026-08-01" in n for n in xlsx), xlsx
    assert zf.read(xlsx[0])[:2] == b"PK"


def test_bulk_export_uses_each_project_own_master(client) -> None:
    """批量导出**默认用各项目自己那份母版** —— 出来的要和平时单项目导的一样。

    这条很要紧：现场每个基坑的日报都是拿自己那份母版套的（表头、封面、概况页都不同）。
    批量导出若悄悄换成系统标准版式，用户拿到手才发现"这批跟我平时发的不一样"。
    《说明.txt》要把"每个文件用的哪份母版"写出来，不许含糊。
    """
    a = _mk_project_with_data(client, "PRJ-15", "十五号")
    bare = _mk_project_with_data(client, "PRJ-16", "十六号（没配母版）")
    _mk_template(client, "十五号专用母版", [a])

    r = client.post("/api/v1/bulk/projects/export", json={"project_ids": [a, bare]})
    assert r.status_code == 200, r.text
    note = zipfile.ZipFile(io.BytesIO(r.content)).read("说明.txt").decode("utf-8")
    assert "十五号专用母版" in note, note
    assert "系统标准版式" in note, f"没母版的项目要走标准版式，并且如实写出来：{note}"


def test_bulk_export_honours_explicit_template(client) -> None:
    """显式指定母版时，就统一用那一份（覆盖"各项目自己的"）。"""
    a = _mk_project_with_data(client, "PRJ-17", "十七号")
    _mk_template(client, "十七号自己的母版", [a])
    shared = _mk_template(client, "统一母版", [a])

    r = client.post(
        "/api/v1/bulk/projects/export", json={"project_ids": [a], "template_id": shared}
    )
    assert r.status_code == 200, r.text
    note = zipfile.ZipFile(io.BytesIO(r.content)).read("说明.txt").decode("utf-8")
    assert f"指定母版 #{shared}" in note, note
    assert "十七号自己的母版" not in note


def test_bulk_export_names_skipped_projects(client) -> None:
    """没数据的项目**跳过但点名** —— 用户拿到 ZIP 才发现少一个才是最坑的。"""
    a = _mk_project_with_data(client, "PRJ-12", "十二号")
    empty = _mk_project(client, "PRJ-13", "十三号（没数据）")

    r = client.post("/api/v1/bulk/projects/export", json={"project_ids": [a, empty]})
    assert r.status_code == 200, r.text
    zf = zipfile.ZipFile(io.BytesIO(r.content))
    names = zf.namelist()
    assert len([n for n in names if n.endswith(".xlsx")]) == 1
    assert "说明.txt" in names, names
    note = zf.read("说明.txt").decode("utf-8")
    assert "十三号" in note


def test_bulk_export_all_empty_is_422(client) -> None:
    """全都没数据 → 422 说清原因，而不是给一个空壳 ZIP。"""
    a = _mk_project(client, "PRJ-14", "十四号（空）")
    r = client.post("/api/v1/bulk/projects/export", json={"project_ids": [a]})
    assert r.status_code == 422, r.text
    assert "没" in r.json()["detail"] or "失败" in r.json()["detail"]


def test_bulk_endpoints_validate_input(client) -> None:
    """空选择 → 422；项目不存在 → 404（不要"部分静默跳过"）。"""
    assert client.post("/api/v1/bulk/projects/archive", json={"project_ids": [], "active": False}).status_code == 422
    assert client.post("/api/v1/bulk/projects/archive", json={"project_ids": [9999], "active": False}).status_code == 404
    assert (
        client.post("/api/v1/bulk/projects/rule-set", json={"project_ids": [], "set_id": None}).status_code
        == 422
    )


def test_bulk_features_are_advertised(client) -> None:
    """能力标记要在 —— 前后端版本对不上时界面才能提醒"重启后端"。"""
    feats = client.get("/api/health").json()["features"]
    assert "app_settings" in feats
    assert "bulk_project_ops" in feats
