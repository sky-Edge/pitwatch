# -*- coding: utf-8 -*-
"""测试：概况/分析/成果提交三页的**文字按本期数据自动生成**。

背景（用户 2026-10-08 实测取证）：五份母版的这三页**逐字完全相同**，内容却是别的工程的 ——

    分析!C3  = "(1)地表沉降监测：…DBC05-01，DBC06-05…"     ← 用户根本没这类测点
    分析!C14 = "(12)测斜QCX1监测：…"
    成果提交!B3 = "本周通过对**屿前站**的监测数据和巡视信息综合分析"
    成果提交!B9~B20 = 12 类监测项清单（用户只有竖向位移 + 水平位移两类）

即：这三页**每期都在往外发别人的东西**。用户选了"系统自动按本期数据生成"。

两条硬口径（本文件重点守）：
1. **只写数据能支持的话**：极值、极值点号、达到预警/报警的点数与点号、类别、点次、期次。
   现场巡视**是否发现异常**是人的判断（系统没有巡视数据）→ 绝不写"未发现异常"。
2. **多余格子必须清空**（母版 12 段、本项目 2 类 → 剩下 10 格清空并隐藏行）。
   只填不改 = 别人的内容继续留在报表上。
"""
from __future__ import annotations

import json
from datetime import date
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote

from openpyxl import Workbook, load_workbook

from backend.core.config import settings
from backend.core.db import SessionLocal
from backend.models.export_template import ExportTemplate
from tests.conftest import put_control_rules
from backend.reports.text_blocks import (
    PURPOSE_ANALYSIS,
    PURPOSE_CONCLUSION,
    PURPOSE_DELIVERABLES,
    PURPOSE_OVERVIEW,
    PURPOSE_PROGRESS,
    PURPOSE_WORKTABLE,
    GenContext,
    generate,
    write_block,
)

SETTLE = "竖向位移"
OVERVIEW_SHEET = "分析"
SUBMIT_SHEET = "成果提交"
SURVEY_SHEET = "概况"

PROJECT_NAME = "十一号基坑"
FOREIGN = ["DBC05-01", "QCX1", "屿前站", "砼支撑轴力", "水中拉伸钢板桩"]


# ─────────────────────────── 纯生成（不碰 Excel） ───────────────────────────


def _row(point_no, cum, level=None, kind="settlement", component="value"):
    return {
        "point_no": point_no,
        "series_type": kind,
        "component": component,
        "cum_mm": cum,
        "warning_level": level,
    }


def _ctx(rows, kinds=("settlement",), is_initial=False):
    return GenContext(
        project_name=PROJECT_NAME,
        obs_date=date(2026, 10, 8),
        period_text="第3期",
        is_initial=is_initial,
        rows=list(rows),
        kinds=list(kinds),
    )


def test_analysis_paragraph_uses_real_extremes_and_points():
    rows = [_row("JJC1", -5.2), _row("JJC2", 3.4), _row("JJC3", 1.0)]
    text = generate(PURPOSE_ANALYSIS, _ctx(rows))[0]
    assert text.startswith("(1)竖向位移监测：")
    assert "[-5.2,3.4]" in text, text
    # 极值点号：先最小值点、后最大值点（照母版句式）
    assert "点号分别为：JJC1，JJC2" in text, text
    assert text.rstrip().endswith("数据未发现异常情况。")


def test_analysis_reports_warning_points_by_level():
    rows = [
        _row("JJC1", 1.0),
        _row("JJC2", 31.0, "control"),
        _row("JJC3", 41.0, "alarm"),
    ]
    text = generate(PURPOSE_ANALYSIS, _ctx(rows))[0]
    assert "达到预警指标的测点共 2 个" in text, text
    assert "达到报警指标的有 1 个，点号为：JJC3" in text, text
    assert "达到预警指标的有 1 个，点号为：JJC2" in text, text
    assert "数据未发现异常情况" not in text


def test_single_level_warning_is_not_repeated():
    """只有一种等级时别写成"共 N 个:…有 N 个"（同义反复，报表上很啰嗦）。"""
    text = generate(PURPOSE_ANALYSIS, _ctx([_row("JJC2", 31.0, "control")]))[0]
    assert text.rstrip().endswith("达到预警指标的测点共 1 个，点号为：JJC2。"), text
    assert ":" not in text.split("点号分别为")[-1].replace("：", ""), text


def test_analysis_says_so_when_no_reading():
    """本期一条有效读数都没有 → 如实说"没有读数"，**不编**一个正常结论。"""
    text = generate(PURPOSE_ANALYSIS, _ctx([], kinds=["settlement"]))[0]
    assert "未能给出分析结论" in text
    assert "数据未发现异常" not in text


def test_displacement_points_carry_component_suffix():
    """水平位移每点 X/Y 两条，不带后缀分不清是哪条出的极值。"""
    rows = [
        _row("JJC1", -1.0, kind="displacement", component="X"),
        _row("JJC1", -9.9, kind="displacement", component="Y"),
    ]
    text = generate(PURPOSE_ANALYSIS, _ctx(rows, kinds=["displacement"]))[0]
    assert "JJC1-Y" in text and "JJC1-X" in text, text
    assert "立柱结构水平位移监测" in text


def test_overview_never_claims_no_abnormality():
    """巡视那句**不能**写"未发现异常" —— 系统没有巡视数据，那是人的判断。"""
    text = "\n".join(generate(PURPOSE_OVERVIEW, _ctx([_row("JJC1", 0.0)])))
    assert "未发现异常" not in text
    assert "现场巡视 1 次" in text
    assert PROJECT_NAME in text


def test_conclusion_switches_on_warning():
    calm = generate(PURPOSE_CONCLUSION, _ctx([_row("JJC1", 1.0)]))[0]
    assert "稳定状态" in calm and PROJECT_NAME in calm and "屿前站" not in calm

    busy = generate(PURPOSE_CONCLUSION, _ctx([_row("JJC1", 31.0, "control")]))[0]
    assert "1 个测点达到预警指标" in busy and "JJC1" in busy
    assert "未发现新增的异常情况" not in busy


def test_deliverables_only_lists_kinds_we_actually_produce():
    texts = generate(PURPOSE_DELIVERABLES, _ctx([_row("JJC1", 1.0)], kinds=["settlement", "displacement"]))
    assert texts == ["（1）竖向位移监测日报表（共 1 次）；", "（2）立柱结构水平位移监测日报表（共 1 次）；"]
    joined = "".join(texts)
    # 母版里那 12 类（管线/建筑物倾斜/支撑轴力/地下水位/测斜）一个都不能出现
    for word in ("管线", "倾斜", "轴力", "水位", "测斜"):
        assert word not in joined


def test_progress_has_date_and_project_name():
    text = generate(PURPOSE_PROGRESS, _ctx([_row("JJC1", 0.0)], is_initial=True))[0]
    assert "2026年10月8日" in text and PROJECT_NAME in text
    assert "测点初始状态" in text
    assert "工作量如下表" in text


def test_worktable_counts_real_readings():
    """点次按本期**实际读数条数**算（沉降每点 1 条、水平位移每点 X/Y 各 1 条）。"""
    rows = [
        _row("JJC1", 1.0),
        _row("JJC2", 2.0),
        _row("JJC1", -1.0, kind="displacement", component="X"),
        _row("JJC1", -2.0, kind="displacement", component="Y"),
    ]
    texts = generate(PURPOSE_WORKTABLE, _ctx(rows, kinds=["settlement", "displacement"]))
    assert texts[0] == "1"
    assert texts[1] == "竖向位移、立柱结构水平位移"
    assert texts[2] == "4点次"
    assert texts[4] == "第3期"


# ─────────────────────────── 写入行为 ───────────────────────────


def test_write_block_clears_extras_and_hides_their_rows():
    wb = Workbook()
    ws = wb.active
    for r in range(1, 6):
        ws.cell(row=r, column=3, value=f"外来内容{r}")

    written, cleared, overflow = write_block(ws, ["Sheet!C1", "Sheet!C2", "Sheet!C3", "Sheet!C4", "Sheet!C5"],
                                             ["本项目第1段", "本项目第2段"])
    assert (written, cleared, overflow) == (2, 3, [])
    assert ws["C1"].value == "本项目第1段"
    assert all(ws.cell(row=r, column=3).value is None for r in (3, 4, 5)), "外来内容没清干净"
    assert all(ws.row_dimensions[r].hidden for r in (3, 4, 5)), "多余的格子该隐藏行，否则留一大片空白"


def test_write_block_reports_overflow_instead_of_dropping():
    wb = Workbook()
    ws = wb.active
    written, _cleared, overflow = write_block(ws, ["Sheet!C1"], ["第一段", "第二段", "第三段"])
    assert written == 1
    assert overflow == ["第二段", "第三段"], "装不下的内容必须报出来，绝不静默丢"


# ─────────────────────────── 端到端：真实形态的母版 ───────────────────────────


def _layout(with_text_blocks: bool) -> dict:
    cols = {"point_col": 1, "v0_m": 2, "prev_m": 3, "current_m": 4, "prev_cum_mm": 5,
            "cum_mm": 6, "delta_mm": 7, "rate": 8, "warning_level": 11,
            "control_cum": 9, "control_rate": 10, "data_start_row": 7}
    lay = {
        "meta": {"sign_date_cell": "签字!E5"},
        "sheets": [{"sheet_name": SETTLE, "kind": "settlement", "layout": dict(cols), "mode": "per_row"}],
    }
    if with_text_blocks:
        lay["text_blocks"] = [
            {"sheet": SURVEY_SHEET, "purpose": PURPOSE_PROGRESS, "cells": [f"{SURVEY_SHEET}!C8"]},
            {"sheet": SURVEY_SHEET, "purpose": PURPOSE_WORKTABLE, "header_row": 10,
             "first_row": 11, "last_row": 11,
             "cols": {"index": 3, "item": 4, "week": 5, "total": 6, "note": 7}},
            {"sheet": OVERVIEW_SHEET, "purpose": PURPOSE_OVERVIEW, "cells": [f"{OVERVIEW_SHEET}!C1"],
             "keep_until_line_containing": "现场巡视"},
            {"sheet": OVERVIEW_SHEET, "purpose": PURPOSE_ANALYSIS,
             "cells": [f"{OVERVIEW_SHEET}!C{r}" for r in range(3, 7)]},
            {"sheet": SUBMIT_SHEET, "purpose": PURPOSE_CONCLUSION, "cells": [f"{SUBMIT_SHEET}!B3"]},
            {"sheet": SUBMIT_SHEET, "purpose": PURPOSE_DELIVERABLES,
             "cells": [f"{SUBMIT_SHEET}!B{r}" for r in range(9, 13)]},
        ]
    return lay


def _build_master(path: Path) -> None:
    """照真实母版的形态造样表：三页外来文字 + 一张数据表。"""
    wb = Workbook()
    ws = wb.active
    ws.title = SETTLE
    ws["A5"] = "监测点号"
    for i in (1, 2):
        r = 6 + i
        ws.cell(row=r, column=1, value=f"JJC{i}")
        ws.cell(row=r, column=9, value=30)
        ws.cell(row=r, column=10, value=10)

    ov = wb.create_sheet(OVERVIEW_SHEET)
    ov["C1"] = ("4、监测成果及分析说明\n4.1 现场巡视情况\n"
                "        本周对屿前站巡视过程中暂未发现异常。\n"
                "    本周现场巡视主要情况详见本周现场巡视记录。")
    ov.merge_cells("C1:R1")
    ov["C2"] = "4.2 监测成果分析"
    ov.merge_cells("C2:R2")
    foreign_segments = [
        "(1)地表沉降监测：\n   经分析本期累积变化量介于：[-30.1,8.3]之间(单位为：mm，点号分别为：DBC05-01，DBC06-05),数据未发现异常情况。",
        "(2)管线竖向位移监测：\n   经分析本期累积变化量介于：[-21.7,-7.2]之间(单位为：mm，点号分别为：GDC03-10，GGC01-01),数据未发现异常情况。",
        "(9)砼支撑轴力监测：\n   经分析本期累积变化量介于：[197.01,4625.5]之间(单位为：mm，点号分别为：TZL02-02，TZL03-02),数据未发现异常情况。",
        "(12)测斜QCX1监测：\n   经分析本期累积变化量介于：[-3.6,22.37]之间(单位为：mm，点号分别为：0.0，14.0),数据未发现异常情况。",
    ]
    for i, seg in enumerate(foreign_segments):
        r = 3 + i
        ov.cell(row=r, column=3, value=seg)
        ov.merge_cells(f"C{r}:R{r}")

    sub = wb.create_sheet(SUBMIT_SHEET)
    sub["B1"] = "5、监测结论与建议"
    sub["B2"] = "5.1监测结论："
    sub["B3"] = "     本周通过对屿前站的监测数据和巡视信息综合分析：本期监测数据处于稳定状态，未发现新增的异常情况。"
    sub.merge_cells("B3:K3")
    sub["B4"] = "5.2相关建议："
    sub["B5"] = "（1）建议合理安排机械、材料等堆放位置，避免监测点遮挡，并注意测点保护。"
    sub["B8"] = "6、监测成果"
    for i, item in enumerate(("（1）地表沉降监测日报表（共1次）、时程曲线图；",
                             "（9）砼支撑轴力监测日报表（共1次）、时程曲线图；",
                             "（11）地下水位监测日报表（共1次）、时程曲线图；",
                             "（12）测斜QCX1监测日报表（共1次）、时程曲线图；"), start=9):
        sub.cell(row=i, column=2, value=item)

    sur = wb.create_sheet(SURVEY_SHEET)
    sur["B7"] = "1.3监测完成情况："
    sur["C8"] = "（1）本周（2026年5月21日）右幅Y11#基坑水平竖向位移共监测1次测点初始状态。\n（2）本周共巡查1次，监测项目工作量如下表："
    sur.merge_cells("C8:O8")
    for c, v in enumerate(("序号", "监测项目", "本周完成工作量", "累积完成工作量", "备注"), start=3):
        sur.cell(row=10, column=c, value=v)
    sur["C11"], sur["D11"], sur["E11"], sur["G11"] = 1, "水中拉伸钢板桩", "4点次", "初始值"
    wb.save(path)


def _mk_project(client, n_points: int = 2) -> int:
    p = client.post("/api/v1/projects", json={"code": "PRJ-TXT", "name": PROJECT_NAME}).json()
    pid = p["id"]
    put_control_rules(client, pid)
    for i in range(1, n_points + 1):
        pt = client.post(f"/api/v1/projects/{pid}/points",
                         json={"point_no": f"JJC{i}", "point_type": "settlement"}).json()
        for d, v in (("2026-09-01", 3500.0), ("2026-09-02", 3500.0 - i)):
            client.post(f"/api/v1/points/{pt['id']}/records",
                        json={"series_type": "settlement", "component": "value",
                              "obs_date": d, "value_mm": v})
    return pid


def _mk_template(pid: int, path: Path, layout: dict) -> int:
    with SessionLocal() as db:
        tpl = ExportTemplate(project_id=pid, name="文字页母版", template_type="excel",
                             is_builtin=False, file_path=str(path), layout_json=layout)
        db.add(tpl)
        db.commit()
        db.refresh(tpl)
        return tpl.id


def _export(client, pid: int, tid: int):
    r = client.get(f"/api/v1/projects/{pid}/reports/excel",
                   params={"template_id": tid, "obs_date": "2026-09-02"})
    assert r.status_code == 200, r.text
    summary = json.loads(unquote(r.headers.get("x-export-summary", "{}") or "{}"))
    return load_workbook(BytesIO(r.content)), summary


def _all_text(ws) -> str:
    return "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


def test_export_rewrites_text_pages_and_clears_the_rest(client, tmp_path):
    path = tmp_path / "文字页.xlsx"
    _build_master(path)
    pid = _mk_project(client)
    tid = _mk_template(pid, path, _layout(with_text_blocks=True))

    wb, summary = _export(client, pid, tid)

    # 分析页：标题保留、正文换成项目自己的一段（本例只有沉降一类）
    ov = wb[OVERVIEW_SHEET]
    assert "4.1 现场巡视情况" in str(ov["C1"].value), "标题行要保留"
    assert "4.2 监测成果分析" in str(ov["C2"].value)
    assert "竖向位移监测" in str(ov["C3"].value)
    assert "JJC1" in str(ov["C3"].value) and "JJC2" in str(ov["C3"].value)
    for r in (4, 5, 6):
        assert ov.cell(row=r, column=3).value is None, f"母版第{r}段外来内容没清掉"
        assert ov.row_dimensions[r].hidden, "清空的格子要隐藏行"
    assert "DBC05-01" not in _all_text(ov) and "QCX1" not in _all_text(ov)

    # 成果提交页：结论换成本项目、清单只列本项目有的类别
    sub = wb[SUBMIT_SHEET]
    assert PROJECT_NAME in str(sub["B3"].value) and "屿前站" not in str(sub["B3"].value)
    assert "竖向位移监测日报表" in str(sub["B9"].value)
    for r in (10, 11, 12):
        assert sub.cell(row=r, column=2).value is None
        assert sub.row_dimensions[r].hidden
    assert "砼支撑轴力" not in _all_text(sub)

    # 概况页：完成情况与工作量表都按本期数据
    sur = wb[SURVEY_SHEET]
    assert PROJECT_NAME in str(sur["C8"].value) and "2026年9月2日" in str(sur["C8"].value)
    assert sur["D11"].value == "竖向位移"
    assert sur["E11"].value == "2点次"
    assert "水中拉伸钢板桩" not in _all_text(sur)

    # 摘要要告诉界面做了什么
    assert summary["text_generated"]
    assert summary["text_cleared"]
    wb.close()


def test_export_summary_lists_generated_pages(client, tmp_path):
    path = tmp_path / "文字页2.xlsx"
    _build_master(path)
    pid = _mk_project(client)
    tid = _mk_template(pid, path, _layout(with_text_blocks=True))
    _wb, summary = _export(client, pid, tid)
    assert OVERVIEW_SHEET in summary["text_generated"]
    assert OVERVIEW_SHEET in summary["text_cleared"]
    assert not summary.get("text_overflow"), "本例内容装得下，不该报溢出"


def test_old_template_learns_text_blocks_on_export(client, tmp_path):
    """升级前上传的母版 layout 里没有 text_blocks → **导出时自动补学一次并写回**。

    这样用户不用为了这个功能重新上传一遍母版（原件本来就在盘上）。
    """
    path = tmp_path / "老母版.xlsx"
    _build_master(path)
    pid = _mk_project(client)
    tid = _mk_template(pid, path, _layout(with_text_blocks=False))

    wb, summary = _export(client, pid, tid)
    assert summary.get("text_generated"), "老母版应当被自动补学并生成"
    assert "DBC05-01" not in _all_text(wb[OVERVIEW_SHEET])
    wb.close()

    with SessionLocal() as db:
        lay = db.get(ExportTemplate, tid).layout_json
    assert "text_blocks" in lay, "补学结果要写回，避免每次导出都重学"
    assert lay["text_blocks"], "这份母版确实有可识别的文字页"


def test_master_without_text_pages_is_left_alone(client, tmp_path):
    """母版里没有概况/分析/成果提交这几页 → 什么都不做（宁可不改，也不能瞎改）。"""
    path = tmp_path / "无文字页.xlsx"
    wb0 = Workbook()
    ws = wb0.active
    ws.title = SETTLE
    ws["A5"] = "监测点号"
    ws["A7"] = "JJC1"
    wb0.save(path)
    wb0.close()

    pid = _mk_project(client)
    tid = _mk_template(pid, path, _layout(with_text_blocks=False))
    _wb, summary = _export(client, pid, tid)
    assert not summary.get("text_generated")
    assert not summary.get("text_cleared")


def test_relearn_only_touches_text_blocks(client, tmp_path):
    """补学**只补 text_blocks 一个键** —— 已识别好的 meta / 数据表布局不许被改坏。"""
    path = tmp_path / "只补文字块.xlsx"
    _build_master(path)
    pid = _mk_project(client)
    lay = _layout(with_text_blocks=False)
    lay["meta"] = {"sign_date_cell": "签字!E5", "period_cell": "签字!A11"}
    tid = _mk_template(pid, path, lay)

    _export(client, pid, tid)

    with SessionLocal() as db:
        after = db.get(ExportTemplate, tid).layout_json
    assert after["meta"] == lay["meta"], "meta 被改动了"
    assert after["sheets"] == lay["sheets"], "数据表布局被改动了"
