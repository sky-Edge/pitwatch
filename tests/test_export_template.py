# -*- coding: utf-8 -*-
"""测试：导出模板选择（需求③）。"""
import io
from pathlib import Path

import openpyxl
from openpyxl import Workbook

from backend.core.config import settings
from tests.conftest import put_control_rules


def _make_template_xlsx() -> bytes:
    """模拟用户报送模板：竖向位移 + r.立柱结构水平位移，双层表头 R5/R6，数据 R7 起。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "竖向位移"
    ws["A1"] = "基坑沉降监测报表（XX 报送版）"
    ws["A5"] = "监测\n点号"
    ws["B5"] = "初始值（m)"
    ws["C5"] = "上次测值（m）"
    ws["D5"] = "本次测值（m）"
    ws["E5"] = "上次累计（mm)"
    ws["F5"] = "本次累计（mm)"
    ws["G5"] = "本次变化（mm)"
    ws["H5"] = "速率（mm/d）"
    ws["I5"] = "控制值"
    ws["I6"] = "累计"
    ws["J6"] = "速率"
    ws["K5"] = "预警等级"
    ws2 = wb.create_sheet("r.立柱结构水平位移")
    for cell, val in [
        ("A5", "监测\n点号"),
        ("B5", "初始值（m)"),
        ("C5", "上次测值（m）"),
        ("D5", "本次测值（m）"),
        ("F5", "本次累计（mm)"),
        ("H5", "速率（mm/d）"),
        ("K5", "预警等级"),
    ]:
        ws2[cell] = val
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _setup_project(client):
    p = client.post("/api/v1/projects", json={"code": "PRJ001", "name": "模板测试"}).json()
    pid = p["id"]
    put_control_rules(client, pid)
    pt = client.post(f"/api/v1/projects/{pid}/points", json={"point_no": "JJC1", "point_type": "settlement"}).json()
    for d, v in [("2026-07-30", 4907.0), ("2026-07-31", 4906.0)]:
        client.post(
            f"/api/v1/points/{pt['id']}/records",
            json={"series_type": "settlement", "component": "value", "obs_date": d, "value_mm": v},
        )
    return pid, pt["id"]


def test_builtin_templates_seeded(client):
    """启动后内置模板存在（标准 Excel + Word）。"""
    pid, _ = _setup_project(client)
    lst = client.get(f"/api/v1/projects/{pid}/export-templates").json()
    builtin = [t for t in lst if t["is_builtin"]]
    assert len(builtin) == 2
    assert any(t["template_type"] == "word" for t in builtin)


def test_upload_template_auto_detect_layout(client):
    """上传模板 → 自动识别列映射。"""
    pid, _ = _setup_project(client)
    r = client.post(
        f"/api/v1/projects/{pid}/export-templates/upload",
        data={"name": "报送版模板"},
        files={"file": ("报送版.xlsx", _make_template_xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 200
    d = r.json()
    assert d["detected_layout"] is not None
    # 竖向位移 Sheet 找到点号列、初始值列
    assert d["id"] > 0
    lst = client.get(f"/api/v1/projects/{pid}/export-templates").json()
    mine = [t for t in lst if not t["is_builtin"]]
    assert len(mine) == 1
    assert mine[0]["layout_json"]["sheets"][0]["layout"]["point_col"] == 1
    assert mine[0]["layout_json"]["sheets"][0]["layout"]["v0_m"] == 2


def test_export_excel_with_template_id(client):
    """带 template_id 导出 → 以模板原件为底模，数据写入（保留模板标题）。"""
    pid, pt = _setup_project(client)
    r = client.post(
        f"/api/v1/projects/{pid}/export-templates/upload",
        data={"name": "报送版"},
        files={"file": ("报送版.xlsx", _make_template_xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    tid = r.json()["id"]

    resp = client.get(f"/api/v1/projects/{pid}/reports/excel", params={"template_id": tid})
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb["竖向位移"]
    # 模板标题保留（证明以模板为底模）
    assert "报送版" in str(ws["A1"].value)
    # 数据写入 R7：JJC1 本次测值 4.906
    assert ws["A7"].value == "JJC1"
    assert abs(float(ws["D7"].value) - 4.906) < 1e-6
    wb.close()


def test_export_without_template_id_still_builtin(client):
    """不传 template_id → 系统内置标准版式（回归保障）。"""
    pid, _ = _setup_project(client)
    resp = client.get(f"/api/v1/projects/{pid}/reports/excel")
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["竖向位移", "r.立柱结构水平位移"]
    assert wb["竖向位移"]["A1"].value == "沉降监测报表"  # 内置版式标题
    wb.close()


def test_upload_same_name_replaces_and_no_orphan(client, tmp_path):
    """同名上传自动替换：不产生记录堆积与孤儿文件；上传中转目录无活跃临时文件。"""

    import openpyxl

    pid = client.post("/api/v1/projects", json={"code": "PRJ001", "name": "同名替换测试"}).json()["id"]

    def _mk(extra_cell_value: float):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "竖向位移"
        ws["A1"] = "沉降监测报表"
        ws["A5"] = "监测\n点号"
        ws["B5"] = "初始值（m)"
        ws["C5"] = "上次测值（m）"
        ws["D5"] = "本次测值（m）"
        ws["I5"] = "控制值"
        ws["I6"] = "累计变化量（mm)"
        ws["J6"] = "变化速率（mm/d）"
        ws["A7"] = "JJC1"
        ws["B7"] = 4.907
        ws["D7"] = extra_cell_value
        import io

        b = io.BytesIO()
        wb.save(b)
        return b.getvalue()

    # 上传两次同名
    r1 = client.post(f"/api/v1/projects/{pid}/export-templates/upload",
                     data={"name": "报送母版"},
                     files={"file": ("a.xlsx", _mk(4.909), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r1.status_code == 200
    r2 = client.post(f"/api/v1/projects/{pid}/export-templates/upload",
                     data={"name": "报送母版"},
                     files={"file": ("a.xlsx", _mk(4.91), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r2.status_code == 200
    # 同名后只有 1 条项目模板（自动替换）
    lst = client.get(f"/api/v1/projects/{pid}/export-templates").json()
    mine = [t for t in lst if not t["is_builtin"]]
    assert len(mine) == 1
    assert mine[0]["name"] == "报送母版"
    # 上传中转目录（模板独立化后统一为 templates/_incoming，不再按项目分目录）
    # 断言：不留"活跃"的裸临时文件；且确实做过 rename（存在 .deleted 变体）
    incoming = settings.templates_dir / "_incoming"
    leftovers = [
        p for p in incoming.glob("_upload_tmp*")
        if not p.name.endswith(".deleted") and ".deleted_" not in p.name
    ]
    assert not leftovers, f"上传临时文件没清理干净：{leftovers}"
    assert any(".deleted" in p.name for p in incoming.glob("_upload_tmp*")), "rename 未生效"
    # 新文件 file_ok=True
    assert mine[0]["file_ok"] is True
    # 原件落在模板库目录下（模板不再归属某个项目）
    assert "_library" in (mine[0]["file_path"] or ""), mine[0]["file_path"]


# ---------- 上传与改名不该被"顺序"卡住（2026-10-06 用户反馈"很反人类"） ----------
#
# 原流程：弹窗里"母版名称"在上、"选择整卷文件上传"在下，点上传却没填名字就被退回
# 「请先填写模板名称」—— 用户想"先挑文件再想名字"做不到；而且传完就再也改不了名。
# 现在的规则：**名字可留空（用文件名兜底），传完还能重命名**。

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SRC = Path(__file__).resolve().parent.parent / "frontend" / "src"


def _upload(client, pid, *, name=None, filename="八号基坑报送版.xlsx"):
    """name=None 表示"完全不带这个字段"（模拟用户没填）。"""
    data = {} if name is None else {"name": name}
    return client.post(
        f"/api/v1/projects/{pid}/export-templates/upload",
        data=data,
        files={"file": (filename, _make_template_xlsx(), XLSX_MIME)},
    )


def test_upload_without_name_uses_filename(client):
    """不填名称也能传；名称取文件名（去扩展名）。"""
    pid, _ = _setup_project(client)
    r = _upload(client, pid, name=None, filename="八号基坑报送版.xlsx")
    assert r.status_code == 200, r.text
    assert r.json()["name"] == "八号基坑报送版"


def test_upload_with_blank_name_uses_filename(client):
    """只敲了空格也算"没填"，别把空名字存进去（下拉框里会变成一行空标签）。"""
    pid, _ = _setup_project(client)
    r = _upload(client, pid, name="   ", filename="十号基坑整卷.xlsm")
    assert r.status_code == 200, r.text
    assert r.json()["name"] == "十号基坑整卷"


def test_upload_name_can_be_renamed_afterwards(client):
    """**先上传、后改名**（用户要的就是这个顺序）。"""
    pid, _ = _setup_project(client)
    tid = _upload(client, pid, name=None, filename="随手存的文件.xlsx").json()["id"]

    r = client.put(f"/api/v1/export-templates/{tid}", json={"name": "八号基坑报送母版"})
    assert r.status_code == 200, r.text
    assert r.json()["name"] == "八号基坑报送母版"
    lst = client.get(f"/api/v1/projects/{pid}/export-templates").json()
    assert [t["name"] for t in lst if not t["is_builtin"]] == ["八号基坑报送母版"]


def test_rename_to_duplicate_is_rejected(client):
    """改成一个已存在的名字 → 拦下并说清为什么（两个同名母版在下拉框里分不出来）。"""
    pid, _ = _setup_project(client)
    _upload(client, pid, name="母版A")
    b = _upload(client, pid, name="母版B").json()["id"]

    r = client.put(f"/api/v1/export-templates/{b}", json={"name": "母版A"})
    assert r.status_code == 400
    detail = r.json()["detail"]
    assert "母版A" in detail and "换个名字" in detail, detail
    # 原有两个都还在、都没被改
    names = sorted(t["name"] for t in client.get(f"/api/v1/projects/{pid}/export-templates").json() if not t["is_builtin"])
    assert names == ["母版A", "母版B"]


def test_builtin_template_cannot_be_renamed(client):
    pid, _ = _setup_project(client)
    builtin = next(
        t for t in client.get(f"/api/v1/projects/{pid}/export-templates").json() if t["is_builtin"]
    )
    r = client.put(f"/api/v1/export-templates/{builtin['id']}", json={"name": "乱改"})
    assert r.status_code == 400


def test_rename_does_not_touch_template_file(client, tmp_path):
    """改名只改显示名称：原件仍在、布局不变（改名不该把母版弄丢）。"""
    pid, _ = _setup_project(client)
    tid = _upload(client, pid, name="旧名字").json()["id"]
    before = next(t for t in client.get(f"/api/v1/projects/{pid}/export-templates").json() if t["id"] == tid)

    client.put(f"/api/v1/export-templates/{tid}", json={"name": "新名字"})
    after = next(t for t in client.get(f"/api/v1/projects/{pid}/export-templates").json() if t["id"] == tid)
    assert after["name"] == "新名字"
    assert after["file_path"] == before["file_path"]
    assert after["file_ok"] is True
    assert after["layout_json"] == before["layout_json"]


def _code_only_tsx(src: str) -> str:
    """剥掉注释（含跨行的 `{/* … */}` 块注释），只留会执行的代码。

    否则"解释为什么不再这么拦"的注释会把守卫自己绊倒 —— 注释里写着的旧提示语
    并不代表代码里还在拦人（这条守卫头两次跑都是这么挂的：先被 `//` 绊、再被块注释绊）。
    """
    out: list[str] = []
    in_block = False
    for line in src.splitlines():
        s = line.strip()
        if in_block:
            if "*/" in s:
                in_block = False
            continue
        if s.startswith(("/*", "{/*")):
            if "*/" not in s:
                in_block = True
            continue
        if s.startswith(("//", "*")):
            continue
        out.append(line)
    return "\n".join(out)


def test_frontend_upload_flow_not_blocked_by_name() -> None:
    """源码守卫：上传不再被"名字"拦住；且界面上提供事后改名入口。"""
    code = _code_only_tsx((SRC / "pages" / "ReportsPage.tsx").read_text(encoding="utf-8"))
    assert "请先填写模板名称" not in code, "还在拦没填名字的用户"
    assert "renameTemplate" in code and "重命名" in code, "没有上传后改名的入口"
    assert "留空就用文件名" in code, "界面上没告诉用户名字可以留空"
