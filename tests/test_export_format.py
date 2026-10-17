# -*- coding: utf-8 -*-
"""导出报表的「数值显示格式 + 列宽」收拾（2026-10-08 用户反馈）。

用户原话：「你改完后没有顺手把单元格宽度控制在最佳位置 我看了会有小数点呢」

**实测到的根因**（拿用户真实母版导出、逐格读 number_format）：
母版是从真实报表另存来的，Excel 给数值格自动挂了 `0.00000_ ` 格式 —— 数据表 B~H、
水平位移表的 I/J 控制值、原始数据区的坐标格**全是**。系统写数值时只改值、不改格式：
- 控制值写进去 → 显示 `20.00000`、`3.00000`；
- 坐标 3413192.929 → 显示 `3413192.92900`；
- 而且水平位移表的 C/D 列（测值/坐标）母版没设宽度 → 默认 9，11 个字符的坐标根本放不下。

**现在三条规则**（`backend/reports/export_format.py`）：
1. 写数值的格子一律「常规」格式 —— 整数显示整数、小数按实际精度；
2. 列宽**只加宽、不放窄**，且**跳过自动换行的表头** —— 否则会把表头文字算进去，
   实测能把 A 列算到 70 宽（A 列 9 宽是母版为了 A4 版式定的）；
3. 写进去的数一律先收到**纸上口径**（毫米 1 位 = 米 4 位），且累计/变化/速率由
   "纸上那个读数"反算 —— 甲方拿计算器逐行核也必须对得上（2026-10-16 用户报障）。
"""
import openpyxl
import pytest

from tests.test_control_values import (
    XLSX_MIME,
    _ctl,
    _download_master,
    _put_rules,
    _rules,
    _upload_master,
)
from tests.test_master_report import _make_master, _seed_project

SRC = __import__("pathlib").Path(__file__).resolve().parent.parent / "backend"


def _master_with_legacy_format(tmp_path, wide_cols=("B", "L")) -> "object":
    """同构母版 + **真实母版那种残留格式**（数值列全挂 `0.00000_ `）。

    wide_cols 里的列显式设成较宽（模拟母版刻意设过的列）—— 导出后必须原样保留。
    """
    wb = openpyxl.load_workbook(_make_master())
    for sn in ("竖向位移", "r.立柱结构水平位移"):
        ws = wb[sn]
        for r in range(7, 18):
            for c in range(2, 11):  # B~J：初始值/测值/累计/变化/速率/控制值
                ws.cell(row=r, column=c).number_format = "0.00000_ "
        # 表头带自动换行（真实母版就是这样）——列宽不该被表头文字拉宽
        ws.cell(row=6, column=9).alignment = openpyxl.styles.Alignment(wrap_text=True)
        ws.cell(row=5, column=2).alignment = openpyxl.styles.Alignment(wrap_text=True)
    wb["竖向位移"].column_dimensions["B"].width = 15.6
    wb["竖向位移"].column_dimensions["L"].width = 16.5
    wb["r.立柱结构水平位移"].column_dimensions["B"].width = 15.6
    wb["r.立柱结构水平位移"].column_dimensions["L"].width = 17.6
    out = tmp_path / "master_legacy.xlsx"
    wb.save(out)
    wb.close()
    return out


def _width(ws, col: str) -> float:
    dim = ws.column_dimensions.get(col)
    return (dim.width if dim is not None and dim.width else None) or (ws.sheet_format.defaultColWidth or 8.43)


def _prepare(client, tmp_path):
    pid = _seed_project(client)
    _put_rules(client, pid, [_rules("settlement", 30, 10), _rules("displacement", 20, 3)])
    tid = _upload_master(client, pid, _master_with_legacy_format(tmp_path))
    return pid, tid


# ---------- 数值格式 ----------


def test_legacy_decimal_format_is_cleared(client, tmp_path):
    """母版残留的 `0.00000_ ` 必须被清掉，且换成本项目定的那套格式。

    - B/C/D（测值，米）→ 「常规」：3.119 就是 3.119，不是 3.11900；
    - E/F/G/H（累计/变化/速率，mm）→ `0.0##`：**要小数点**（整数也显示 2.0）；
    - I/J（控制值）→ 「常规」：20、3 不带小数点。

    判据取"有值的格子"而不是固定列号：测试母版里没被识别进 layout 的列本来就不会写。
    """
    from backend.reports.export_format import GENERAL, MM_DELTA_FORMAT

    pid, tid = _prepare(client, tmp_path)
    wb = _download_master(client, pid, tid)

    checked = 0
    for sn in ("竖向位移", "r.立柱结构水平位移"):
        ws = wb[sn]
        for r in range(7, (ws.max_row or 0) + 1):
            for c in range(2, 12):
                cell = ws.cell(row=r, column=c)
                v = cell.value
                if not isinstance(v, (int, float)) or isinstance(v, bool):
                    continue
                want = MM_DELTA_FORMAT if c in (5, 6, 7, 8) else GENERAL
                assert cell.number_format == want, (
                    f"[{sn}] {r}行{c}列 格式应为 {want}，实际 {cell.number_format}（值 {v}）"
                )
                checked += 1
    assert checked >= 40, f"只检查到 {checked} 个数值格，这条测试等于没写"

    ws = wb["竖向位移"]
    assert _ctl(ws, 7, 9) == 30 and _ctl(ws, 7, 10) == 10, "控制值仍应写成规则值"
    wb.close()


# ---------- 列宽 ----------


def test_narrow_columns_are_widened_to_fit(client, tmp_path):
    """坐标 `3413192.929`（11 字符）在默认 9 宽的列里放不下 → 必须加宽。

    母版的 C/D 列（测值）没设宽度、吃 sheet 默认值 9；水平位移表里那两列装的是
    绝对坐标，宽度不够会被 Excel 显示成科学计数/掉小数 —— 用户说的"列宽没在最佳位置"。
    """
    pid, tid = _prepare(client, tmp_path)
    wb = _download_master(client, pid, tid)
    wsd = wb["r.立柱结构水平位移"]
    for col in ("C", "D"):
        assert _width(wsd, col) >= 12, f"{col} 列宽 {_width(wsd, col)} 放不下 3413192.929"
    wb.close()


def test_master_column_widths_are_never_narrowed(client, tmp_path):
    """母版刻意设过的列宽一律不动 —— 版式与 A4 打印分页是它的决定。"""
    pid, tid = _prepare(client, tmp_path)
    wb = _download_master(client, pid, tid)
    assert _width(wb["竖向位移"], "B") == pytest.approx(15.6)
    assert _width(wb["竖向位移"], "L") == pytest.approx(16.5)
    assert _width(wb["r.立柱结构水平位移"], "B") == pytest.approx(15.6)
    assert _width(wb["r.立柱结构水平位移"], "L") == pytest.approx(17.6)
    wb.close()


def test_wrapped_header_does_not_drive_column_width(client, tmp_path):
    """带自动换行的表头不参与测宽。

    表头文字（「累计变化量（mm)」= 15 宽）如果算进去，会把窄列一律撑到 17+；
    而母版给那些列的就是 8~9 —— 表头本来就折行显示，不需要加宽。
    """
    pid, tid = _prepare(client, tmp_path)
    wb = _download_master(client, pid, tid)
    wsd = wb["r.立柱结构水平位移"]
    # 这两列母版没设宽度（默认 9）且内容只有 20 / 3 → 不该因为表头文字被拉宽
    for col in ("I", "J"):
        assert _width(wsd, col) <= 10.5, f"{col} 列被表头文字拉宽到 {_width(wsd, col)}"
    wb.close()


def test_export_is_idempotent_on_width(client, tmp_path):
    """连导两次，列宽一致（不会一次比一次宽）。"""
    pid, tid = _prepare(client, tmp_path)
    a = _download_master(client, pid, tid)
    b = _download_master(client, pid, tid)
    for col in ("A", "B", "C", "D", "I", "J", "L"):
        assert _width(a["竖向位移"], col) == pytest.approx(_width(b["竖向位移"], col)), col
    a.close()
    b.close()


def test_mm_columns_keep_a_decimal_point(client, tmp_path):
    """「上次累计 / 本次累计 / 本次变化 / 变化速率」四列**必须有小数点**（用户 2026-10-08 要求）。

    用户看到的是一列整数（-4、-5、2、2），觉得"怎么没有小数点"。这四列是毫米读数，
    按现场习惯要带小数 —— 整数也显示成 `2.0`。

    ⚠ 为什么用 `0.0##` 而不是固定两位小数：固定两位会把 0.125 显示成 0.13，
    报表上印的数与系统判定用的值就对不上了（那是我最不想制造的问题）。
    """
    import io

    from backend.reports.export_format import MM_DELTA_FORMAT

    assert MM_DELTA_FORMAT.startswith("0.0"), "至少要有一位小数，否则整数还是不带小数点"
    assert "#" in MM_DELTA_FORMAT, "多余位数应可省略（0.5 不该写成 0.500）"

    pid, tid = _prepare(client, tmp_path)
    MM_COLS = (5, 6, 7, 8)

    wb = _download_master(client, pid, tid)
    checked = 0
    for sn in ("竖向位移", "r.立柱结构水平位移"):
        ws = wb[sn]
        for r in range(7, (ws.max_row or 0) + 1):
            for c in MM_COLS:
                cell = ws.cell(row=r, column=c)
                if not isinstance(cell.value, (int, float)):
                    continue
                assert cell.number_format == MM_DELTA_FORMAT, (
                    f"[{sn}] {r}行{c}列 应带小数点，实际格式 {cell.number_format}"
                )
                checked += 1
    assert checked >= 8, f"只检查到 {checked} 个单元格，这条测试等于没写"
    wb.close()

    resp = client.get(f"/api/v1/projects/{pid}/reports/excel", params={"obs_date": "2026-08-21"})
    assert resp.status_code == 200, resp.text
    std = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)
    for sn in ("竖向位移", "r.立柱结构水平位移"):
        for r in range(7, 11):
            for c in MM_COLS:
                assert std[sn].cell(row=r, column=c).number_format == MM_DELTA_FORMAT, f"{sn} {r}行{c}列"
    std.close()


# ---------- 纸上口径：读数(m) 与 累计/变化(mm) 必须能手算复核 ----------


def test_paper_values_reconcile_by_hand() -> None:
    """同一行里米列和毫米列必须能用计算器对上（2026-10-16 用户报"手算对不上"）。

    拿用户台账上真实出现过的数：初始 4908.83mm、本次 4909.23mm。
    老口径把米列 round 到 3 位 → 纸上印 `4.909 / 4.909`，而累计那格印 `0.4`，
    (4.909−4.909)×1000 = 0 ≠ 0.4 —— 报表自己打自己脸。
    """
    from decimal import Decimal

    from backend.reports.export_format import paper_values

    p = paper_values(4908.83, 4909.06, 4909.23, skipped_days=1)
    assert (p["v0_m"], p["prev_m"], p["current_m"]) == (4.9088, 4.9091, 4.9092)
    assert (p["prev_cum_mm"], p["cum_mm"], p["delta_mm"], p["rate"]) == (0.3, 0.4, 0.1, 0.1)

    # 用 Decimal 复核 —— 人拿计算器就是这么算的；float 会把差 1e-13 判成"约等于"
    d = {k: Decimal(str(v)) for k, v in p.items()}
    assert (d["current_m"] - d["prev_m"]) * 1000 == d["delta_mm"]
    assert (d["current_m"] - d["v0_m"]) * 1000 == d["cum_mm"]
    assert (d["prev_m"] - d["v0_m"]) * 1000 == d["prev_cum_mm"]

    # 首期：没有"上次"，变化量/上次累计留空，累计 = 0（不是 None）
    first = paper_values(4908.83, None, 4908.83, None)
    assert first["prev_m"] is None and first["delta_mm"] is None
    assert first["prev_cum_mm"] is None and first["cum_mm"] == 0.0
    assert first["rate"] is None

    # 微小回弹不许在纸上印出 "-0.0"
    assert str(paper_values(4908.8, 4908.84, 4908.8, 1)["cum_mm"]) == "0.0"


def _dec(ws, row: int, col: int):
    from decimal import Decimal

    v = ws.cell(row=row, column=col).value
    if not isinstance(v, (int, float)) or isinstance(v, bool):
        return None
    return Decimal(str(v))


def _has_col(ws, col: int) -> bool:
    """这张表有没有这一列（表头两行里有没有字）—— 测试母版是精简版，水平位移表
    就没排「上次累计 / 本次变化」两列；表上没有的列不能拿来判"对不上"。"""
    for r in (5, 6):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and v.strip():
            return True
    return False


def _hand_check_ledger(ws, label: str, first_row: int = 7) -> int:
    """逐行做甲方那笔手算。返回核到的行数（0 行说明这条测试等于没写）。"""
    cols = {c: _has_col(ws, c) for c in (3, 5, 7)}   # 上次测值 / 上次累计 / 本次变化
    checked = 0
    for r in range(first_row, (ws.max_row or 0) + 1):
        v0, prev, cur = _dec(ws, r, 2), _dec(ws, r, 3), _dec(ws, r, 4)
        prev_cum, cum, delta = _dec(ws, r, 5), _dec(ws, r, 6), _dec(ws, r, 7)
        if cur is None or v0 is None or cum is None:
            continue
        assert (cur - v0) * 1000 == cum, f"{label} {r}行：本次累计 {cum} ≠ (本次 {cur} − 初始 {v0})×1000"
        if cols[3] and prev is not None:
            if cols[7]:
                assert (cur - prev) * 1000 == delta, (
                    f"{label} {r}行：本次变化量 {delta} ≠ (本次 {cur} − 上次 {prev})×1000"
                )
            if cols[5]:
                assert (prev - v0) * 1000 == prev_cum, (
                    f"{label} {r}行：上次累计 {prev_cum} ≠ (上次 {prev} − 初始 {v0})×1000"
                )
        checked += 1
    return checked


def _seed_sub_mm_project(client, pno="PRJ-SUBMM"):
    """三期读数，期期都带 0.01mm 级的零头 —— 就是"手算对不上"的最小复现。"""
    pid = client.post("/api/v1/projects", json={"code": pno, "name": "口径测试"}).json()["id"]
    _put_rules(client, pid, [_rules("settlement", 30, 10), _rules("displacement", 20, 3)])
    readings = {
        "JJC1": {"value": [4908.83, 4909.06, 4909.23], "X": [3413055119.37, 3413055119.61, 3413055119.84],
                 "Y": [510174159.42, 510174159.18, 510174159.07]},
        "JJC2": {"value": [4912.27, 4912.27, 4912.61], "X": [3413056233.05, 3413056233.44, 3413056233.92],
                 "Y": [510175002.66, 510175002.71, 510175002.35]},
    }
    for no, comps in readings.items():
        pt = client.post(f"/api/v1/projects/{pid}/points", json={"point_no": no, "point_type": "both"}).json()["id"]
        for comp, series in (("value", "settlement"), ("X", "displacement"), ("Y", "displacement")):
            for i, d in enumerate(("2026-08-19", "2026-08-20", "2026-08-21")):
                r = client.post(f"/api/v1/points/{pt}/records",
                                json={"series_type": series, "component": comp,
                                      "obs_date": d, "value_mm": comps[comp][i]})
                assert r.status_code == 201, r.text
    return pid


def test_exported_ledger_reconciles_by_hand(client, tmp_path) -> None:
    """**导出的纸面上**逐行手算复核 —— 母版套打与标准版式两条路都要过。

    判据写在导出件上（而不是写在 `_series_rows` 的返回值上）：用户看的是纸，
    中间层对得上、落到 Excel 又被格式或二次 round 改掉，一样是错。
    """
    import io

    pid = _seed_sub_mm_project(client)

    tid = _upload_master(client, pid, _master_with_legacy_format(tmp_path))
    wb = _download_master(client, pid, tid)
    total = 0
    for sn in ("竖向位移", "r.立柱结构水平位移"):
        total += _hand_check_ledger(wb[sn], f"母版套打/{sn}")
    wb.close()
    assert total >= 4, f"只核到 {total} 行，这条测试等于没写"

    resp = client.get(f"/api/v1/projects/{pid}/reports/excel", params={"obs_date": "2026-08-21"})
    assert resp.status_code == 200, resp.text
    std = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)
    total = 0
    for sn in ("竖向位移", "r.立柱结构水平位移"):
        total += _hand_check_ledger(std[sn], f"标准版式/{sn}")
    std.close()
    assert total >= 4, f"只核到 {total} 行，这条测试等于没写"


def test_rate_is_derivable_from_the_printed_delta(client, tmp_path) -> None:
    """速率那一格必须能用纸上的「本次变化量 ÷ 间隔天数」复核出来（隔 1 天时应完全相等）。"""
    pid = _seed_sub_mm_project(client)
    tid = _upload_master(client, pid, _master_with_legacy_format(tmp_path))
    wb = _download_master(client, pid, tid)
    ws = wb["竖向位移"]
    checked = 0
    for r in range(7, (ws.max_row or 0) + 1):
        delta, rate = _dec(ws, r, 7), _dec(ws, r, 8)
        if delta is None or rate is None:
            continue
        # 三期都是隔天观测 → Δt = 1 天，速率就该等于变化量
        assert rate == delta, f"{r}行：速率 {rate} ≠ 变化量 {delta} ÷ 1 天"
        checked += 1
    wb.close()
    assert checked >= 2, f"只核到 {checked} 行，这条测试等于没写"


# ---------- 源码守卫 ----------


def test_all_export_paths_share_the_format_helpers() -> None:
    """三条导出路径里"要写数值"的两条（母版套打 / 旧模板）都得走同一套格式处理。

    否则又会是"换个模板就漏了"——这个项目里已经栽过（导出不显示预警那次）。
    """
    master = (SRC / "reports" / "master_report.py").read_text(encoding="utf-8")
    legacy = (SRC / "reports" / "template_excel.py").read_text(encoding="utf-8")
    assert "export_format" in master and "_set_num" in master
    assert "export_format" in legacy and "_set_num_cell" in legacy
    assert "MM_DELTA_FORMAT" in master and "MM_DELTA_FORMAT" in legacy, "「带小数点」的口径要三条路径一致"
    # 格式只允许在 export_format 里定义
    others = [
        p for p in (SRC / "reports").glob("*.py")
        if p.name != "export_format.py" and "number_format =" in p.read_text(encoding="utf-8")
    ]
    assert not others, f"这些文件自己设了 number_format，应统一走 export_format：{others}"


def test_paper_caliber_has_a_single_source() -> None:
    """「读数 + 派生量」只能出自 `paper_values`，mm→m 只能出自 `mm_to_m`。

    这次的病根就是各条路各写一份：母版套打/旧模板把米列 round 到 3 位，
    毫米列却直接抄引擎原值 —— 两份都"对"，摆在一行里就是错的。
    """
    import re

    for name in ("template_excel.py", "excel_report.py"):
        src = (SRC / "reports" / name).read_text(encoding="utf-8")
        assert "paper_values" in src, f"{name} 没走 paper_values，纸上会手算对不上"

    offenders = [
        p.name for p in (SRC / "reports").glob("*.py")
        if p.name != "export_format.py" and re.search(r"/\s*1000", p.read_text(encoding="utf-8"))
    ]
    assert not offenders, f"这些文件自己把 mm 换成 m，应统一走 export_format.mm_to_m：{offenders}"

    # 口径常量也只有一处；报表文字（text_blocks）用的是同一个数
    tb = (SRC / "reports" / "text_blocks.py").read_text(encoding="utf-8")
    assert "MM_DECIMALS" in tb, "分析段落里的小数位必须与数据表同口径"
