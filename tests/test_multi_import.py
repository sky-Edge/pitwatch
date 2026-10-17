# -*- coding: utf-8 -*-
"""跨项目导入（一次上传、自动分到各项目）测试。

覆盖的口径见 `backend/services/multi_import_service.py` 模块注释。这里额外焊死两条
**出过事就会很贵**的规则：

1. 一个批次只管自己那几张表 —— 否则改日期重新核对会把别的项目的表一起拉进来，
   用户点确认时别人的数据就串进了这个项目（`importer.sheets_of_batch`）；
2. 认不出来的表**不建批次、不落库**，只摆出候选让用户手选 —— 认错比认不出糟得多，
   两边的点号还常常恰好同名（都叫 Z1、Z2…），校验会一路绿灯。
"""
import io

import openpyxl

HEADER_MAIN = [
    "监测\n点号", "初始值（m)", "上次测值（m）", "本次测值（m）",
    "上次累计变化量（mm)", "本次累计变化量（mm)", "本次变化量（mm)", "变化速率（mm/d）",
    "控制值", None, "预警\n等级", "备注",
]
HEADER_SUB = [None] * 8 + ["累计变化量（mm)", "变化速率（mm/d）", None, None]


def _add_ledger_sheet(wb, sheet_name: str, proj_name: str | None, points: list[str],
                      obs_date: str = "2026-07-31") -> None:
    """加一张与真实台账同构的「竖向位移」表（双层表头，数据从第 7 行起）。

    `proj_name=None` = 表里不写「监测工程名称」（老模板就是这样），用来试文件名兜底。
    """
    ws = wb.create_sheet(sheet_name)
    meta = ["沉降监测报表"]
    if proj_name:
        meta.append(f"监测工程名称：{proj_name}")
    meta += [f"本次监测时间：{obs_date}", "仪器型号：Leica TZ05"]
    for i, v in enumerate(meta, start=1):
        ws.cell(row=i, column=1, value=v)
    for i, v in enumerate(HEADER_MAIN, start=1):
        ws.cell(row=5, column=i, value=v)
    for i, v in enumerate(HEADER_SUB, start=1):
        ws.cell(row=6, column=i, value=v)
    for r, no in enumerate(points, start=7):
        ws.cell(row=r, column=1, value=no)
        ws.cell(row=r, column=2, value=4.907)
        ws.cell(row=r, column=3, value=4.907)
        ws.cell(row=r, column=4, value=4.906)


def _xlsx(*specs, filename: str = "台账.xlsx") -> tuple[str, bytes]:
    """按 [(表名, 工程名称, [点号]), ...] 造一份台账，返回 (文件名, 字节)。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for spec in specs:
        _add_ledger_sheet(wb, spec[0], spec[1], spec[2], *(spec[3:] or ()))
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return filename, buf.getvalue()


_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _project(client, code: str, name: str, points: list[str]) -> int:
    p = client.post("/api/v1/projects", json={"code": code, "name": name}).json()
    for no in points:
        client.post(
            f"/api/v1/projects/{p['id']}/points",
            json={"point_no": no, "point_type": "settlement", "location_desc": "桩顶"},
        )
    return p["id"]


def _upload(client, *files):
    """files: (名字, 字节) 或 (名字, 字节, MIME)。

    非 Excel 要如实报 MIME —— 拿 xlsx 的糊弄过去，测的就不是"用户拖错文件"这条路了。
    """
    parts = [("files", (f[0], f[1], f[2] if len(f) > 2 else _XLSX_MIME)) for f in files]
    return client.post("/api/v1/import/multi/upload", files=parts)


def _two_pits(client):
    """八号基坑(JJC1/JJC2) + 十号基坑(Z1/Z2) —— 点号**故意不重名**，便于断言串没串。"""
    a = _project(client, "PRJ-08", "八号基坑", ["JJC1", "JJC2"])
    b = _project(client, "PRJ-10", "十号基坑", ["Z1", "Z2"])
    return a, b


# ────────────────────────────── 认项目 ──────────────────────────────


def test_one_file_two_pits_routes_each_sheet_to_its_own_project(client):
    """一份文件里混着两个基坑的表 → 各归各的项目，各自一个批次。"""
    a, b = _two_pits(client)
    name, data = _xlsx(
        ("八号竖向位移", "兴善公路 八号基坑围护工程", ["JJC1", "JJC2"]),
        ("十号竖向位移", "十号基坑", ["Z1", "Z2"]),
    )
    resp = _upload(client, (name, data))
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["errors"] == [] and body["pending"] == []
    assert len(body["groups"]) == 2
    by_pid = {g["project_id"]: g for g in body["groups"]}
    assert set(by_pid) == {a, b}
    # 表里写的是带地名前缀的长名字，也要认出来（且不该只算"相近"）
    g8 = by_pid[a]
    assert [s["sheet_name"] for s in g8["sheets"]] == ["八号竖向位移"]
    assert g8["needs_eyeball"] is False
    assert "八号" in g8["sheets"][0]["how"] or g8["sheets"][0]["score"] >= 0.9
    # 每组只含自己项目的点
    assert {r["point_no"] for r in g8["preview"]["valid_rows"]} == {"JJC1", "JJC2"}
    assert {r["point_no"] for r in by_pid[b]["preview"]["valid_rows"]} == {"Z1", "Z2"}


def test_upload_alone_writes_nothing(client):
    """上传 + 认项目**一条数据都不落库** —— 落库只能是用户点过的那个 confirm。"""
    a, _b = _two_pits(client)
    name, data = _xlsx(("竖向位移", "八号基坑", ["JJC1", "JJC2"]))
    assert _upload(client, (name, data)).status_code == 200
    pts = client.get(f"/api/v1/projects/{a}/points").json()
    for p in pts:
        assert client.get(f"/api/v1/points/{p['id']}/records").json() == []


def test_confirm_through_the_shared_endpoint_lands_in_the_right_project(client):
    """分好组之后走**已有的** confirm，数据落在认出来的那个项目里。"""
    a, b = _two_pits(client)
    name, data = _xlsx(
        ("八号竖向位移", "八号基坑", ["JJC1", "JJC2"]),
        ("十号竖向位移", "十号基坑", ["Z1", "Z2"]),
    )
    groups = {g["project_id"]: g for g in _upload(client, (name, data)).json()["groups"]}
    for pid in (a, b):
        g = groups[pid]
        rows = g["preview"]["valid_rows"]
        conf = client.post(
            f"/api/v1/import/batches/{g['batch_id']}/confirm", json={"valid_rows": rows}
        )
        assert conf.status_code == 200, conf.text
        assert conf.json()["imported_rows"] == 2
    # 八号基坑里只能有八号的点，每条读数都出自八号那个批次；十号同理。
    # （records 接口不带 point_no，点号要从 points 上取）
    for pid in (a, b):
        pts = client.get(f"/api/v1/projects/{pid}/points").json()
        assert {p["point_no"] for p in pts} == ({"JJC1", "JJC2"} if pid == a else {"Z1", "Z2"})
        for p in pts:
            recs = client.get(f"/api/v1/points/{p['id']}/records").json()
            assert [r["obs_date"] for r in recs] == ["2026-07-31"]
            assert {r["import_batch_id"] for r in recs} == {groups[pid]["batch_id"]}


def test_filename_fallback_is_flagged_for_a_second_look(client):
    """表里没写工程名称 → 只能按文件名认，并且**明说**是按文件名认的。"""
    _a, b = _two_pits(client)
    name, data = _xlsx(("竖向位移", None, ["Z1", "Z2"]), filename="十号基坑台账2026.7.31.xlsx")
    body = _upload(client, (name, data)).json()
    assert len(body["groups"]) == 1
    g = body["groups"][0]
    assert g["project_id"] == b
    assert g["needs_eyeball"] is True
    assert g["sheets"][0]["from_filename"] is True
    assert "文件名" in g["sheets"][0]["how"]


def test_unnamed_sheet_in_a_mixed_file_is_not_guessed(client):
    """混着两个基坑的文件里，有一张表自己没写名称 → 不猜，摆出候选让用户手选。"""
    _two_pits(client)
    name, data = _xlsx(
        ("八号竖向位移", "八号基坑", ["JJC1"]),
        ("十号竖向位移", "十号基坑", ["Z1"]),
        ("没写名字的表", None, ["JJC2"]),
    )
    body = _upload(client, (name, data)).json()
    assert len(body["groups"]) == 2  # 两张写了名字的照常分组
    assert len(body["pending"]) == 1
    p = body["pending"][0]
    assert p["sheet"]["sheet_name"] == "没写名字的表"
    assert p["sheet"]["project_id"] is None
    assert "混着好几个基坑" in p["sheet"]["reason"]
    assert {c["project_name"] for c in p["sheet"]["candidates"]} == {"八号基坑", "十号基坑"}


def test_unroutable_sheet_gets_no_batch_until_assigned(client):
    """认不出的表**不建批次**；用户手选项目之后才建，并且能正常导进去。"""
    a, _b = _two_pits(client)
    name, data = _xlsx(("竖向位移", "某个系统里没有的基坑", ["JJC1", "JJC2"]))
    body = _upload(client, (name, data)).json()
    assert body["groups"] == []
    assert len(body["pending"]) == 1
    p = body["pending"][0]

    r = client.post(
        "/api/v1/import/multi/assign",
        json={
            "upload_name": p["upload_name"],
            "sheet_names": ["竖向位移"],
            "project_id": a,
            "filename": p["filename"],
        },
    )
    assert r.status_code == 200, r.text
    g = r.json()
    assert g["project_id"] == a
    assert g["sheets"][0]["how"] == "你手动指定的项目"
    assert {x["point_no"] for x in g["preview"]["valid_rows"]} == {"JJC1", "JJC2"}
    conf = client.post(
        f"/api/v1/import/batches/{g['batch_id']}/confirm", json={"valid_rows": g["preview"]["valid_rows"]}
    )
    assert conf.json()["imported_rows"] == 2


def test_assign_never_reads_outside_the_upload_dir(client):
    """前端回传的暂存名不能当路径用 —— 带目录成分的、指到外面的，一律 404。"""
    a, _b = _two_pits(client)
    for bad in ("../../monitor.db", "..\\..\\monitor.db", "/etc/passwd", "不存在的文件.xlsx"):
        r = client.post(
            "/api/v1/import/multi/assign",
            json={"upload_name": bad, "sheet_names": ["竖向位移"], "project_id": a},
        )
        assert r.status_code == 404, f"{bad} → {r.status_code}"


def test_unnamed_sheets_inherit_the_one_project_of_that_file(client):
    """一份文件里别的表都指向同一个基坑 → 没写名字的表跟着走（比文件名更硬的证据）。"""
    a, _b = _two_pits(client)
    name, data = _xlsx(
        ("八号竖向位移", "八号基坑", ["JJC1"]),
        ("另一张没名字的", None, ["JJC2"]),
    )
    body = _upload(client, (name, data)).json()
    assert body["pending"] == []
    assert len(body["groups"]) == 1
    g = body["groups"][0]
    assert g["project_id"] == a
    by_name = {s["sheet_name"]: s for s in g["sheets"]}
    assert by_name["另一张没名字的"]["from_filename"] is False
    assert "同一份文件" in by_name["另一张没名字的"]["how"]
    assert g["needs_eyeball"] is False


# ─────────────────────── ★ 批次只管自己那几张表 ───────────────────────


def test_recheck_only_touches_the_sheets_of_that_batch(client):
    """改一组的观测日期，**不能**把同一份文件里别的项目的表拉进来。

    这是跨项目导入最贵的一条：`recheck` 是重解析整个文件的，不过滤就会串数据，
    而且全程不报错 —— 用户点确认时别人的读数已经进了这个项目。
    """
    a, b = _two_pits(client)
    name, data = _xlsx(
        ("八号竖向位移", "八号基坑", ["JJC1", "JJC2"]),
        ("十号竖向位移", "十号基坑", ["Z1", "Z2"]),
    )
    groups = {g["project_id"]: g for g in _upload(client, (name, data)).json()["groups"]}
    ga, gb = groups[a], groups[b]

    r = client.post(
        f"/api/v1/import/batches/{ga['batch_id']}/recheck",
        json={"dates": [{"sheet_name": "八号竖向位移", "obs_date": "2026-08-05"}]},
    )
    assert r.status_code == 200, r.text
    prev = r.json()
    assert [s["sheet_name"] for s in prev["sheets"]] == ["八号竖向位移"]
    assert {x["point_no"] for x in prev["valid_rows"]} == {"JJC1", "JJC2"}
    assert {x["obs_date"] for x in prev["valid_rows"]} == {"2026-08-05"}
    # 十号基坑那组一个字都没变
    assert {x["point_no"] for x in gb["preview"]["valid_rows"]} == {"Z1", "Z2"}
    assert {x["obs_date"] for x in gb["preview"]["valid_rows"]} == {"2026-07-31"}


def test_single_project_import_still_rechecks_the_whole_file(client):
    """老路径（项目内 Excel 导入）行为一个字不变：整个文件的表都归它。"""
    a, _b = _two_pits(client)
    name, data = _xlsx(
        ("八号竖向位移", "八号基坑", ["JJC1"]),
        ("另一张", "八号基坑", ["JJC2"]),
    )
    up = client.post(
        f"/api/v1/projects/{a}/import/upload", files={"file": (name, data, _XLSX_MIME)}
    )
    assert up.status_code == 200, up.text
    body = up.json()
    assert len(body["sheets"]) == 2
    rc = client.post(
        f"/api/v1/import/batches/{body['batch_id']}/recheck",
        json={"dates": [{"sheet_name": "八号竖向位移", "obs_date": "2026-08-05"}]},
    )
    assert rc.status_code == 200, rc.text
    assert len(rc.json()["sheets"]) == 2


# ────────────────────────────── 其余口径 ──────────────────────────────


def test_a_broken_file_does_not_block_the_others(client):
    """一次拖十几份是常态：其中一份坏了要**点名报出来**，其余照常分组。"""
    _two_pits(client)
    good = _xlsx(("竖向位移", "八号基坑", ["JJC1"]))
    resp = _upload(client, good, ("说明.txt", b"not an excel", "text/plain"))
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert len(body["groups"]) == 1
    assert len(body["errors"]) == 1
    assert "说明.txt" in body["errors"][0]


def test_point_list_sheets_are_not_routed(client):
    """点位清单（点号+坐标/高程）是「批量建项目」的活，这里只挑出来说明、不导。"""
    _two_pits(client)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("点位清单")
    for i, v in enumerate(["点号", "X", "Y", "Z"], start=1):
        ws.cell(row=1, column=i, value=v)
    ws.cell(row=2, column=1, value="JJC1")
    ws.cell(row=2, column=2, value=3413055.12)
    ws.cell(row=2, column=3, value=412345.67)
    ws.cell(row=2, column=4, value=4.907)
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    body = _upload(client, ("八号基坑点位清单.xlsx", buf.getvalue())).json()
    assert body["groups"] == [] and body["pending"] == []
    assert [p["sheet_name"] for p in body["point_lists"]] == ["点位清单"]


def test_no_projects_yet_says_so_instead_of_500(client):
    """系统里一个项目都没有 → 说人话，别让用户对着 500 猜。"""
    name, data = _xlsx(("竖向位移", "八号基坑", ["JJC1"]))
    r = _upload(client, (name, data))
    assert r.status_code == 422
    assert "还没有项目" in r.json()["detail"]
