# -*- coding: utf-8 -*-
"""Phase 5 导入引擎测试。

测试数据动态生成与用户真实模板（八号基坑）同构的 Excel：
- 双层表头（R5 主表头 + R6 控制值子表头），数据从 R7 起
- R3 含「本次监测时间：YYYY-MM-DD」
- 竖向位移 Sheet：每测点 1 行（component=value）
- 水平位移 Sheet：每测点 2 行（X 行有点号 / Y 行点号空，继承）
- 读值单位 m（解析后 ×1000 → mm）
"""
from datetime import date

import openpyxl
import pytest

from backend.importers.parser import parse_workbook

HEADER_META = ["沉降监测报表", "监测工程名称：8号基坑", "本次监测时间：2026-07-31",
               "仪器型号：Leica TZ05", "监测点号\t初始值（m)\t上次测值（m）\t本次测值（m）\t上次累计变化量（mm)\t本次累计变化量（mm)\t本次变化量（mm)\t变化速率（mm/d）\t控制值\t\t预警等级\t备注"]
HEADER_SUB = ["", "", "", "", "", "", "", "", "累计变化量（mm)", "变化速率（mm/d）", "", ""]


def _make_wb() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---- 竖向位移 Sheet ----
    ws1 = wb.create_sheet("竖向位移")
    meta = ["沉降监测报表", "监测工程名称：8号基坑", "本次监测时间：2026-07-31", "仪器型号：Leica TZ05"]
    header_main = ["监测\n点号", "初始值（m)", "上次测值（m）", "本次测值（m）",
                   "上次累计变化量（mm)", "本次累计变化量（mm)", "本次变化量（mm)", "变化速率（mm/d）",
                   "控制值", None, "预警\n等级", "备注"]
    header_sub = [None] * 8 + ["累计变化量（mm)", "变化速率（mm/d）", None, None]
    for i, v in enumerate(meta, start=1):
        ws1.cell(row=i, column=1, value=v)
    for i, v in enumerate(header_main, start=1):
        ws1.cell(row=5, column=i, value=v)
    for i, v in enumerate(header_sub, start=1):
        ws1.cell(row=6, column=i, value=v)
    # 数据行（初始值期：只有初始值与本次测值；第一期：全列）
    ws1.append([])
    ws1.cell(row=7, column=1, value="JJC1")
    ws1.cell(row=7, column=2, value=4.907)
    ws1.cell(row=7, column=3, value=4.907)
    ws1.cell(row=7, column=4, value=4.906)
    ws1.cell(row=7, column=5, value=0)
    ws1.cell(row=7, column=6, value=-1.0)
    ws1.cell(row=7, column=7, value=-1.0)
    ws1.cell(row=7, column=8, value=-0.5)
    ws1.cell(row=7, column=9, value=30)
    ws1.cell(row=7, column=10, value=10)
    ws1.cell(row=8, column=1, value="JJC2")
    ws1.cell(row=8, column=2, value=4.86)
    ws1.cell(row=8, column=3, value=4.86)
    ws1.cell(row=8, column=4, value=4.859)
    ws1.cell(row=8, column=5, value=0)
    ws1.cell(row=8, column=6, value=-1.0)
    ws1.cell(row=8, column=7, value=-1.0)
    ws1.cell(row=8, column=8, value=-0.5)
    ws1.cell(row=8, column=9, value=30)
    ws1.cell(row=8, column=10, value=10)

    # ---- 水平位移 Sheet（每点 2 行：X 有点号，Y 点号空）----
    ws2 = wb.create_sheet("r.立柱结构水平位移")
    meta2 = ["基坑水平位移监测日报表", "监测工程名称：8号基坑", "本次监测时间：2026-07-31", "仪器型号：Leica TZ05"]
    header2 = ["监测\n点号", "初始值（m)", "上次测值（m)", "本次测值（m)",
               "上次累计变化量（mm)", "本次累计变化量（mm)", "本次变化量（mm)", "变化速率（mm/d）",
               "控制值", None, "预警\n等级", "备注"]
    for i, v in enumerate(meta2, start=1):
        ws2.cell(row=i, column=1, value=v)
    for i, v in enumerate(header2, start=1):
        ws2.cell(row=5, column=i, value=v)
    for i, v in enumerate(header_sub, start=1):
        ws2.cell(row=6, column=i, value=v)
    # JJC1 X 行
    ws2.cell(row=7, column=1, value="JJC1")
    ws2.cell(row=7, column=2, value=3413055.12)
    ws2.cell(row=7, column=3, value=3413055.12)
    ws2.cell(row=7, column=4, value=3413055.117)
    ws2.cell(row=7, column=5, value=0)
    ws2.cell(row=7, column=6, value=-3.0)
    ws2.cell(row=7, column=7, value=-3.0)
    ws2.cell(row=7, column=8, value=-1.5)
    ws2.cell(row=7, column=9, value=20)
    ws2.cell(row=7, column=10, value=3)
    # JJC1 Y 行（点号空）
    ws2.cell(row=8, column=2, value=510174.166)
    ws2.cell(row=8, column=3, value=510174.166)
    ws2.cell(row=8, column=4, value=510174.159)
    ws2.cell(row=8, column=5, value=0)
    ws2.cell(row=8, column=6, value=-7.0)
    ws2.cell(row=8, column=7, value=-7.0)
    ws2.cell(row=8, column=8, value=-3.5)
    ws2.cell(row=8, column=9, value=20)
    ws2.cell(row=8, column=10, value=3)
    # JJC2 X/Y
    ws2.cell(row=9, column=1, value="JJC2")
    ws2.cell(row=9, column=2, value=3413155.12)
    ws2.cell(row=9, column=4, value=3413155.10)
    ws2.cell(row=9, column=6, value=-20.0)
    ws2.cell(row=10, column=2, value=510284.166)
    ws2.cell(row=10, column=4, value=510284.16)

    # ---- 非数据 Sheet（应被跳过）----
    ws3 = wb.create_sheet("封面")
    ws3.cell(row=1, column=1, value="基坑监测日报")
    ws4 = wb.create_sheet("T.立柱结构水平位移")  # 测斜历史，表头不匹配 → 跳过
    ws4.cell(row=1, column=1, value="深度(m)")
    ws4.cell(row=2, column=1, value="0.5")
    return wb


@pytest.fixture
def template_xlsx(tmp_path):
    wb = _make_wb()
    p = tmp_path / "八号基坑（第一期）2026.7.31.xlsx"
    wb.save(p)
    wb.close()
    return p


def test_parse_sheet_identification(template_xlsx):
    """Sheet 识别：只解析 2 个数据 Sheet，跳过封面。"""
    sheets = parse_workbook(template_xlsx)
    assert [s.sheet_name for s in sheets] == ["竖向位移", "r.立柱结构水平位移"]


def test_parse_obs_date(template_xlsx):
    """日期提取：R3「本次监测时间：2026-07-31」。"""
    sheets = parse_workbook(template_xlsx)
    for s in sheets:
        assert s.obs_date == date(2026, 7, 31)


def test_parse_settlement_rows(template_xlsx):
    """沉降表：每测点 1 行，component=value，单位 m→mm（×1000）。"""
    s = parse_workbook(template_xlsx)[0]
    assert s.series_type == "settlement"
    assert s.total_rows == 2
    r1 = s.rows[0]
    assert r1.point_no == "JJC1"
    assert r1.component == "value"
    assert r1.current_mm == pytest.approx(4906.0)      # 4.906 m × 1000
    assert r1.initial_mm == pytest.approx(4907.0)
    assert r1.control_cum_mm == 30
    assert r1.control_rate_mm_d == 10  # 「控制值」合并组第二列：变化速率控制值
    assert r1.remark is None


def test_parse_displacement_xy_rows(template_xlsx):
    """位移表：X 行有点号、Y 行点号空（继承），坐标 ×1000。"""
    s = parse_workbook(template_xlsx)[1]
    assert s.series_type == "displacement"
    rows = s.rows
    assert len(rows) == 4  # JJC1 X/Y + JJC2 X/Y
    # JJC1 X
    assert rows[0].point_no == "JJC1" and rows[0].component == "X"
    assert rows[0].current_mm == pytest.approx(3413055.117 * 1000, rel=1e-9)
    assert rows[0].control_cum_mm == 20
    assert rows[0].control_rate_mm_d == 3  # 位移：累计控制值 20 / 速率控制值 3
    # JJC1 Y（点号继承）
    assert rows[1].point_no == "JJC1" and rows[1].component == "Y"
    assert rows[1].current_mm == pytest.approx(510174.159 * 1000, rel=1e-9)
    # JJC2 X（无上次测值 → 本次值仍在）
    assert rows[2].point_no == "JJC2" and rows[2].component == "X"
    assert rows[2].current_mm == pytest.approx(3413155.10 * 1000, rel=1e-9)


def test_parse_missing_value(template_xlsx):
    """缺测：值为空/'-' → current_mm=None。"""
    wb = _make_wb()
    ws = wb["竖向位移"]
    ws.cell(row=9, column=1, value="JJC3")
    ws.cell(row=9, column=2, value=5.0)
    ws.cell(row=9, column=4, value="-")   # 本次测值缺测
    p = template_xlsx.parent / "missing.xlsx"
    wb.save(p)
    wb.close()
    s = parse_workbook(p)[0]
    r3 = s.rows[2]
    assert r3.point_no == "JJC3"
    assert r3.current_mm is None


def test_skip_non_data_sheets(template_xlsx):
    """非数据 Sheet（封面）被跳过。"""
    sheets = parse_workbook(template_xlsx)
    names = [s.sheet_name for s in sheets]
    assert "封面" not in names
    assert "T.立柱结构水平位移" not in names  # 表头不含点号/测值


def test_v0_period_sheet(tmp_path):
    """初始值期：整 Sheet 无本次测值但有初始值 → 初始值即该期读数（V0 基准）。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("竖向位移")
    for i, v in enumerate(["沉降监测报表", "监测工程名称：8号基坑", "本次监测时间：2026-07-30", "仪器型号：Leica TZ05"], start=1):
        ws.cell(row=i, column=1, value=v)
    header_main = ["监测\n点号", "初始值（m)", "上次测值（m）", "本次测值（m）",
                   "上次累计变化量（mm)", "本次累计变化量（mm)", "本次变化量（mm)", "变化速率（mm/d）",
                   "控制值", None, "预警\n等级", "备注"]
    for i, v in enumerate(header_main, start=1):
        ws.cell(row=5, column=i, value=v)
    # 初始值期：只填初始值列（4.907），本次测值留空
    ws.cell(row=7, column=1, value="JJC1")
    ws.cell(row=7, column=2, value=4.907)
    ws.cell(row=8, column=1, value="JJC2")
    ws.cell(row=8, column=2, value=4.86)
    p = tmp_path / "初始值.xlsx"
    wb.save(p)
    wb.close()

    sheets = parse_workbook(p)
    assert len(sheets) == 1
    s = sheets[0]
    assert s.obs_date == date(2026, 7, 30)
    assert s.total_rows == 2
    # V0 期：current_mm 取自初始值 ×1000
    assert s.rows[0].point_no == "JJC1"
    assert s.rows[0].component == "value"
    assert s.rows[0].current_mm == pytest.approx(4907.0)
    assert s.rows[1].current_mm == pytest.approx(4860.0)


def test_v0_period_displacement_xy(tmp_path):
    """初始值期位移表：X/Y 两行都只有初始值，均落库为 V0 读数。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("r.立柱结构水平位移")
    for i, v in enumerate(["基坑水平位移监测日报表", "监测工程名称：8号基坑", "本次监测时间：2026-07-30", "仪器型号：Leica TZ05"], start=1):
        ws.cell(row=i, column=1, value=v)
    header2 = ["监测\n点号", "初始值（m)", "上次测值（m)", "本次测值（m)",
               "上次累计变化量（mm)", "本次累计变化量（mm)", "本次变化量（mm)", "变化速率（mm/d）",
               "控制值", None, "预警\n等级", "备注"]
    for i, v in enumerate(header2, start=1):
        ws.cell(row=5, column=i, value=v)
    # JJC1 X 行（点号）+ Y 行（点号空，仅初始值）
    ws.cell(row=7, column=1, value="JJC1")
    ws.cell(row=7, column=2, value=3413055.12)
    ws.cell(row=8, column=2, value=510174.166)
    p = tmp_path / "位移初始值.xlsx"
    wb.save(p)
    wb.close()

    sheets = parse_workbook(p)
    s = sheets[0]
    assert s.total_rows == 2
    assert s.rows[0].point_no == "JJC1" and s.rows[0].component == "X"
    assert s.rows[0].current_mm == pytest.approx(3413055.12 * 1000, rel=1e-9)
    assert s.rows[1].point_no == "JJC1" and s.rows[1].component == "Y"
    assert s.rows[1].current_mm == pytest.approx(510174.166 * 1000, rel=1e-9)


def _make_alt_wb(unit: str = "m") -> openpyxl.Workbook:
    """构造"另一种格式"台账：表头行位置不同、列顺序打乱、列名用同义写法、日期写法不同。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("沉降观测表")
    ws.cell(row=1, column=1, value="某某工程沉降观测记录")
    ws.cell(row=2, column=1, value="工程名称：测试工程")
    ws.cell(row=3, column=1, value="观测日期：2026/8/15")  # 非"本次监测时间"写法
    # 表头在第 6 行（原模板第 5 行），列顺序打乱、用词不同
    hdr = {
        2: "测点编号",
        5: f"本次读数（{unit}）",
        1: "备注",
        4: f"初值（{unit}）",
        3: f"前次测值（{unit}）",
        8: "累计变化（mm）",
        9: "变化量（mm）",
        10: "速率（mm/d）",
        11: "控制值（mm）",
    }
    for col, text in hdr.items():
        ws.cell(row=6, column=col, value=text)
    ws.cell(row=7, column=2, value="LZ-01")
    ws.cell(row=7, column=4, value=4.907 if unit == "m" else 4907)
    ws.cell(row=7, column=5, value=4.906 if unit == "m" else 4906)
    ws.cell(row=7, column=8, value=-1.0)
    ws.cell(row=7, column=9, value=-1.0)
    ws.cell(row=7, column=10, value=-0.5)
    ws.cell(row=7, column=11, value=30)
    return wb


def test_parse_alt_template_layout_and_wording(tmp_path):
    """换模板容忍度：表头行位置变化 + 列顺序打乱 + 列名同义 + 日期写法不同 → 仍能识别。"""
    p = tmp_path / "alt.xlsx"
    wb = _make_alt_wb("m")
    wb.save(p)
    wb.close()

    sheets = parse_workbook(p)
    assert len(sheets) == 1
    s = sheets[0]
    assert s.series_type == "settlement"
    assert s.obs_date == date(2026, 8, 15)
    assert s.total_rows == 1
    r = s.rows[0]
    assert r.point_no == "LZ-01"
    assert r.component == "value"
    assert r.current_mm == pytest.approx(4906.0)   # 4.906 m × 1000
    assert r.initial_mm == pytest.approx(4907.0)
    assert r.control_cum_mm == 30


def test_parse_read_values_in_mm_not_scaled(tmp_path):
    """读值列标注 (mm) 时按毫米处理、不再 ×1000（防换模板造成 1000 倍误差）。"""
    p = tmp_path / "alt_mm.xlsx"
    wb = _make_alt_wb("mm")
    wb.save(p)
    wb.close()

    sheets = parse_workbook(p)
    r = sheets[0].rows[0]
    assert r.current_mm == pytest.approx(4906.0)
    assert r.initial_mm == pytest.approx(4907.0)


def test_is_valid_point_no_rejects_numeric():
    """纯数值（-1/0/2/2026/3.5）不是测点号；带字母前缀的编号才是。"""
    from backend.importers.parser import is_valid_point_no

    assert is_valid_point_no("JJC1")
    assert is_valid_point_no("JJC1-1")
    assert is_valid_point_no("QC01")
    for bad in ("-1", "0", "2", "2026", "3.5", "+1", ""):
        assert not is_valid_point_no(bad), bad


def test_parse_ignores_numeric_tail_table(tmp_path):
    """真实台账常在数据表下方附一小块数值表（差值/坐标），不得被识别为测点。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "竖向位移"
    ws.cell(row=1, column=1, value="沉降监测报表")
    ws.cell(row=3, column=1, value="本次监测时间：2026-08-01")
    for i, v in enumerate(
        ["监测点号", "初始值（m)", "上次测值（m）", "本次测值（m）",
         "上次累计变化量（mm)", "本次累计变化量（mm)", "本次变化量（mm）", "变化速率（mm/d）"],
        start=1,
    ):
        ws.cell(row=5, column=i, value=v)
    ws.cell(row=6, column=1, value="JJC1")
    ws.cell(row=6, column=2, value=4.907)
    ws.cell(row=6, column=4, value=4.906)
    ws.cell(row=7, column=1, value="JJC2")
    ws.cell(row=7, column=2, value=4.860)
    ws.cell(row=7, column=4, value=4.859)
    # 说明行 + 结论行
    ws.cell(row=9, column=1, value="变化量正负号说明：-1为X正数朝左线")
    ws.cell(row=11, column=1, value="监测结论及建议：")
    # 表格下方的数值小表 —— 这些不是点号
    ws.cell(row=14, column=1, value=-1)
    ws.cell(row=14, column=2, value=3413069.85)
    ws.cell(row=15, column=1, value=0)
    ws.cell(row=15, column=2, value=3413068.55)
    ws.cell(row=16, column=1, value=2)
    ws.cell(row=16, column=2, value=3413073.11)
    p = tmp_path / "tail.xlsx"
    wb.save(p)
    wb.close()

    sheets = parse_workbook(p)
    assert len(sheets) == 1
    s = sheets[0]
    assert s.total_rows == 2
    assert [r.point_no for r in s.rows] == ["JJC1", "JJC2"]


def test_parse_simple_xyz_sheet(tmp_path):
    """「名字 + X/Y/Z」简易初始值清单：Z→沉降、X/Y→位移，无日期（前端补填）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["名字", "Y", "X", "Z"])
    ws.append(["JJC1", 510199.345, 3413030.175, 4.411])
    ws.append(["JJC2", 510200.279, 3413033.238, 4.427])
    p = tmp_path / "xyz.xlsx"
    wb.save(p)
    wb.close()

    sheets = parse_workbook(p)
    assert len(sheets) == 1
    s = sheets[0]
    assert s.series_type == "both"
    assert s.obs_date is None  # 简易表没有日期 → 走"补填日期"流程
    assert s.total_rows == 6   # 2 点 × (X/Y/Z)
    m = {(r.point_no, r.series_type, r.component): r.current_mm for r in s.rows}
    assert m[("JJC1", "settlement", "value")] == pytest.approx(4411.0)
    assert m[("JJC1", "displacement", "X")] == pytest.approx(3413030175.0)
    assert m[("JJC1", "displacement", "Y")] == pytest.approx(510199345.0)


def test_is_legacy_xls_detection(tmp_path):
    """老版 .xls（OLE2 头）能识别；普通 xlsx 不误判。"""
    from backend.importers.parser import is_legacy_xls

    xls = tmp_path / "old.xls"
    xls.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    assert is_legacy_xls(xls)

    wb = openpyxl.Workbook()
    xlsx = tmp_path / "new.xlsx"
    wb.save(xlsx)
    wb.close()
    assert not is_legacy_xls(xlsx)


# ---------------------------------------------------------------------------
# 自适应解析增强（2026-09-23）：控制值双列 / 表头行放宽 / 中文点号 / 裸日期单元格
# ---------------------------------------------------------------------------


def test_control_rate_two_level_header_with_merge(tmp_path):
    """「控制值」为合并表头时，第二层 累计控制值 / 速率控制值 两列都要取到。

    真实台账：R5 I=「控制值」(合并 I:J)、R6 I=累计变化量 J=变化速率、数据 I=30 J=10。
    此前只取到 I 列，速率控制值（沉降 10 / 位移 3 mm·d⁻¹）被丢弃。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "竖向位移"
    ws.cell(row=1, column=1, value="本次监测时间：2026-09-20")
    for i, v in enumerate(["监测点号", "初始值（m）", "本次测值（m）", "控制值", None, "预警等级", "备注"], start=1):
        if v:
            ws.cell(row=3, column=i, value=v)
    ws.merge_cells(start_row=3, start_column=4, end_row=3, end_column=5)
    ws.cell(row=4, column=4, value="累计变化量（mm）")
    ws.cell(row=4, column=5, value="变化速率（mm/d）")
    ws.cell(row=5, column=1, value="JJC1")
    ws.cell(row=5, column=2, value=3.184)
    ws.cell(row=5, column=3, value=3.189)
    ws.cell(row=5, column=4, value=30)
    ws.cell(row=5, column=5, value=10)
    p = tmp_path / "ctrl_merge.xlsx"
    wb.save(p)
    wb.close()

    s = parse_workbook(p)[0]
    assert s.column_map["control_cum_mm"] == 4
    assert s.column_map["control_rate_mm_d"] == 5
    assert s.rows[0].control_cum_mm == 30
    assert s.rows[0].control_rate_mm_d == 10


def test_single_level_header_does_not_invent_rate_column(tmp_path):
    """单层表头（控制值只有一列）：不得凭空造出速率控制值，也不能把首行数据当子表头。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "竖向位移"
    ws.cell(row=1, column=1, value="本次监测时间：2026-09-20")
    for i, v in enumerate(["监测点号", "初始值（m）", "本次测值（m）", "控制值"], start=1):
        ws.cell(row=3, column=i, value=v)
    ws.cell(row=4, column=1, value="JJC1")
    ws.cell(row=4, column=2, value=3.184)
    ws.cell(row=4, column=3, value=3.189)
    ws.cell(row=4, column=4, value=30)
    p = tmp_path / "ctrl_single.xlsx"
    wb.save(p)
    wb.close()

    s = parse_workbook(p)[0]
    assert s.column_map["control_cum_mm"] == 4
    assert s.column_map.get("control_rate_mm_d") is None
    assert s.rows[0].control_cum_mm == 30
    assert s.rows[0].control_rate_mm_d is None
    assert s.total_rows == 1


def test_header_found_beyond_15_rows(tmp_path):
    """表头在第 16 行（上方是标题与说明块）也能识别——扫描范围放宽到 30 行。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "竖向位移"
    ws.cell(row=2, column=1, value="某某工程 基坑监测日报")
    ws.cell(row=3, column=4, value="本次监测时间：2026-09-20")
    for r in range(4, 16):
        ws.cell(row=r, column=1, value=f"说明文字第{r}行")
    for i, v in enumerate(["监测点号", "初始值（m）", "本次测值（m）"], start=1):
        ws.cell(row=16, column=i, value=v)
    ws.cell(row=17, column=1, value="JJC1")
    ws.cell(row=17, column=2, value=4.411)
    ws.cell(row=17, column=3, value=4.410)
    p = tmp_path / "header16.xlsx"
    wb.save(p)
    wb.close()

    sheets = parse_workbook(p)
    assert len(sheets) == 1
    assert sheets[0].obs_date == date(2026, 9, 20)
    assert [r.point_no for r in sheets[0].rows] == ["JJC1"]


def test_chinese_point_no_accepted_and_footer_rejected(tmp_path):
    """中文点号（左1号点）按"表内编号形状一致"接纳；页脚/说明文字不得被当成测点。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "竖向位移"
    ws.cell(row=1, column=1, value="本次监测时间：2026-09-20")
    for i, v in enumerate(["监测点号", "初始值（m）", "本次测值（m）"], start=1):
        ws.cell(row=3, column=i, value=v)
    ws.cell(row=4, column=1, value="左1号点")
    ws.cell(row=4, column=2, value=4.411)
    ws.cell(row=4, column=3, value=4.410)
    ws.cell(row=5, column=1, value="左2号点")
    ws.cell(row=5, column=2, value=4.427)
    ws.cell(row=5, column=3, value=4.426)
    ws.cell(row=6, column=1, value="共 2 个测点")
    ws.cell(row=7, column=1, value="监测结论及建议")
    p = tmp_path / "cn_point.xlsx"
    wb.save(p)
    wb.close()

    s = parse_workbook(p)[0]
    assert [r.point_no for r in s.rows] == ["左1号点", "左2号点"]


def test_bare_date_cell_becomes_obs_date(tmp_path):
    """模板把日期做成日期格式且没写「日期/时间」标签时，也能取到观测日期。"""
    import datetime as _dt

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "竖向位移"
    ws.cell(row=1, column=1, value="工程名称：测试")
    ws.cell(row=2, column=1, value=_dt.datetime(2026, 9, 20))
    ws.cell(row=2, column=5, value="检定日期")
    ws.cell(row=2, column=6, value=_dt.datetime(2025, 10, 20))  # 仪器检定日期 ≠ 观测日期
    for i, v in enumerate(["监测点号", "初始值（m）", "本次测值（m）"], start=1):
        ws.cell(row=4, column=i, value=v)
    ws.cell(row=5, column=1, value="JJC1")
    ws.cell(row=5, column=2, value=4.411)
    ws.cell(row=5, column=3, value=4.410)
    p = tmp_path / "bare_date.xlsx"
    wb.save(p)
    wb.close()

    s = parse_workbook(p)[0]
    assert s.obs_date == date(2026, 9, 20)


def test_paragraph_sheet_not_parsed_as_data(tmp_path):
    """说明页里出现「点号/测值/沉降」整段文字，不得被当成数据表。"""
    wb = openpyxl.Workbook()
    ws0 = wb.active
    ws0.title = "概况"
    ws0.cell(row=1, column=1, value="1、工程概况")
    ws0.cell(row=8, column=1, value="本工程监测点号共 8 个，本次沉降测值与累计变化量详见竖向位移报表。")
    ws0.cell(row=20, column=1, value="监测点号与本次测值汇总：见附表")

    ws1 = wb.create_sheet("竖向位移")
    ws1.cell(row=1, column=1, value="本次监测时间：2026-09-20")
    for i, v in enumerate(["监测点号", "初始值（m）", "本次测值（m）"], start=1):
        ws1.cell(row=3, column=i, value=v)
    ws1.cell(row=4, column=1, value="JJC1")
    ws1.cell(row=4, column=2, value=4.411)
    ws1.cell(row=4, column=3, value=4.410)
    p = tmp_path / "paragraph.xlsx"
    wb.save(p)
    wb.close()

    sheets = parse_workbook(p)
    assert [s.sheet_name for s in sheets] == ["竖向位移"]


def test_sheet1_with_cjk_points_infers_type_from_title(tmp_path):
    """Sheet 名是 Sheet1、表头也没有「沉降/位移」字样：靠标题行「沉降监测报表」定类型。

    中文点号（2号右边蹲观测点1 / 8号蹲观测点1）按"该行有读数"接纳；
    表格下方的说明行（含冒号、无读数）必须被排除。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(row=1, column=1, value="沉降监测报表")
    ws.cell(row=2, column=1, value="监测工程名称：路面监测")
    ws.cell(row=3, column=1, value="本次监测时间：2026-07-31")
    headers = ["监测点号", "初始值（m)", "上次测值（m）", "本次测值（m）",
               "上次累计变化量（mm)", "本次累计变化量（mm)", "本次变化量（mm）",
               "变化速率（mm/d）", "控制值", None]
    for i, v in enumerate(headers, start=1):
        if v:
            ws.cell(row=5, column=i, value=v)
    ws.cell(row=6, column=9, value="累计变化量（mm)")
    ws.cell(row=6, column=10, value="变化速率（mm/d）")

    ws.cell(row=7, column=1, value="2号右边蹲观测点1")
    ws.cell(row=7, column=2, value=3.433)
    ws.cell(row=7, column=4, value=3.432)
    ws.cell(row=7, column=9, value=30)
    ws.cell(row=7, column=10, value=10)
    ws.cell(row=8, column=1, value="8号蹲观测点1")
    ws.cell(row=8, column=2, value=4.846)
    ws.cell(row=8, column=4, value=4.847)
    ws.cell(row=8, column=9, value=30)
    ws.cell(row=8, column=10, value=10)
    ws.cell(row=9, column=1, value='变化量正负号说明：\n"-"为上升"+"为下沉')
    ws.cell(row=10, column=1, value="监测结论及建议：\n本次监测工作为涉路基坑路面竖向位移初始值监测工作。")
    ws.cell(row=11, column=1, value="现场监测人：")
    p = tmp_path / "sheet1_cjk.xlsx"
    wb.save(p)
    wb.close()

    sheets = parse_workbook(p)
    assert len(sheets) == 1
    s = sheets[0]
    assert s.series_type == "settlement"          # 由标题行判定
    assert s.obs_date == date(2026, 7, 31)
    assert [r.point_no for r in s.rows] == ["2号右边蹲观测点1", "8号蹲观测点1"]
    assert s.rows[0].current_mm == pytest.approx(3432.0)
    assert s.rows[0].control_cum_mm == 30
    assert s.rows[0].control_rate_mm_d == 10


def test_headerless_name_xyz_list(tmp_path):
    """无表头的「名字 + 坐标 + 高程」仪器导出清单：按数值量级定 X/Y，小数列当高程。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet0"
    rows = [
        ("6号右左边基坑", 3412997.533, 510197.577, 4.619),
        ("6号右左边基坑1", 3412997.578, 510193.872, 4.731),
        ("6号右左边基坑2", 3413000.670, 510192.019, 4.795),
    ]
    for r, (name, x, y, z) in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=x)
        ws.cell(row=r, column=3, value=y)
        ws.cell(row=r, column=4, value=z)
    p = tmp_path / "headerless.xlsx"
    wb.save(p)
    wb.close()

    sheets = parse_workbook(p)
    assert len(sheets) == 1
    s = sheets[0]
    assert s.series_type == "both"
    assert s.obs_date is None
    assert s.column_map == {"point_no": 1, "x": 2, "y": 3, "z": 4}
    assert s.total_rows == 9  # 3 点 × (X / Y / Z)
    m = {(r.point_no, r.series_type, r.component): r.current_mm for r in s.rows}
    assert m[("6号右左边基坑", "displacement", "X")] == pytest.approx(3412997.533 * 1000)
    assert m[("6号右左边基坑", "displacement", "Y")] == pytest.approx(510197.577 * 1000)
    assert m[("6号右左边基坑", "settlement", "value")] == pytest.approx(4619.0)


def test_headerless_list_does_not_misfire_on_text_sheets(tmp_path):
    """普通文本页 / 纯数值列不得被误判成"无表头清单"。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "说明"
    ws.cell(row=1, column=1, value="1、工程概况")
    ws.cell(row=2, column=1, value="本工程监测点号共 8 个")
    ws.cell(row=3, column=1, value=3.431)  # 纯数值列，不是点号
    p = tmp_path / "notlist.xlsx"
    wb.save(p)
    wb.close()

    assert parse_workbook(p) == []


def test_minimal_two_column_table_is_accepted(tmp_path):
    """最小可用台账「点号 + 本次测值」两列也要能导入（新防线不得过严）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "竖向位移"
    ws.cell(row=1, column=1, value="本次监测时间：2026-09-20")
    ws.cell(row=3, column=1, value="监测点号")
    ws.cell(row=3, column=2, value="本次测值（m）")
    ws.cell(row=4, column=1, value="JJC1")
    ws.cell(row=4, column=2, value=4.410)
    p = tmp_path / "two_col.xlsx"
    wb.save(p)
    wb.close()

    sheets = parse_workbook(p)
    assert len(sheets) == 1
    assert sheets[0].total_rows == 1
    assert sheets[0].rows[0].point_no == "JJC1"
    assert sheets[0].rows[0].current_mm == pytest.approx(4410.0)


def _find_legacy_sample(need_parsable: bool = False):
    """在样本库里找一个老版 .xls（OLE2）。need_parsable=True 时要求它能解析出数据表。"""
    import glob

    from backend.importers.parser import is_legacy_xls

    for p in sorted(glob.glob("database/uploads/*.xls*")):
        if not is_legacy_xls(p):
            continue
        if not need_parsable:
            return p
        try:
            if parse_workbook(p):
                return p
        except Exception:  # noqa: BLE001
            continue
    return None


def test_xls_conversion_writes_to_system_temp_not_source_dir(tmp_path):
    """老版 .xls 的转换临时文件必须落在系统临时目录，不得写进源文件所在目录。

    否则进程中断会在 uploads/ 里留下 *.converted.xlsx，既不干净，也可能被后续扫描
    误当成用户上传的文件。
    """
    import glob
    import os
    import tempfile
    from pathlib import Path

    from backend.importers.parser import xls_to_xlsx

    src = _find_legacy_sample()
    if src is None:
        pytest.skip("样本库里没有老版 .xls 文件")

    before = set(glob.glob(os.path.join(os.path.dirname(src), "*.converted.xlsx")))
    out = xls_to_xlsx(src)
    try:
        assert out.exists()
        assert str(out.parent.resolve()) == str(Path(tempfile.gettempdir()).resolve())
        after = set(glob.glob(os.path.join(os.path.dirname(src), "*.converted.xlsx")))
        assert after == before  # 源目录没多出临时文件
    finally:
        out.unlink(missing_ok=True)


def test_legacy_xls_file_parses_end_to_end():
    """老版 .xls（扩展名写着 .xlsx 的那种）能一路解析出数据表。"""
    src = _find_legacy_sample(need_parsable=True)
    if src is None:
        pytest.skip("样本库里没有能解析出数据表的老版 .xls 文件")

    sheets = parse_workbook(src)
    assert sheets
    assert all(s.total_rows > 0 for s in sheets)
