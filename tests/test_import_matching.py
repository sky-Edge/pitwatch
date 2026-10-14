# -*- coding: utf-8 -*-
"""测试：导入测点匹配/自动建点（需求②）。"""
import io

from openpyxl import Workbook

from backend.importers.point_matcher import normalize_point_no, suggest_match


def _make_xlsx(rows) -> bytes:
    """生成模拟模板 xlsx（单 Sheet：竖向位移，双层表头，数据从 R7）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "竖向位移"
    ws["A1"] = "沉降监测报表"
    ws["A2"] = "工程"
    ws["A3"] = "本次监测时间：2026-08-01"
    ws["A5"] = "监测\n点号"
    ws["B5"] = "初始值（m)"
    ws["D5"] = "本次测值（m）"
    ws["A6"] = None
    ws["B6"] = None
    ws["D6"] = None
    for i, (no, init_m, cur_m) in enumerate(rows):
        r = 7 + i
        ws.cell(row=r, column=1, value=no)
        ws.cell(row=r, column=2, value=init_m)
        ws.cell(row=r, column=4, value=cur_m)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client, pid, xlsx_bytes, fname="test.xlsx"):
    return client.post(
        f"/api/v1/projects/{pid}/import/upload",
        files={"file": (fname, xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


# ---------- matcher 单测 ----------

def test_normalize():
    assert normalize_point_no("jjc-01") == "JJC01"
    assert normalize_point_no("JJC1") == "JJC1"
    assert normalize_point_no("JJC 01") == "JJC01"
    assert normalize_point_no("1＃") == "1"  # 全角 # → NFKC 转 # → 剥除
    assert normalize_point_no("ＪＪＣ-１") == "JJC1"  # 全角字符


def test_suggest_match_exact_and_strong():
    existing = ["JJC1", "JJC2", "LZ-01", "A10"]
    # 规范化后完全相等 → exact
    assert suggest_match("jjc1", existing)[0]["confidence"] == "exact"
    # 前缀相同、数字尾数值相等 → strong
    hit = suggest_match("JJC-01", existing)
    assert any(h["point_no"] == "JJC1" for h in hit)
    # 数字尾同但前缀不同 → weak（排在最后）
    hit = suggest_match("1#", ["A1", "JJC1"])
    assert hit[0]["point_no"] in ("A1", "JJC1")


# ---------- 端到端：预览 point_issues ----------

def test_upload_unknown_points_reported(client):
    """上传含未知点号的初始 Excel → 预览返回 point_issues + 推断类型。"""
    p = client.post("/api/v1/projects", json={"code": "PRJ001", "name": "导入匹配测试"}).json()
    pid = p["id"]
    # 项目里先建 JJC1（沉降+位移 both）
    client.post(f"/api/v1/projects/{pid}/points", json={"point_no": "JJC1", "point_type": "both"}).json()

    data = _make_xlsx([("JJC1", 4.907, None), ("JJC-2", 5.2, None), ("LZ-9", 4.4, None)])
    resp = _upload(client, pid, data)
    assert resp.status_code == 200
    d = resp.json()
    assert len(d["error_rows"]) == 2  # JJC-2 与 LZ-9
    # JJC-2 有 suggested 候选 JJC2? 不存在 → 候选为空；LZ-9 无候选
    issues = {i["point_no"]: i for i in d["point_issues"]}
    assert "JJC-2" in issues
    assert "LZ-9" in issues
    assert issues["LZ-9"]["infer_point_type"] == "settlement"  # 只有沉降列


def test_confirm_auto_create_points(client):
    """确认导入时勾选自动建点 → 建点成功且数据入库。"""
    p = client.post("/api/v1/projects", json={"code": "PRJ001", "name": "导入匹配测试2"}).json()
    pid = p["id"]
    data = _make_xlsx([("LZ-9", 4.4, 4.401)])
    resp = _upload(client, pid, data)
    d = resp.json()
    bid = d["batch_id"]
    # 把该行作为 valid 回传 + auto_create LZ-9
    row = {k: v for k, v in d["error_rows"][0].items() if k != "error" and k != "suggested"}
    row.pop("error_kind", None)
    r = client.post(
        f"/api/v1/import/batches/{bid}/confirm",
        json={
            "valid_rows": [row],
            "point_actions": [{"excel_no": "LZ-9", "auto_create": True, "infer_point_type": "settlement", "map_to": None}],
        },
    )
    assert r.status_code == 200
    body = r.json()
    assert body["imported_rows"] == 1
    # 测点已建且类型 settlement
    pts = client.get(f"/api/v1/projects/{pid}/points").json()
    lz = next(x for x in pts if x["point_no"] == "LZ-9")
    assert lz["point_type"] == "settlement"


def test_confirm_map_to_existing(client):
    """确认导入时映射到已有测点 → 数据落到已有测点。"""
    p = client.post("/api/v1/projects", json={"code": "PRJ001", "name": "导入匹配测试3"}).json()
    pid = p["id"]
    # 项目已有 JJC1
    pt = client.post(f"/api/v1/projects/{pid}/points", json={"point_no": "JJC1", "point_type": "settlement"}).json()
    data = _make_xlsx([("JJC-01", 4.907, 4.906)])
    d = _upload(client, pid, data).json()
    bid = d["batch_id"]
    issue = d["point_issues"][0]
    assert issue["point_no"] == "JJC-01"
    # JJC-01 与 JJC1 数字尾相同 → 应给 strong 候选
    assert any(s["point_no"] == "JJC1" for s in issue["suggested"])

    row = {k: v for k, v in d["error_rows"][0].items() if k != "error" and k != "suggested"}
    row.pop("error_kind", None)
    row["point_no"] = "JJC1"  # 前端重写为映射目标
    r = client.post(
        f"/api/v1/import/batches/{bid}/confirm",
        json={
            "valid_rows": [row],
            "point_actions": [{"excel_no": "JJC-01", "auto_create": False, "infer_point_type": "settlement", "map_to": "JJC1"}],
        },
    )
    assert r.status_code == 200
    assert r.json()["imported_rows"] == 1
    recs = client.get(f"/api/v1/points/{pt['id']}/records", params={"from_date": "2026-08-01", "to_date": "2026-08-01"}).json()
    assert len(recs) == 1
    assert abs(recs[0]["value_mm"] - 4906.0) < 1e-6


def test_confirm_row_without_obs_date_skipped(client):
    """确认导入：valid_rows 中含 obs_date=None 的行（如初始值/无日期 Sheet 的未知点行）
    不应让 confirm 崩溃——直接跳过该行，其余正常导入。"""
    p = client.post("/api/v1/projects", json={"code": "PRJ001", "name": "无日期行防御"}).json()
    pid = p["id"]
    data = _make_xlsx([("LZ-9", 4.4, 4.401)])
    resp = _upload(client, pid, data)
    assert resp.status_code == 200
    d = resp.json()
    bid = d["batch_id"]
    row = {k: v for k, v in d["error_rows"][0].items() if k != "error" and k != "suggested"}
    row.pop("error_kind", None)
    # 构造：obs_date 置 None（模拟前端把无法识别日期的行回传）
    row["obs_date"] = None
    r = client.post(
        f"/api/v1/import/batches/{bid}/confirm",
        json={
            "valid_rows": [row],
            "point_actions": [{"excel_no": "LZ-9", "auto_create": True, "infer_point_type": "settlement", "map_to": None}],
        },
    )
    assert r.status_code == 200
    assert r.json()["imported_rows"] == 0  # 无日期行被跳过，不崩溃


def _make_xlsx_no_date(rows) -> bytes:
    """生成无"本次监测时间"表头的模板 xlsx（模拟用户表头漏填日期）。"""
    from io import BytesIO
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "竖向位移"
    ws["A1"] = "沉降监测报表"
    ws["A2"] = "工程"
    # 故意不写 A3"本次监测时间"
    ws["A5"] = "监测\n点号"
    ws["B5"] = "初始值（m)"
    ws["D5"] = "本次测值（m）"
    for i, (no, init_m, cur_m) in enumerate(rows):
        r = 7 + i
        ws.cell(row=r, column=1, value=no)
        ws.cell(row=r, column=2, value=init_m)
        ws.cell(row=r, column=4, value=cur_m)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_upload_no_date_reported_not_rejected(client):
    """上传无观测日期的 Excel → 行归入 no_date_rows（不再整批拒绝），
    前端补日期后确认可导入；不补日期则不导入。"""
    p = client.post("/api/v1/projects", json={"code": "PRJ001", "name": "无日期导入"}).json()
    pid = p["id"]
    # 项目已有测点
    client.post(f"/api/v1/projects/{pid}/points", json={"point_no": "JJC1", "point_type": "settlement"}).json()

    data = _make_xlsx_no_date([("JJC1", 4.907, 4.908)])
    resp = _upload(client, pid, data)
    assert resp.status_code == 200
    d = resp.json()
    assert len(d["no_date_rows"]) == 1
    assert d["no_date_rows"][0]["obs_date"] is None
    assert len(d["valid_rows"]) == 0  # 无日期不被当作有效行

    bid = d["batch_id"]
    # 场景A：不补日期确认 → 0 条导入（行被安全跳过，不崩溃；批次转为 confirmed）
    row = {k: v for k, v in d["no_date_rows"][0].items() if k != "error" and k != "suggested"}
    row.pop("error_kind", None)
    r = client.post(f"/api/v1/import/batches/{bid}/confirm", json={"valid_rows": [row]})
    assert r.status_code == 200
    assert r.json()["imported_rows"] == 0

    # 场景B：补日期后确认（新批次）→ 导入成功
    data2 = _make_xlsx_no_date([("JJC1", 4.907, 4.908)])
    d2 = _upload(client, pid, data2).json()
    bid2 = d2["batch_id"]
    row2 = {k: v for k, v in d2["no_date_rows"][0].items() if k != "error" and k != "suggested"}
    row2.pop("error_kind", None)
    row2["obs_date"] = "2026-08-01"
    r2 = client.post(f"/api/v1/import/batches/{bid2}/confirm", json={"valid_rows": [row2]})
    assert r2.status_code == 200
    assert r2.json()["imported_rows"] == 1
    # 数据确实入库（V0 即首期读数 4908mm）
    recs = client.get(f"/api/v1/projects/{pid}/record-dates").json()
    assert "2026-08-01" in recs["dates"]
